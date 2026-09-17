# -*- coding: utf-8 -*-
"""
SGNL / СОД Project Setup GUI
--------------------------------------------------
Утилита для создания папочной структуры, синхронизации прав на папки
и создания пользовательских атрибутов проекта docs.sgnl.pro.

Файлы, под которые написана версия:
    1) Excel с выбираемым пользователем листом: колонки "Уровень N" + колонки ролей.
    2) Дизайн/UX перенесён из приложенного Larix Platform проекта: карточки, вкладки, тема и иконки.

Зависимости:
    pip install requests openpyxl PySide6
Если используется PyQt5:
    pip install requests openpyxl PyQt5

Важно по авторизации:
    Сначала авторизация выполняется фоновыми HTTP-запросами через requests.Session.
    Если CAPTCHA не требуется, никаких веб-страниц пользователь не видит.
    Если SGNL требует CAPTCHA, программа показывает защищённую форму авторизации
    во встроенном Qt WebEngine прямо внутри основного окна. Внешний Chrome/Edge
    не запускается. После прохождения CAPTCHA OAuth PKCE автоматически выдаёт
    access_token для клиента docs_web.
    Hub вручную открывать не нужно, но hub.sgnl.pro остаётся OAuth-сервером.

Логика прав по легенде Excel:
    пусто      -> не трогать эту роль на этой папке
    -          -> снять существующую запись права через DELETE /api/permissions
    П          -> read
    С          -> read + download
    З          -> read + download + create
    Р          -> read + download + create + update + delete

Основные API-маршруты:
    GET  {docs}/api/folders/project/{projectId}/tree
    PUT  {docs}/api/folders
    GET  {docs}/api/permissions/tree?projectId={projectId}
    GET  {hub}/api/v1/hub/companies/{companyId}/custom/roles
    GET  {hub}/api/v1/hub/companies/{companyId}/users
    OAuth использует client_id docs_web и scopes из реального DOCS web-клиента; hub-запросы выполняются с browser-like заголовками.
    POST   {docs}/api/permissions/folders/{folderId}/upsert
    DELETE {docs}/api/permissions  {projectId, permissionIds}
    POST   {docs}/api/attributes/types/find
    PUT    {docs}/api/attributes/types
"""

from __future__ import annotations

import base64
import hashlib
import json
import os
import re
import secrets
import sys
import traceback
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Sequence, Tuple
from urllib.parse import parse_qs, urlencode, urlparse

import requests
import urllib3
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

WEBENGINE_AVAILABLE = False
QtWebEngineWidgets = None

try:
    from PySide6 import QtCore, QtGui, QtWidgets

    Signal = QtCore.Signal
    Slot = QtCore.Slot
    try:
        from PySide6 import QtWebEngineWidgets as _QtWebEngineWidgets

        QtWebEngineWidgets = _QtWebEngineWidgets
        WEBENGINE_AVAILABLE = True
    except ImportError:
        pass
except ImportError:  # pragma: no cover
    from PyQt5 import QtCore, QtGui, QtWidgets

    Signal = QtCore.pyqtSignal
    Slot = QtCore.pyqtSlot
    try:
        from PyQt5 import QtWebEngineWidgets as _QtWebEngineWidgets

        QtWebEngineWidgets = _QtWebEngineWidgets
        WEBENGINE_AVAILABLE = True
    except ImportError:
        pass

from sgnl_ui_theme import (
    ACCENT,
    ACCENT_HOVER,
    STATUS_DANGER,
    STATUS_SUCCESS,
    ThemeToggle,
    asset_path,
    build_qss,
    themed_icon,
    shared_asset_dir,
    initial_dark_theme,
    persist_dark_theme,
    apply_windows_titlebar_theme,
)
from shared.ui_components import (
    NoWheelTabBar,
    configure_preview_table,
    fit_preview_height,
    make_details_button,
    show_table_details,
    add_standard_header_controls,
    install_status_bar,
    mark_destructive_buttons,
)
from shared.theme_core import shared_style_overrides, install_window_state


# -----------------------------------------------------------------------------
# Runtime defaults. Company/project are intentionally empty until login.
# -----------------------------------------------------------------------------
DEFAULT_DOCS_URL = "https://docs.sgnl.pro"
DEFAULT_HUB_URL = "https://hub.sgnl.pro"
DEFAULT_PROJECT_ID = ""
DEFAULT_COMPANY_ID = ""
DEFAULT_COMPANY_NAME = ""
DEFAULT_PROJECT_NAME = ""

OAUTH_CLIENT_ID = "docs_web"
OAUTH_SCOPE = (
    "offline_access cl.front cl.sys hub.r hub.e "
    "doc.c doc.r doc.u doc.d doc.e doc.x "
    "conv.x conv.r cnr.r cnr.e lg.c"
)

# Browser User-Agent for OAuth and API requests.
# It must exist before SgnlClient methods build browser-like headers.
BROWSER_UA = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
    "AppleWebKit/537.36 (KHTML, like Gecko) "
    "Chrome/124.0.0.0 Safari/537.36"
)

VERIFY_SSL_CERTIFICATE = True
if not VERIFY_SSL_CERTIFICATE:
    urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

RIGHT_ORDER = ["П", "С", "З", "Р"]
RIGHT_LABELS = {
    "-": "Снять права",
    "П": "Просмотр",
    "С": "Скачивание",
    "З": "Загрузка",
    "Р": "Редактирование",
}
RIGHT_FLAGS: Dict[str, Dict[str, bool]] = {
    "П": {"read": True, "download": False, "create": False, "update": False, "delete": False},
    "С": {"read": True, "download": True, "create": False, "update": False, "delete": False},
    "З": {"read": True, "download": True, "create": True, "update": False, "delete": False},
    "Р": {"read": True, "download": True, "create": True, "update": True, "delete": True},
}
ZERO_FLAGS = {"read": False, "download": False, "create": False, "update": False, "delete": False}

RIGHT_ALIASES = {
    "п": "П",
    "p": "П",
    "read": "П",
    "view": "П",
    "просмотр": "П",
    "чтение": "П",
    "с": "С",
    "c": "С",
    "download": "С",
    "скачивание": "С",
    "скачать": "С",
    "з": "З",
    "z": "З",
    "upload": "З",
    "create": "З",
    "создание": "З",
    "загрузка": "З",
    "загрузить": "З",
    "р": "Р",
    "r": "Р",
    "edit": "Р",
    "update": "Р",
    "modify": "Р",
    "редактирование": "Р",
    "изменение": "Р",
    "изменить": "Р",
    "полный доступ": "Р",
}


# -----------------------------------------------------------------------------
# Text helpers
# -----------------------------------------------------------------------------
def clean_text(value: object) -> str:
    if value is None:
        return ""
    text = str(value).replace("\xa0", " ").replace("\r", " ").replace("\n", " ").strip()
    text = re.sub(r"\s+", " ", text)
    if text.casefold() in {"none", "nan", "null"}:
        return ""
    return text


def normalize_text(value: object) -> str:
    text = clean_text(value)
    text = unicodedata.normalize("NFKC", text)
    text = text.replace("ё", "е").replace("Ё", "Е")
    text = text.casefold()
    text = re.sub(r"\s+", " ", text).strip()
    return text


def normalize_path_parts(parts: Sequence[object]) -> str:
    return "/".join(normalize_text(part) for part in parts if clean_text(part))


def display_path(parts: Sequence[object]) -> str:
    return " / ".join(clean_text(part) for part in parts if clean_text(part))


def best_close_matches(value: str, variants: Iterable[str], limit: int = 5, cutoff: float = 0.68) -> List[str]:
    import difflib

    norm_variants = list(variants)
    return difflib.get_close_matches(value, norm_variants, n=limit, cutoff=cutoff)


def bearer_header(token: str) -> str:
    token = clean_text(token)
    if not token:
        return ""
    if token.lower().startswith("bearer "):
        return token
    return "Bearer " + token


def decode_jwt_payload(token: str) -> dict:
    """Диагностически декодирует payload JWT без проверки подписи.

    Нужен только для лога: client_id/scopes сразу показывают,
    каким web-клиентом получен token.
    """
    token = clean_text(token)
    if token.lower().startswith("bearer "):
        token = token.split(" ", 1)[1].strip()
    parts = token.split(".")
    if len(parts) < 2:
        return {}
    payload = parts[1]
    payload += "=" * ((4 - len(payload) % 4) % 4)
    try:
        return json.loads(base64.urlsafe_b64decode(payload.encode("ascii")).decode("utf-8"))
    except Exception:
        return {}


def base64url_no_padding(data: bytes) -> str:
    return base64.urlsafe_b64encode(data).decode("ascii").rstrip("=")


def make_pkce_pair() -> Tuple[str, str]:
    # RFC 7636: verifier 43..128 chars. token_urlsafe(64) is normally 86 chars.
    verifier = secrets.token_urlsafe(64)[:128]
    challenge = base64url_no_padding(hashlib.sha256(verifier.encode("ascii")).digest())
    return verifier, challenge


PERMISSION_FLAG_KEYS = ("read", "download", "create", "update", "delete")


def permission_source_id(item: dict) -> str:
    """Возвращает ID роли/пользователя из записи дерева прав."""
    if not isinstance(item, dict):
        return ""
    direct = clean_text(item.get("sourceId") or item.get("roleId") or item.get("userId"))
    if direct:
        return direct
    source = item.get("source")
    if isinstance(source, dict):
        return clean_text(source.get("id") or source.get("sourceId"))
    return ""


def permission_flag_value(item: dict, key: str) -> bool:
    """Читает флаг права с поддержкой плоского и вложенного ответа API."""
    if not isinstance(item, dict):
        return False
    if key in item:
        return bool(item.get(key))
    nested = item.get("permissions") or item.get("flags")
    if isinstance(nested, dict):
        return bool(nested.get(key))
    return False


def permission_flags_equal(existing: dict, desired: Dict[str, bool]) -> bool:
    """True только при полном совпадении всех пяти флагов прав.

    Любое отличие считается изменением. Поэтому более высокий или более низкий
    текущий уровень будет перезаписан точным набором флагов из Excel.
    """
    if not isinstance(existing, dict):
        return False
    return all(
        permission_flag_value(existing, key) == bool(desired.get(key))
        for key in PERMISSION_FLAG_KEYS
    )


def index_permissions_by_source(permissions: Sequence[dict]) -> Dict[str, dict]:
    result: Dict[str, dict] = {}
    for item in permissions or []:
        source_id = permission_source_id(item)
        if source_id:
            result[source_id] = item
    return result


# -----------------------------------------------------------------------------
# Data models
# -----------------------------------------------------------------------------
@dataclass
class PermissionFlags:
    code: str
    flags: Dict[str, bool]

    @property
    def is_remove(self) -> bool:
        return self.code == "-"

    @property
    def label(self) -> str:
        return RIGHT_LABELS.get(self.code, self.code)

    @property
    def text(self) -> str:
        if not self.code:
            return ""
        return f"{self.code} ({self.label})"


@dataclass
class PermissionEntry:
    row_number: int
    path_parts: List[str]
    permissions: Dict[str, PermissionFlags]

    @property
    def path_text(self) -> str:
        return display_path(self.path_parts)


@dataclass
class FolderItem:
    id: str
    name: str
    path_parts: List[str]
    permissions: List[dict] = field(default_factory=list)

    @property
    def path_text(self) -> str:
        return display_path(self.path_parts)

    @property
    def norm_path(self) -> str:
        return normalize_path_parts(self.path_parts)


@dataclass
class Principal:
    name: str
    id: str
    source_type: str = "Role"  # Role / User


@dataclass
class PrincipalResolution:
    requested_name: str
    principal: Optional[Principal] = None
    message: str = ""

    @property
    def ok(self) -> bool:
        return bool(self.principal and self.principal.id)


@dataclass
class PlanRow:
    entry: PermissionEntry
    folder: Optional[FolderItem]
    status: str
    message: str
    principals: Dict[str, PrincipalResolution] = field(default_factory=dict)
    pending_roles: List[str] = field(default_factory=list)
    unchanged_roles: List[str] = field(default_factory=list)

    @property
    def is_valid(self) -> bool:
        return (
            self.status in {"К применению", "Без изменений"}
            and self.folder is not None
            and all(res.ok for res in self.principals.values())
        )

    @property
    def can_apply(self) -> bool:
        return self.is_valid and bool(self.pending_roles)

    @property
    def no_changes(self) -> bool:
        return self.is_valid and not self.pending_roles

    @property
    def permission_count(self) -> int:
        return len(self.entry.permissions)

    @property
    def pending_permission_count(self) -> int:
        return len(self.pending_roles)

    @property
    def unchanged_permission_count(self) -> int:
        return len(self.unchanged_roles)


@dataclass
class FolderPathEntry:
    row_number: int
    path_parts: List[str]

    @property
    def path_text(self) -> str:
        return display_path(self.path_parts)


@dataclass
class FolderCreatePlanRow:
    row_number: int
    path_parts: List[str]
    status: str
    parent_path: str
    folder_id: str = ""
    message: str = ""

    @property
    def path_text(self) -> str:
        return display_path(self.path_parts)

    @property
    def name(self) -> str:
        return self.path_parts[-1] if self.path_parts else ""

    @property
    def can_create(self) -> bool:
        return self.status == "Создать"


# -----------------------------------------------------------------------------
# Excel template generation
# -----------------------------------------------------------------------------
TEMPLATE_ORANGE = "F7921E"
TEMPLATE_DARK = "2F2F2F"
TEMPLATE_LIGHT = "FFF4E8"
TEMPLATE_BORDER = "D9D9D9"
TEMPLATE_MUTED = "666666"


def _template_header_style(cell) -> None:
    cell.fill = PatternFill("solid", fgColor=TEMPLATE_ORANGE)
    cell.font = Font(color="FFFFFF", bold=True)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin = Side(style="thin", color=TEMPLATE_BORDER)
    cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)


def _template_body_style(cell) -> None:
    thin = Side(style="thin", color=TEMPLATE_BORDER)
    cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    cell.alignment = Alignment(vertical="top", wrap_text=True)


def _prepare_template_sheet(ws, title: str, subtitle: str, headers: Sequence[str], widths: Sequence[float]) -> None:
    last_col = max(1, len(headers))
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    ws.cell(1, 1, title)
    ws.cell(1, 1).font = Font(size=15, bold=True, color="FFFFFF")
    ws.cell(1, 1).fill = PatternFill("solid", fgColor=TEMPLATE_DARK)
    ws.cell(1, 1).alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 28

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
    ws.cell(2, 1, subtitle)
    ws.cell(2, 1).font = Font(size=9, color=TEMPLATE_MUTED, italic=True)
    ws.cell(2, 1).fill = PatternFill("solid", fgColor=TEMPLATE_LIGHT)
    ws.cell(2, 1).alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[2].height = 34

    for col, header in enumerate(headers, 1):
        cell = ws.cell(3, col, header)
        _template_header_style(cell)
    ws.row_dimensions[3].height = 26
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:{ws.cell(3, last_col).coordinate}"

    for idx, width in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(3, idx).column_letter].width = width

    ws.sheet_view.showGridLines = False


def _add_template_notes_sheet(wb: Workbook, rows: Sequence[Tuple[str, str]]) -> None:
    ws = wb.create_sheet("Инструкция")
    ws.sheet_view.showGridLines = False
    ws["A1"] = "Как заполнять шаблон"
    ws["A1"].font = Font(size=15, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor=TEMPLATE_DARK)
    ws["A1"].alignment = Alignment(vertical="center")
    ws.merge_cells("A1:B1")
    ws.row_dimensions[1].height = 28
    ws["A3"] = "Поле / правило"
    ws["B3"] = "Описание"
    _template_header_style(ws["A3"])
    _template_header_style(ws["B3"])
    for row_idx, (name, description) in enumerate(rows, 4):
        ws.cell(row_idx, 1, name)
        ws.cell(row_idx, 2, description)
        _template_body_style(ws.cell(row_idx, 1))
        _template_body_style(ws.cell(row_idx, 2))
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 90
    ws.freeze_panes = "A4"


def build_excel_template(template_kind: str, target_path: str) -> None:
    """Creates one of the three user-facing SGNL Excel templates."""
    kind = clean_text(template_kind).casefold()
    wb = Workbook()
    ws = wb.active

    if kind == "permissions":
        ws.title = "2. Ролевая матрица"
        headers = ["Уровень 1", "Уровень 2", "Уровень 3", "Роль 1", "Роль 2", "Роль 3"]
        _prepare_template_sheet(
            ws,
            "SGNL — шаблон ролевой матрицы",
            "Колонки «Роль 1/2/3» замените точными названиями ролей из SGNL. Значения прав: П, С, З, Р, - или пусто.",
            headers,
            [28, 28, 34, 22, 22, 22],
        )
        examples = [
            ["01. Проектирование", "АР", "Рабочая документация", "Р", "П", ""],
            ["", "КР", "Рабочая документация", "С", "П", ""],
            ["", "ОВ", "Рабочая документация", "З", "", "П"],
            ["02. Общая папка", "Обмен", "Входящие", "П", "С", "-"],
        ]
        for r, row in enumerate(examples, 4):
            for c, value in enumerate(row, 1):
                ws.cell(r, c, value)
                _template_body_style(ws.cell(r, c))
        validation = DataValidation(type="list", formula1='"П,С,З,Р,-"', allow_blank=True)
        validation.error = "Допустимые значения: П, С, З, Р, - или пустая ячейка."
        validation.errorTitle = "Некорректное право"
        validation.prompt = "П — просмотр; С — скачивание; З — загрузка; Р — полный доступ; - — снять права."
        validation.promptTitle = "Права SGNL"
        validation.showErrorMessage = True
        validation.showInputMessage = True
        ws.add_data_validation(validation)
        validation.add("D4:F503")
        for col in range(4, 7):
            ws.cell(3, col).comment = Comment("Замените название на точное имя роли из SGNL.", "SGNL Platform")
        _add_template_notes_sheet(wb, [
            ("Уровень N", "Иерархия папок. Можно добавлять дополнительные колонки «Уровень 4», «Уровень 5» и т.д. перед колонками ролей."),
            ("Название роли", "Должно точно совпадать с названием роли или пользователя, доступного проекту SGNL."),
            ("П", "Просмотр: read."),
            ("С", "Просмотр + скачивание: read + download."),
            ("З", "Просмотр + скачивание + создание: read + download + create."),
            ("Р", "Полный доступ: read + download + create + update + delete."),
            ("-", "Удалить существующую запись права для этой роли на папке."),
            ("Пусто", "Не менять право этой роли на этой папке."),
        ])

    elif kind == "folders":
        ws.title = "Папочная структура"
        headers = ["Уровень 1", "Уровень 2", "Уровень 3", "Уровень 4", "Уровень 5"]
        _prepare_template_sheet(
            ws,
            "SGNL — шаблон папочной структуры",
            "Заполняйте иерархию слева направо. Пустая ячейка наследует предыдущий заполненный уровень; новые папки создаются только если их ещё нет.",
            headers,
            [30, 30, 34, 34, 34],
        )
        examples = [
            ["01. Проектирование", "АР", "Рабочая документация", "", ""],
            ["", "", "Исходные данные", "", ""],
            ["", "КР", "Рабочая документация", "", ""],
            ["02. Общая папка", "Обмен", "Входящие", "От заказчика", ""],
            ["", "", "Исходящие", "Проектировщику", ""],
        ]
        for r, row in enumerate(examples, 4):
            for c, value in enumerate(row, 1):
                ws.cell(r, c, value)
                _template_body_style(ws.cell(r, c))
        _add_template_notes_sheet(wb, [
            ("Уровень 1", "Корневой уровень создаваемой структуры внутри проекта."),
            ("Уровень 2..N", "Вложенные папки. Можно добавить столько колонок «Уровень N», сколько требуется."),
            ("Пустая ячейка", "Сохраняет последнее значение соответствующего родительского уровня до тех пор, пока в этой колонке не появится новое значение."),
            ("Создание", "Программа предварительно сверяет дерево SGNL и создаёт только отсутствующие папки; существующие не удаляет."),
        ])

    elif kind == "attributes":
        ws.title = "Атрибуты"
        headers = ["Название", "Тип", "Область", "Обязательный", "Значения списка"]
        _prepare_template_sheet(
            ws,
            "SGNL — шаблон атрибутов",
            "Одна строка = один атрибут. Для типа «Список» варианты указываются через ;. Допустимые типы: Текст, Да/Нет, Дата, Список.",
            headers,
            [34, 20, 28, 18, 55],
        )
        examples = [
            ["Шифр документа", "Текст", "Элемент / файл", "Нет", ""],
            ["На согласовании", "Да/Нет", "Элемент / файл", "Нет", ""],
            ["Дата выпуска", "Дата", "Элемент / файл", "Нет", ""],
            ["Раздел", "Список", "Папка", "Да", "АР; КР; ОВ; ВК"],
        ]
        for r, row in enumerate(examples, 4):
            for c, value in enumerate(row, 1):
                ws.cell(r, c, value)
                _template_body_style(ws.cell(r, c))

        type_validation = DataValidation(type="list", formula1='"Текст,Да/Нет,Дата,Список"', allow_blank=False)
        area_validation = DataValidation(type="list", formula1='"Без области,Элемент / файл,Папка"', allow_blank=True)
        required_validation = DataValidation(type="list", formula1='"Да,Нет"', allow_blank=False)
        for validation in (type_validation, area_validation, required_validation):
            validation.showErrorMessage = True
            ws.add_data_validation(validation)
        type_validation.add("B4:B503")
        area_validation.add("C4:C503")
        required_validation.add("D4:D503")
        ws["E3"].comment = Comment("Только для типа «Список». Несколько вариантов разделяйте точкой с запятой (;).", "SGNL Platform")
        _add_template_notes_sheet(wb, [
            ("Название", "Название создаваемого атрибута. Активный дубликат с тем же именем программа не должна создавать повторно."),
            ("Тип", "Текст / Да/Нет / Дата / Список."),
            ("Область", "Без области, Элемент / файл или Папка."),
            ("Обязательный", "Да или Нет."),
            ("Значения списка", "Заполняется только для типа «Список». Значения перечисляются через ; в требуемом порядке."),
        ])

    else:
        raise ValueError(f"Неизвестный тип Excel-шаблона: {template_kind}")

    # Main working sheet stays first, so selecting/opening the template immediately lands on data.
    wb.active = 0
    wb.save(target_path)


# -----------------------------------------------------------------------------
# Excel parser
# -----------------------------------------------------------------------------
def right_from_cell(value: object) -> Optional[PermissionFlags]:
    text = clean_text(value)
    if not text:
        # Пустая ячейка означает: роль на этой папке не трогать.
        return None
    if text in {"-", "—", "–"}:
        # Дефис — отдельная команда: удалить существующую запись права.
        # Это НЕ upsert пяти False-флагов: web-клиент SGNL использует
        # DELETE /api/permissions с permissionIds.
        return PermissionFlags("-", dict(ZERO_FLAGS))
    if text in {"нет", "Нет", "0"}:
        return None

    # Keep one-letter Cyrillic codes intact, but also accept words / English aliases.
    tokens = [clean_text(part) for part in re.split(r"[,;\n\r/]+", text) if clean_text(part)]
    if not tokens:
        tokens = [text]

    codes: List[str] = []
    for token in tokens:
        norm = normalize_text(token)
        code = RIGHT_ALIASES.get(norm)
        if not code and len(norm) == 1:
            code = RIGHT_ALIASES.get(norm)
        if code and code not in codes:
            codes.append(code)

    if not codes:
        return None

    # If several rights are written, use the strongest level from the legend.
    strongest = max(codes, key=lambda code: RIGHT_ORDER.index(code) if code in RIGHT_ORDER else -1)
    return PermissionFlags(strongest, dict(RIGHT_FLAGS[strongest]))


class PermissionExcelParser:
    def __init__(self, excel_path: str, sheet_name: str = ""):
        self.excel_path = excel_path
        self.sheet_name = clean_text(sheet_name)
        self.resolved_sheet_name = ""

    @staticmethod
    def list_sheet_names(excel_path: str) -> List[str]:
        if not os.path.exists(excel_path):
            raise FileNotFoundError(f"Excel-файл не найден: {excel_path}")
        workbook = load_workbook(excel_path, read_only=True, data_only=True)
        try:
            return list(workbook.sheetnames)
        finally:
            workbook.close()

    def _read_matrix(self) -> Tuple[List[List[object]], int, List[Tuple[int, int]], List[Tuple[int, str]]]:
        """Читает лист одним последовательным проходом.

        Для read_only workbook нельзя многократно вызывать sheet.cell(): каждый
        случайный доступ может заново проходить XML-поток и на нескольких сотнях
        строк превращается в минуты ожидания.
        """
        if not os.path.exists(self.excel_path):
            raise FileNotFoundError(f"Excel-файл не найден: {self.excel_path}")

        workbook = load_workbook(self.excel_path, read_only=True, data_only=True)
        try:
            if self.sheet_name:
                if self.sheet_name not in workbook.sheetnames:
                    raise RuntimeError(
                        f"Лист '{self.sheet_name}' не найден. Доступные листы: "
                        + ", ".join(workbook.sheetnames)
                    )
                sheet = workbook[self.sheet_name]
            else:
                sheet = workbook["2. Ролевая матрица"] if "2. Ролевая матрица" in workbook.sheetnames else workbook.active
            self.resolved_sheet_name = clean_text(sheet.title)
            rows = [list(row) for row in sheet.iter_rows(values_only=True)]
        finally:
            workbook.close()

        if not rows:
            raise RuntimeError("Excel-лист пуст.")
        header_index = self._find_header_index(rows)
        header = rows[header_index]

        level_columns: List[Tuple[int, int]] = []
        last_level_col = -1
        for col_idx, value in enumerate(header):
            text = clean_text(value)
            match = re.fullmatch(r"Уровень\s*(\d+)", text, flags=re.IGNORECASE)
            if match:
                level_columns.append((int(match.group(1)), col_idx))
                last_level_col = max(last_level_col, col_idx)
        level_columns.sort(key=lambda pair: pair[0])
        if not level_columns:
            raise RuntimeError("В Excel не найдены колонки вида 'Уровень 2', 'Уровень 3', ...")

        role_columns: List[Tuple[int, str]] = []
        for col_idx in range(last_level_col + 1, len(header)):
            role_name = clean_text(header[col_idx])
            if not role_name:
                continue
            if normalize_text(role_name).startswith("легенда"):
                continue
            if re.match(r"^[ПСЗР]-", role_name, flags=re.IGNORECASE):
                continue
            role_columns.append((col_idx, role_name))

        return rows, header_index, level_columns, role_columns

    def parse(self) -> List[PermissionEntry]:
        rows, header_index, level_columns, role_columns = self._read_matrix()
        if not role_columns:
            raise RuntimeError("В Excel не найдены колонки ролей после колонок уровней.")

        path_stack: List[Optional[str]] = [None] * len(level_columns)
        result: List[PermissionEntry] = []
        for row_index in range(header_index + 1, len(rows)):
            row = rows[row_index]
            for stack_idx, (_level_num, col_idx) in enumerate(level_columns):
                value = clean_text(row[col_idx] if col_idx < len(row) else None)
                if value:
                    path_stack[stack_idx] = value
                    for clear_idx in range(stack_idx + 1, len(path_stack)):
                        path_stack[clear_idx] = None

            path_parts = [part for part in path_stack if part]
            if not path_parts:
                continue

            permissions: Dict[str, PermissionFlags] = {}
            for col_idx, role_name in role_columns:
                cell_value = row[col_idx] if col_idx < len(row) else None
                flags = right_from_cell(cell_value)
                if flags:
                    permissions[role_name] = flags
            if permissions:
                result.append(
                    PermissionEntry(
                        row_number=row_index + 1,
                        path_parts=list(path_parts),
                        permissions=permissions,
                    )
                )

        if not result:
            raise RuntimeError("В Excel не найдено ни одной строки с правами.")
        return result

    def parse_paths(self) -> List[FolderPathEntry]:
        """Читает папочные строки независимо от наличия прав."""
        rows, header_index, level_columns, _role_columns = self._read_matrix()
        path_stack: List[Optional[str]] = [None] * len(level_columns)
        result: List[FolderPathEntry] = []

        for row_index in range(header_index + 1, len(rows)):
            row = rows[row_index]
            changed = False
            for stack_idx, (_level_num, col_idx) in enumerate(level_columns):
                value = clean_text(row[col_idx] if col_idx < len(row) else None)
                if not value:
                    continue
                changed = True
                path_stack[stack_idx] = value
                for clear_idx in range(stack_idx + 1, len(path_stack)):
                    path_stack[clear_idx] = None
            if not changed:
                continue
            path_parts = [part for part in path_stack if part]
            if path_parts:
                result.append(FolderPathEntry(row_number=row_index + 1, path_parts=list(path_parts)))

        if not result:
            raise RuntimeError("В Excel не найдено ни одной строки папочной структуры.")
        return result

    @staticmethod
    def _find_header_index(rows: Sequence[Sequence[object]]) -> int:
        for row_index, row in enumerate(rows[:30]):
            count = sum(
                1
                for value in row
                if re.fullmatch(r"Уровень\s*\d+", clean_text(value), flags=re.IGNORECASE)
            )
            if count >= 2:
                return row_index
        raise RuntimeError("Не найдена строка заголовков с колонками 'Уровень N'.")


# -----------------------------------------------------------------------------
# API client
# -----------------------------------------------------------------------------
class CaptchaRequiredError(RuntimeError):
    """SGNL explicitly required an interactive CAPTCHA challenge."""


class SgnlClient:
    def __init__(self, docs_url: str, hub_url: str, token: str = "", timeout: int = 120, cookies: Optional[Dict[str, str]] = None):
        self.docs_url = docs_url.strip().rstrip("/") or DEFAULT_DOCS_URL
        self.hub_url = hub_url.strip().rstrip("/") or DEFAULT_HUB_URL
        self.token = clean_text(token)
        # Отдельные таймауты: соединение должно установиться быстро, но большое дерево
        # папок разрешаем читать дольше.
        self.timeout = (20, max(60, int(timeout)))
        self.session = requests.Session()
        self.session.verify = VERIFY_SSL_CERTIFICATE
        # В v3 системный proxy/VPN был принудительно отключён через trust_env=False.
        # Из-за этого браузер открывал SGNL, а Python получал WinError 10060.
        # requests/urllib теперь используют системные и переменные окружения proxy.
        self.session.trust_env = True
        retry = Retry(
            total=3,
            connect=3,
            read=2,
            status=2,
            backoff_factor=0.7,
            status_forcelist=(429, 500, 502, 503, 504),
            allowed_methods=frozenset({"GET", "HEAD", "OPTIONS"}),
            raise_on_status=False,
        )
        adapter = HTTPAdapter(max_retries=retry, pool_connections=10, pool_maxsize=10)
        self.session.mount("https://", adapter)
        self.session.mount("http://", adapter)
        if cookies:
            self.session.cookies.update({clean_text(k): clean_text(v) for k, v in cookies.items() if clean_text(k)})
        # Кэш fallback-цепочки /api/companies/find/current-user.
        # У некоторых аккаунтов /api/v1/hub/companies возвращает 403 сразу после OAuth-входа,
        # хотя список компаний доступен через маршрут текущего пользователя.
        self._company_project_ids: Dict[str, List[str]] = {}

    def _request(self, method: str, url: str, **kwargs) -> requests.Response:
        kwargs.setdefault("timeout", self.timeout)
        try:
            return self.session.request(method, url, **kwargs)
        except (requests.exceptions.ConnectTimeout, requests.exceptions.ProxyError) as exc:
            # ConnectTimeout/ProxyError происходят до получения ответа. Если системный
            # proxy существует, один раз пробуем прямое соединение.
            proxies = requests.utils.get_environ_proxies(url) if self.session.trust_env else {}
            if proxies:
                self.session.trust_env = False
                try:
                    return self.session.request(method, url, **kwargs)
                except requests.exceptions.RequestException as direct_exc:
                    raise RuntimeError(
                        f"Не удалось подключиться к {urlparse(url).netloc} ни через системный proxy, "
                        f"ни напрямую. Проверь VPN/proxy и разрешение для python.exe в брандмауэре. "
                        f"Последняя ошибка: {direct_exc}"
                    ) from direct_exc
            raise RuntimeError(
                f"Не удалось подключиться к {urlparse(url).netloc}. Проверь VPN, системный proxy "
                f"и разрешение для python.exe в брандмауэре. Ошибка: {exc}"
            ) from exc
        except requests.exceptions.ConnectionError as exc:
            # Для изменяющих запросов не делаем автоматический повтор: сервер мог
            # получить запрос, даже если клиент не дождался ответа.
            raise RuntimeError(
                f"Соединение с {urlparse(url).netloc} было разорвано. Повтор не выполнен автоматически, "
                f"чтобы не продублировать изменение. Ошибка: {exc}"
            ) from exc
        except requests.exceptions.RequestException as exc:
            raise RuntimeError(f"Сетевая ошибка при обращении к {url}: {exc}") from exc

    @property
    def headers(self) -> Dict[str, str]:
        # Для docs API нужен Bearer access_token.
        return self.docs_api_headers(include_auth=True)

    def docs_api_headers(self, include_auth: bool = True) -> Dict[str, str]:
        headers = {
            "Accept": "application/json, text/plain, */*",
            "Content-Type": "application/json",
            "Origin": self.docs_url,
            "Referer": self.docs_url + "/",
            "User-Agent": BROWSER_UA,
        }
        if include_auth:
            auth = bearer_header(self.token)
            if auth:
                headers["Authorization"] = auth
        return headers

    def hub_docs_headers(self, include_auth: bool = False, method: str = "GET") -> Dict[str, str]:
        # Точные browser-like заголовки для запросов из docs.sgnl.pro в hub.sgnl.pro.
        # В полном HAR эти hub v1 вызовы идут с Origin/Referer docs и без Content-Type на GET.
        headers = {
            "Accept": "application/json, text/plain, */*",
            "Accept-Language": "en-US,en;q=0.9,ru-RU;q=0.8,ru;q=0.7",
            "Origin": self.docs_url,
            "Referer": self.docs_url + "/",
            "Sec-Fetch-Dest": "empty",
            "Sec-Fetch-Mode": "cors",
            "Sec-Fetch-Site": "same-site",
            "User-Agent": BROWSER_UA,
        }
        if method.upper() != "GET":
            headers["Content-Type"] = "application/json"
        if include_auth:
            auth = bearer_header(self.token)
            if auth:
                headers["Authorization"] = auth
        return headers

    def hub_app_headers(self, include_auth: bool = True, method: str = "POST") -> Dict[str, str]:
        # Маршруты hub-приложения, которые в HAR вызываются с Origin hub.sgnl.pro.
        headers = {
            "Accept": "application/json, text/plain, */*",
            "Accept-Language": "en-US,en;q=0.9,ru-RU;q=0.8,ru;q=0.7",
            "Origin": self.hub_url,
            "Referer": self.hub_url + "/hub",
            "Sec-Fetch-Dest": "empty",
            "Sec-Fetch-Mode": "cors",
            "Sec-Fetch-Site": "same-origin",
            "User-Agent": BROWSER_UA,
        }
        if method.upper() != "GET":
            headers["Content-Type"] = "application/json"
        if include_auth:
            auth = bearer_header(self.token)
            if auth:
                headers["Authorization"] = auth
        return headers

    def _request_with_cookie_fallback(self, method: str, url: str, context: str, **kwargs) -> requests.Response:
        method_upper = method.upper()
        # 1) Как в HAR: docs -> hub, без Authorization.
        response = self._request(
            method,
            url,
            headers=self.hub_docs_headers(include_auth=False, method=method_upper),
            **kwargs,
        )
        if response.status_code not in (401, 403):
            return response

        # 2) Тот же контекст, но с Bearer access_token.
        bearer_response = self._request(
            method,
            url,
            headers=self.hub_docs_headers(include_auth=True, method=method_upper),
            **kwargs,
        )
        return bearer_response if bearer_response.status_code not in (401, 403) else response

    def _auth_origin_headers(self) -> Dict[str, str]:
        return {
            "Accept": "*/*",
            "Content-Type": "application/json",
            "Origin": "https://auth.sgnl.pro",
            "Referer": "https://auth.sgnl.pro/",
            "User-Agent": BROWSER_UA,
        }

    def _hub_form_headers(self, referer: str = "") -> Dict[str, str]:
        headers = {
            "Accept": "application/json, text/plain, */*",
            "Content-Type": "application/x-www-form-urlencoded",
            "Origin": self.hub_url,
            "Referer": referer or (self.hub_url + "/"),
            "User-Agent": BROWSER_UA,
        }
        return headers

    def build_authorization_request(self) -> dict:
        verifier, challenge = make_pkce_pair()
        redirect_uri = self.docs_url + "/callback"
        params = {
            "client_id": OAUTH_CLIENT_ID,
            "redirect_uri": redirect_uri,
            "response_type": "code",
            "scope": OAUTH_SCOPE,
            "state": verifier,
            "code_challenge": challenge,
            "code_challenge_method": "S256",
        }
        return {
            "authorize_url": self.hub_url + "/connect/authorize?" + urlencode(params),
            "verifier": verifier,
            "state": verifier,
            "redirect_uri": redirect_uri,
        }

    def exchange_authorization_code(self, code: str, verifier: str, redirect_uri: str, referer: str = "") -> dict:
        code = clean_text(code)
        verifier = clean_text(verifier)
        redirect_uri = clean_text(redirect_uri)
        if not code or not verifier or not redirect_uri:
            raise RuntimeError("Недостаточно данных для OAuth token exchange.")
        response = self._request(
            "POST",
            self.hub_url + "/connect/token",
            data={
                "grant_type": "authorization_code",
                "client_id": OAUTH_CLIENT_ID,
                "redirect_uri": redirect_uri,
                "code": code,
                "code_verifier": verifier,
            },
            headers=self._hub_form_headers(referer=referer or redirect_uri),
        )
        self._raise_for_response(response, "Не удалось получить OAuth access_token")
        data = response.json()
        token = clean_text(data.get("access_token"))
        if not token:
            raise RuntimeError(f"В ответе /connect/token нет access_token: {data}")
        self.token = token
        return data

    def login_email_password(self, email: str, password: str) -> dict:
        """Получает OAuth access_token без браузера через requests.Session.

        Последовательность соответствует docs_авторизация.har:
        1) GET  /connect/authorize (OAuth PKCE);
        2) GET  страницы auth.sgnl.pro/account/login для web-сессии;
        3) POST /api/v1/auth/is-captcha-challenge-required;
        4) POST /api/v1/auth/login-with-credentials только с email/password;
        5) GET  /connect/authorize/callback;
        6) POST /connect/token.

        В приложенном HAR проверка CAPTCHA возвращает false (тело ответа
        размером 5 байт), поэтому challengeKey в запросе входа отсутствует.
        """
        email = clean_text(email)
        if not email or not password:
            raise RuntimeError("Укажи email и пароль.")

        verifier, challenge = make_pkce_pair()
        redirect_uri = self.docs_url + "/callback"
        state = secrets.token_urlsafe(24)
        params = {
            "client_id": OAUTH_CLIENT_ID,
            "redirect_uri": redirect_uri,
            "response_type": "code",
            "scope": OAUTH_SCOPE,
            "state": state,
            "code_challenge": challenge,
            "code_challenge_method": "S256",
        }

        # 1) Начинаем OAuth. Сервер должен перенаправить на auth.sgnl.pro.
        authorize_url = self.hub_url + "/connect/authorize?" + urlencode(params)
        response = self._request(
            "GET",
            authorize_url,
            allow_redirects=False,
            headers={
                "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
                "Referer": self.docs_url + "/",
                "User-Agent": BROWSER_UA,
            },
        )
        if response.status_code not in (301, 302, 303, 307, 308):
            self._raise_for_response(response, "Не удалось начать OAuth-авторизацию")

        login_page_url = clean_text(response.headers.get("Location"))
        if not login_page_url:
            raise RuntimeError("OAuth-сервер не вернул адрес страницы входа.")

        # 2) Открываем страницу входа в той же requests.Session, чтобы получить
        # необходимые cookies web-сессии. Никакое окно пользователю не показывается.
        login_page = self._request(
            "GET",
            login_page_url,
            allow_redirects=True,
            headers={
                "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
                "Referer": self.docs_url + "/",
                "User-Agent": BROWSER_UA,
            },
        )
        if login_page.status_code >= 400:
            self._raise_for_response(login_page, "Не удалось подготовить сессию входа")

        # 3) Реальная форма сначала спрашивает, требуется ли CAPTCHA для email.
        captcha_response = self._request(
            "POST",
            self.hub_url + "/api/v1/auth/is-captcha-challenge-required",
            json={"email": email},
            headers=self._auth_origin_headers(),
        )
        self._raise_for_response(captcha_response, "Не удалось проверить необходимость CAPTCHA")
        try:
            captcha_required = bool(captcha_response.json())
        except Exception:
            captcha_required = normalize_text(captcha_response.text) in {"true", "1", "yes"}

        if captcha_required:
            raise CaptchaRequiredError(
                "SGNL потребовал CAPTCHA для этой попытки входа."
            )

        # 4) В HAR этот запрос содержит только email и password — challengeKey нет.
        login_response = self._request(
            "POST",
            self.hub_url + "/api/v1/auth/login-with-credentials",
            json={"email": email, "password": password},
            headers=self._auth_origin_headers(),
        )
        if not login_response.ok:
            body = normalize_text(login_response.text)
            if "captcha" in body or "challenge" in body:
                raise CaptchaRequiredError(
                    "SGNL отклонил фоновый вход и запросил CAPTCHA."
                )
            self._raise_for_response(login_response, "Не удалось войти по email/паролю")

        # 5) После успешного login-with-credentials web-сессия уже авторизована.
        callback_url = self.hub_url + "/connect/authorize/callback?" + urlencode(params)
        callback_response = self._request(
            "GET",
            callback_url,
            allow_redirects=False,
            headers={
                "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
                "Referer": "https://auth.sgnl.pro/",
                "User-Agent": BROWSER_UA,
            },
        )
        if callback_response.status_code not in (301, 302, 303, 307, 308):
            self._raise_for_response(callback_response, "Не удалось получить OAuth authorization code")

        location = clean_text(callback_response.headers.get("Location"))
        query = parse_qs(urlparse(location).query)
        returned_state = clean_text(query.get("state", [""])[0])
        if returned_state and returned_state != state:
            raise RuntimeError("OAuth state не совпал; авторизация остановлена.")

        code = clean_text(query.get("code", [""])[0])
        if not code:
            error = clean_text(query.get("error", [""])[0])
            description = clean_text(query.get("error_description", [""])[0])
            details = ": ".join(part for part in (error, description) if part)
            raise RuntimeError(
                "OAuth callback не вернул code"
                + (f": {details}" if details else ". Возможно, изменилась схема авторизации.")
            )

        # 6) Меняем одноразовый authorization code на access_token.
        return self.exchange_authorization_code(
            code=code,
            verifier=verifier,
            redirect_uri=redirect_uri,
            referer=location or callback_url,
        )

    def _raise_for_response(self, response: requests.Response, context: str) -> None:
        if response.ok:
            return
        text = response.text or ""
        if len(text) > 2500:
            text = text[:2500] + "..."
        raise RuntimeError(f"{context}. HTTP {response.status_code}: {text}")

    def get_companies(self) -> List[dict]:
        response = self._request_with_cookie_fallback("GET", f"{self.hub_url}/api/v1/hub/companies", "Не удалось загрузить компании")
        if response.status_code == 403:
            # В расширенном HAR после входа компания подтягивается так:
            # POST /api/companies/find/current-user {"deleted": false}.
            # Поэтому 403 на общем списке компаний не считаем фатальной ошибкой.
            return self.get_current_user_companies()
        self._raise_for_response(response, "Не удалось загрузить компании")
        data = response.json()
        if not isinstance(data, list):
            raise RuntimeError(f"Ожидался список компаний, получено: {type(data).__name__}")
        self._remember_company_project_ids(data)
        return data

    def get_current_user_companies(self) -> List[dict]:
        response = self._request(
            "POST",
            f"{self.hub_url}/api/companies/find/current-user",
            headers=self.hub_app_headers(include_auth=True, method="POST"),
            json={"deleted": False},
        )
        self._raise_for_response(response, "Не удалось загрузить компании текущего пользователя")
        data = response.json()
        if not isinstance(data, list):
            raise RuntimeError(f"Ожидался список компаний текущего пользователя, получено: {type(data).__name__}")

        companies: List[dict] = []
        for row in data:
            if not isinstance(row, dict):
                continue
            company = row.get("company") if isinstance(row.get("company"), dict) else row
            company_id = clean_text(company.get("id"))
            if not company_id:
                continue
            item = dict(company)
            if row.get("role") and not item.get("applicationRole"):
                item["applicationRole"] = row.get("role")
            project_ids = row.get("projectIds") or row.get("project_ids") or item.get("projectIds") or []
            if isinstance(project_ids, list):
                item["projectIds"] = [clean_text(x) for x in project_ids if clean_text(x)]
                self._company_project_ids[company_id] = item["projectIds"]
            companies.append(item)
        return companies

    def _remember_company_project_ids(self, companies: Sequence[dict]) -> None:
        for company in companies:
            if not isinstance(company, dict):
                continue
            company_id = clean_text(company.get("id"))
            project_ids = company.get("projectIds") or company.get("project_ids") or []
            if company_id and isinstance(project_ids, list):
                self._company_project_ids[company_id] = [clean_text(x) for x in project_ids if clean_text(x)]

    def get_project(self, project_id: str) -> dict:
        project_id = clean_text(project_id)
        response = self._request_with_cookie_fallback(
            "POST",
            f"{self.hub_url}/api/v1/projects/get-project",
            f"Не удалось загрузить проект {project_id}",
            params={"projectId": project_id},
        )
        self._raise_for_response(response, f"Не удалось загрузить проект {project_id}")
        data = response.json()
        if not isinstance(data, dict):
            raise RuntimeError(f"Ожидался объект проекта, получено: {type(data).__name__}")
        return data

    def get_company_projects(self, company_id: str) -> List[dict]:
        response = self._request_with_cookie_fallback(
            "GET",
            f"{self.hub_url}/api/v1/hub/companies/{company_id}/projects",
            "Не удалось загрузить проекты компании",
        )
        if response.status_code == 403:
            # Fallback из /api/companies/find/current-user: там есть projectIds.
            if company_id not in self._company_project_ids:
                self.get_current_user_companies()
            project_ids = self._company_project_ids.get(clean_text(company_id), [])
            if project_ids:
                projects = []
                for project_id in project_ids:
                    try:
                        projects.append(self.get_project(project_id))
                    except Exception:
                        # Один битый/недоступный проект не должен ломать загрузку списка целиком.
                        continue
                return projects
        self._raise_for_response(response, "Не удалось загрузить проекты компании")
        data = response.json()
        if not isinstance(data, list):
            raise RuntimeError(f"Ожидался список проектов, получено: {type(data).__name__}")
        return data

    def get_custom_roles(self, company_id: str) -> List[dict]:
        response = self._request_with_cookie_fallback(
            "GET",
            f"{self.hub_url}/api/v1/hub/companies/{company_id}/custom/roles",
            "Не удалось загрузить роли компании",
        )
        self._raise_for_response(response, "Не удалось загрузить роли компании")
        data = response.json()
        if not isinstance(data, list):
            raise RuntimeError(f"Ожидался список ролей, получено: {type(data).__name__}")
        return data

    def get_company_users(self, company_id: str) -> List[dict]:
        response = self._request_with_cookie_fallback(
            "GET",
            f"{self.hub_url}/api/v1/hub/companies/{company_id}/users",
            "Не удалось загрузить пользователей компании",
        )
        self._raise_for_response(response, "Не удалось загрузить пользователей компании")
        data = response.json()
        if not isinstance(data, list):
            raise RuntimeError(f"Ожидался список пользователей, получено: {type(data).__name__}")
        return data

    def get_project_roles(self, project_id: str) -> List[dict]:
        response = self._request_with_cookie_fallback(
            "GET",
            f"{self.hub_url}/api/v1/hub/projects/{project_id}/roles",
            "Не удалось загрузить роли проекта",
        )
        self._raise_for_response(response, "Не удалось загрузить роли проекта")
        data = response.json()
        if not isinstance(data, list):
            raise RuntimeError(f"Ожидался список ролей проекта, получено: {type(data).__name__}")
        return data

    def get_project_info(self, project_id: str) -> dict:
        response = self._request(
            "GET",
            f"{self.docs_url}/api/projects/{project_id}/info",
            headers=self.headers,
        )
        self._raise_for_response(response, "Не удалось прочитать информацию о проекте")
        data = response.json()
        if not isinstance(data, dict):
            raise RuntimeError(f"Ожидался объект проекта, получено: {type(data).__name__}")
        return data

    def get_folder_tree(self, project_id: str, root_id: str = "") -> dict:
        project_id = clean_text(project_id)
        if not project_id:
            raise RuntimeError("Не указан Project ID.")
        params = {"rootId": clean_text(root_id)} if clean_text(root_id) else None
        response = self._request(
            "GET",
            f"{self.docs_url}/api/folders/project/{project_id}/tree",
            params=params,
            headers=self.docs_api_headers(include_auth=True),
        )
        self._raise_for_response(response, "Не удалось загрузить дерево папок")
        data = response.json()
        if not isinstance(data, dict):
            raise RuntimeError(f"Ожидался объект дерева папок, получено: {type(data).__name__}")
        return data

    def create_folder(self, name: str, parent_id: str) -> str:
        name = clean_text(name)
        parent_id = clean_text(parent_id)
        if not name or not parent_id:
            raise RuntimeError("Для создания папки нужны name и parentId.")
        response = self._request(
            "PUT",
            f"{self.docs_url}/api/folders",
            headers=self.docs_api_headers(include_auth=True),
            json={"name": name, "parentId": parent_id},
        )
        self._raise_for_response(response, f"Не удалось создать папку '{name}'")
        data = response.json()
        folder_id = clean_text(data.get("data") if isinstance(data, dict) else "")
        if not folder_id:
            raise RuntimeError(f"API создал папку, но не вернул её ID: {data}")
        return folder_id

    def get_permissions_tree(self, project_id: str) -> dict:
        response = self._request(
            "GET",
            f"{self.docs_url}/api/permissions/tree",
            params={"projectId": project_id},
            headers=self.headers,
            timeout=self.timeout,
        )
        self._raise_for_response(response, "Не удалось загрузить дерево прав")
        data = response.json()
        if not isinstance(data, dict):
            raise RuntimeError(f"Ожидался объект дерева, получено: {type(data).__name__}")
        return data

    def upsert_folder_permissions(self, folder_id: str, payload: List[dict]) -> None:
        response = self._request(
            "POST",
            f"{self.docs_url}/api/permissions/folders/{folder_id}/upsert",
            headers=self.headers,
            json=payload,
        )
        self._raise_for_response(response, f"Не удалось применить права для папки {folder_id}")

    def delete_permissions(self, project_id: str, permission_ids: Sequence[str]) -> None:
        project_id = clean_text(project_id)
        ids = [clean_text(value) for value in permission_ids if clean_text(value)]
        if not project_id:
            raise RuntimeError("Для снятия прав не указан Project ID.")
        if not ids:
            return
        response = self._request(
            "DELETE",
            f"{self.docs_url}/api/permissions",
            headers=self.headers,
            json={"projectId": project_id, "permissionIds": ids},
        )
        self._raise_for_response(
            response,
            "Не удалось снять права: " + ", ".join(ids[:10]),
        )

    def attribute_headers(self, project_id: str) -> Dict[str, str]:
        """Заголовки для экрана настроек атрибутов, повторяющие запросы из HAR."""
        project_id = clean_text(project_id)
        headers = self.docs_api_headers(include_auth=True)
        if project_id:
            headers["Referer"] = f"{self.docs_url}/projects/{project_id}/settings/attributes"
        headers.update({
            "Sec-Fetch-Dest": "empty",
            "Sec-Fetch-Mode": "cors",
            "Sec-Fetch-Site": "same-origin",
        })
        return headers

    def get_attribute_types(self, project_id: str) -> List[dict]:
        project_id = clean_text(project_id)
        if not project_id:
            raise RuntimeError("Для загрузки атрибутов не указан Project ID.")
        response = self._request(
            "POST",
            f"{self.docs_url}/api/attributes/types/find",
            headers=self.attribute_headers(project_id),
            json={"projectId": project_id},
        )
        self._raise_for_response(response, "Не удалось загрузить атрибуты проекта")
        data = response.json()
        if not isinstance(data, list):
            raise RuntimeError(f"Ожидался список атрибутов, получено: {type(data).__name__}")
        return data

    def create_attribute_type(
        self,
        project_id: str,
        name: str,
        attribute_type: str,
        is_required: bool = False,
        data_type: str = "",
        values: Optional[Sequence[str]] = None,
    ) -> str:
        """Создаёт тип атрибута по точной форме PUT /api/attributes/types из HAR."""
        project_id = clean_text(project_id)
        name = clean_text(name)
        attribute_type = clean_text(attribute_type)
        data_type = clean_text(data_type)
        if not project_id or not name:
            raise RuntimeError("Для создания атрибута нужны Project ID и название.")
        if attribute_type not in {"Text", "Bool", "Date", "List"}:
            raise RuntimeError(f"Неподдерживаемый тип атрибута: {attribute_type}")
        if data_type and data_type not in {"Item", "Folder"}:
            raise RuntimeError(f"Неподдерживаемая область атрибута: {data_type}")

        list_payload = []
        seen_values = set()
        if attribute_type == "List":
            for value in values or []:
                value = clean_text(value)
                norm = normalize_text(value)
                if not value or norm in seen_values:
                    continue
                seen_values.add(norm)
                list_payload.append({"position": len(list_payload), "name": value})
            if not list_payload:
                raise RuntimeError(f"Для атрибута-списка '{name}' нужно указать хотя бы одно значение.")

        payload = {
            "projectId": project_id,
            "isRequired": bool(is_required),
            "type": attribute_type,
            "order": 0,
            "name": name,
            "list": list_payload,
        }
        # В HAR dataType может быть Item, Folder либо вообще отсутствовать.
        if data_type:
            payload["dataType"] = data_type

        response = self._request(
            "PUT",
            f"{self.docs_url}/api/attributes/types",
            headers=self.attribute_headers(project_id),
            json=payload,
        )
        self._raise_for_response(response, f"Не удалось создать атрибут '{name}'")
        data = response.json()
        attribute_id = clean_text(data.get("data") if isinstance(data, dict) else "")
        if not attribute_id:
            raise RuntimeError(f"API создал атрибут '{name}', но не вернул его ID: {data}")
        return attribute_id


# -----------------------------------------------------------------------------
# Indexes / resolvers
# -----------------------------------------------------------------------------
class FolderStructureIndex:
    def __init__(self, root_id: str):
        self.root_id = clean_text(root_id)
        self.by_path: Dict[str, dict] = {"": {"id": self.root_id, "name": "Root", "path_parts": []}}

    @classmethod
    def from_tree(cls, tree: dict) -> "FolderStructureIndex":
        root_id = clean_text(tree.get("id") or tree.get("folderId"))
        if not root_id:
            raise RuntimeError("В дереве папок не найден ID корневой папки.")
        index = cls(root_id)

        def walk(node: dict, parent_parts: List[str], is_root: bool = False):
            name = clean_text(node.get("name") or node.get("folderName"))
            folder_id = clean_text(node.get("id") or node.get("folderId"))
            current_parts = list(parent_parts)
            if not is_root and name:
                current_parts.append(name)
                norm_path = normalize_path_parts(current_parts)
                if norm_path not in index.by_path:
                    index.by_path[norm_path] = {
                        "id": folder_id,
                        "name": name,
                        "path_parts": current_parts,
                    }
            for child in node.get("children") or []:
                if isinstance(child, dict):
                    walk(child, current_parts, False)

        walk(tree, [], True)
        return index

    def get(self, path_parts: Sequence[str]) -> Optional[dict]:
        return self.by_path.get(normalize_path_parts(path_parts))

    def add(self, path_parts: Sequence[str], folder_id: str):
        parts = [clean_text(part) for part in path_parts if clean_text(part)]
        self.by_path[normalize_path_parts(parts)] = {
            "id": clean_text(folder_id),
            "name": parts[-1] if parts else "Root",
            "path_parts": parts,
        }


def build_required_folder_rows(entries: Sequence[FolderPathEntry]) -> List[FolderPathEntry]:
    """Разворачивает каждую Excel-строку во все префиксы и убирает дубли."""
    result: List[FolderPathEntry] = []
    seen: set = set()
    for entry in entries:
        for depth in range(1, len(entry.path_parts) + 1):
            parts = list(entry.path_parts[:depth])
            norm_path = normalize_path_parts(parts)
            if not norm_path or norm_path in seen:
                continue
            seen.add(norm_path)
            result.append(FolderPathEntry(row_number=entry.row_number, path_parts=parts))
    return result


class FolderIndex:
    def __init__(self):
        self.items: List[FolderItem] = []
        self.path_index: Dict[str, List[FolderItem]] = {}
        self.name_index: Dict[str, List[FolderItem]] = {}

    def add(self, item: FolderItem) -> None:
        self.items.append(item)
        self.path_index.setdefault(item.norm_path, []).append(item)
        self.name_index.setdefault(normalize_text(item.name), []).append(item)

    @classmethod
    def from_tree(cls, tree: dict) -> "FolderIndex":
        index = cls()

        def node_name(node: dict) -> str:
            return clean_text(node.get("folderName") or node.get("name"))

        def node_id(node: dict) -> str:
            return clean_text(node.get("folderId") or node.get("id"))

        def walk(node: dict, parent_path: List[str], is_root: bool = False) -> None:
            name = node_name(node)
            folder_id = node_id(node)
            children = node.get("children") or []
            permissions = node.get("permissions") or []
            current_path = list(parent_path)
            if not is_root and name:
                current_path.append(name)
                if folder_id:
                    index.add(FolderItem(folder_id, name, current_path, permissions if isinstance(permissions, list) else []))
            for child in children:
                if isinstance(child, dict):
                    walk(child, current_path, is_root=False)

        walk(tree, [], is_root=True)
        return index

    def find(self, path_parts: Sequence[str]) -> Tuple[Optional[FolderItem], str]:
        parts = list(path_parts)
        if parts and normalize_text(parts[0]) == "root":
            parts = parts[1:]
        norm = normalize_path_parts(parts)
        exact = self.path_index.get(norm, [])
        if len(exact) == 1:
            return exact[0], "Найдено по полному пути"
        if len(exact) > 1:
            return None, "Неоднозначный полный путь: " + "; ".join(item.path_text for item in exact[:8])

        last = parts[-1] if parts else ""
        same_name = self.name_index.get(normalize_text(last), [])
        if same_name:
            return None, "Папка с таким именем есть, но полный путь отличается: " + "; ".join(item.path_text for item in same_name[:5])

        close = best_close_matches(norm, self.path_index.keys())
        if close:
            variants: List[str] = []
            for key in close:
                variants.extend(item.path_text for item in self.path_index.get(key, []))
            return None, "Путь не найден. Похожие: " + "; ".join(variants[:5])
        return None, "Путь не найден в дереве проекта"


class PrincipalResolver:
    def __init__(self, roles: Sequence[dict], users: Sequence[dict]):
        self.by_norm: Dict[str, Principal] = {}
        self.display_names: Dict[str, str] = {}
        self._load_roles(roles)
        self._load_users(users)

    def _register(self, key: str, principal: Principal) -> None:
        norm = normalize_text(key)
        if norm and norm not in self.by_norm:
            self.by_norm[norm] = principal
            self.display_names[norm] = principal.name

    def _load_roles(self, roles: Sequence[dict]) -> None:
        for role in roles:
            role_id = clean_text(role.get("id") or role.get("roleId"))
            name = clean_text(role.get("name"))
            if role_id and name:
                principal = Principal(name=name, id=role_id, source_type="Role")
                self._register(name, principal)

    def _load_users(self, users: Sequence[dict]) -> None:
        for user in users:
            user_id = clean_text(user.get("userId") or user.get("id"))
            email = clean_text(user.get("userEmail") or user.get("email"))
            last = clean_text(user.get("userLastName") or user.get("lastName"))
            first = clean_text(user.get("userFirstName") or user.get("firstName"))
            middle = clean_text(user.get("userMiddleName") or user.get("middleName"))
            full_name = clean_text(" ".join(part for part in [last, first, middle] if part))
            name = full_name or email
            if user_id and name:
                principal = Principal(name=name, id=user_id, source_type="User")
                self._register(name, principal)
                if email:
                    self._register(email, principal)
                if first and last:
                    self._register(f"{first} {last}", principal)

    def resolve(self, requested_name: str) -> PrincipalResolution:
        norm = normalize_text(requested_name)
        principal = self.by_norm.get(norm)
        if principal:
            return PrincipalResolution(requested_name, principal, "точное совпадение")
        close = best_close_matches(norm, self.by_norm.keys(), limit=5, cutoff=0.72)
        if close:
            variants = ", ".join(self.display_names.get(key, key) for key in close)
            return PrincipalResolution(requested_name, message=f"Не найдено точно. Похожие: {variants}")
        return PrincipalResolution(requested_name, message="Роль/пользователь не найден")


# -----------------------------------------------------------------------------
# Workers
# -----------------------------------------------------------------------------
class OAuthTokenWorker(QtCore.QObject):
    log = Signal(str)
    done = Signal(object)
    failed = Signal(str)

    def __init__(self, docs_url: str, hub_url: str, code: str, verifier: str, redirect_uri: str, callback_url: str):
        super().__init__()
        self.docs_url = docs_url
        self.hub_url = hub_url
        self.code = code
        self.verifier = verifier
        self.redirect_uri = redirect_uri
        self.callback_url = callback_url

    @Slot()
    def run(self):
        try:
            self.log.emit("OAuth code получен, выполняю обмен на access_token\n")
            client = SgnlClient(self.docs_url, self.hub_url)
            data = client.exchange_authorization_code(
                code=self.code,
                verifier=self.verifier,
                redirect_uri=self.redirect_uri,
                referer=self.callback_url,
            )
            token = clean_text(data.get("access_token"))
            payload = decode_jwt_payload(token)
            if payload:
                client_id = clean_text(payload.get("client_id"))
                scopes = payload.get("scope") or []
                scopes_text = scopes if isinstance(scopes, str) else ", ".join(
                    clean_text(item) for item in scopes if clean_text(item)
                )
                self.log.emit(f"  token client_id: {client_id or 'не определён'}\n")
                self.log.emit(f"  token scopes: {scopes_text or 'не определены'}\n")
            self.done.emit({"token": token, "raw": data, "cookies": {}})
        except Exception:
            self.failed.emit(traceback.format_exc())


class LoginWorker(QtCore.QObject):
    log = Signal(str)
    done = Signal(object)
    captcha_required = Signal(str)
    failed = Signal(str)

    def __init__(self, docs_url: str, hub_url: str, email: str, password: str):
        super().__init__()
        self.docs_url = docs_url
        self.hub_url = hub_url
        self.email = email
        self.password = password

    def _log(self, text: str):
        self.log.emit(text)

    @Slot()
    def run(self):
        try:
            self._log("Авторизация через SGNL OAuth\n")
            client = SgnlClient(self.docs_url, self.hub_url)
            data = client.login_email_password(self.email, self.password)
            token = clean_text(data.get("access_token"))
            payload = decode_jwt_payload(token)
            self._log("  access_token получен\n")
            if payload:
                client_id = clean_text(payload.get("client_id"))
                scopes = payload.get("scope") or []
                if isinstance(scopes, str):
                    scopes_text = scopes
                elif isinstance(scopes, list):
                    scopes_text = ", ".join(clean_text(x) for x in scopes if clean_text(x))
                else:
                    scopes_text = ""
                self._log(f"  token client_id: {client_id or 'не определён'}\n")
                self._log(f"  token scopes: {scopes_text or 'не определены'}\n")
            self.done.emit({"token": token, "raw": data, "cookies": requests.utils.dict_from_cookiejar(client.session.cookies)})
        except CaptchaRequiredError as exc:
            self.log.emit("  SGNL запросил интерактивную CAPTCHA; переключаюсь на встроенную форму\n")
            self.captcha_required.emit(str(exc))
        except Exception:
            self.failed.emit(traceback.format_exc())


class LoadCatalogsWorker(QtCore.QObject):
    log = Signal(str)
    done = Signal(object)
    failed = Signal(str)

    def __init__(self, docs_url: str, hub_url: str, token: str, company_id: str = "", cookies: Optional[Dict[str, str]] = None):
        super().__init__()
        self.docs_url = docs_url
        self.hub_url = hub_url
        self.token = token
        self.company_id = company_id
        self.cookies = cookies or {}

    def _log(self, text: str):
        self.log.emit(text)

    @Slot()
    def run(self):
        try:
            client = SgnlClient(self.docs_url, self.hub_url, self.token, cookies=self.cookies)
            self._log("Загрузка компаний и проектов\n")
            companies = client.get_companies()
            selected_company_id = clean_text(self.company_id)
            if not selected_company_id and companies:
                selected_company_id = clean_text(companies[0].get("id"))
            projects = client.get_company_projects(selected_company_id) if selected_company_id else []
            self._log(f"  компаний: {len(companies)}; проектов: {len(projects)}\n")
            self.done.emit({"companies": companies, "projects": projects, "company_id": selected_company_id})
        except Exception:
            self.failed.emit(traceback.format_exc())


class CheckAccessWorker(QtCore.QObject):
    log = Signal(str)
    done = Signal(object)
    failed = Signal(str)

    def __init__(self, docs_url: str, hub_url: str, token: str, project_id: str, company_id: str, cookies: Optional[Dict[str, str]] = None):
        super().__init__()
        self.docs_url = docs_url
        self.hub_url = hub_url
        self.token = token
        self.project_id = project_id
        self.company_id = company_id
        self.cookies = cookies or {}

    def _log(self, text: str):
        self.log.emit(text)

    @Slot()
    def run(self):
        try:
            client = SgnlClient(self.docs_url, self.hub_url, self.token, cookies=self.cookies)
            self._log("Проверка API-доступа\n")
            project = client.get_project_info(self.project_id)
            self._log(f"  проект: {clean_text(project.get('name') or project.get('title') or self.project_id)}\n")
            try:
                roles = client.get_custom_roles(self.company_id)
                self._log(f"  ролей компании: {len(roles)}\n")
            except Exception as exc:
                self._log(f"  WARNING: роли компании не загружены: {exc}\n")
                roles = client.get_project_roles(self.project_id)
                self._log(f"  ролей проекта: {len(roles)}\n")
            tree = client.get_permissions_tree(self.project_id)
            index = FolderIndex.from_tree(tree)
            self._log(f"  папок в дереве прав: {len(index.items)}\n")
            self.done.emit({"roles": len(roles), "folders": len(index.items)})
        except Exception:
            self.failed.emit(traceback.format_exc())


class PreviewWorker(QtCore.QObject):
    log = Signal(str)
    done = Signal(object)
    failed = Signal(str)

    def __init__(
        self,
        docs_url: str,
        hub_url: str,
        token: str,
        project_id: str,
        company_id: str,
        excel_path: str,
        sheet_name: str = "",
        cookies: Optional[Dict[str, str]] = None,
    ):
        super().__init__()
        self.docs_url = docs_url
        self.hub_url = hub_url
        self.token = token
        self.project_id = project_id
        self.company_id = company_id
        self.excel_path = excel_path
        self.sheet_name = clean_text(sheet_name)
        self.cookies = cookies or {}

    def _log(self, text: str):
        self.log.emit(text)

    @Slot()
    def run(self):
        try:
            self._log("=" * 72 + "\n")
            self._log("Проверка Excel-матрицы и API SGNL\n")
            self._log(f"Excel: {self.excel_path}\n")
            self._log(f"Лист Excel: {self.sheet_name or 'автоматически'}\n")
            self._log(f"Project ID: {self.project_id}\n")
            self._log(f"Company ID: {self.company_id}\n")
            self._log("=" * 72 + "\n\n")

            self._log("[1/4] Чтение Excel\n")
            entries = PermissionExcelParser(self.excel_path, self.sheet_name).parse()
            self._log(f"  строк с командами прав: {len(entries)}\n")
            remove_markers = sum(
                1
                for entry in entries
                for permission in entry.permissions.values()
                if permission.is_remove
            )
            self._log(f"  явных маркеров '-' на снятие прав: {remove_markers}\n")
            roles_from_excel = sorted({role for entry in entries for role in entry.permissions.keys()}, key=normalize_text)
            self._log("  роли из Excel: " + ", ".join(roles_from_excel) + "\n")

            client = SgnlClient(self.docs_url, self.hub_url, self.token, cookies=self.cookies)

            self._log("[2/4] Загрузка ролей/пользователей\n")
            try:
                roles = client.get_custom_roles(self.company_id)
            except Exception as exc:
                roles = []
                self._log(f"  WARNING: роли компании не загружены: {exc}\n")
            try:
                project_roles = client.get_project_roles(self.project_id)
                # Merge project roles with company roles without duplicates.
                existing_ids = {clean_text(item.get("id") or item.get("roleId")) for item in roles}
                for role in project_roles:
                    rid = clean_text(role.get("id") or role.get("roleId"))
                    if rid and rid not in existing_ids:
                        roles.append(role)
                        existing_ids.add(rid)
            except Exception as exc:
                self._log(f"  WARNING: роли проекта отдельно не загружены: {exc}\n")
            try:
                users = client.get_company_users(self.company_id)
            except Exception as exc:
                users = []
                self._log(f"  WARNING: пользователи компании не загружены: {exc}\n")
            self._log(f"  ролей: {len(roles)}; пользователей: {len(users)}\n")
            resolver = PrincipalResolver(roles, users)

            self._log("[3/4] Загрузка дерева прав\n")
            tree = client.get_permissions_tree(self.project_id)
            index = FolderIndex.from_tree(tree)
            self._log(f"  папок: {len(index.items)}\n")
            self._log(f"  уникальных путей: {len(index.path_index)}\n")

            self._log("[4/4] Сопоставление путей, ролей и точное сравнение прав\n")
            principal_cache: Dict[str, PrincipalResolution] = {}
            plan: List[PlanRow] = []
            for entry in entries:
                folder, folder_message = index.find(entry.path_parts)
                status = "Папка не найдена" if not folder else "К применению"
                principal_results: Dict[str, PrincipalResolution] = {}
                messages = [folder_message]
                missing: List[str] = []

                for role_name in entry.permissions.keys():
                    if role_name not in principal_cache:
                        principal_cache[role_name] = resolver.resolve(role_name)
                    res = principal_cache[role_name]
                    principal_results[role_name] = res
                    if not res.ok:
                        missing.append(role_name)

                pending_roles: List[str] = []
                unchanged_roles: List[str] = []
                pending_remove_roles: List[str] = []
                pending_upsert_roles: List[str] = []
                if missing:
                    status = "Ошибка сопоставления"
                    messages.append("Не найдены роли/пользователи: " + ", ".join(missing[:10]))
                    for name in missing[:5]:
                        msg = principal_results[name].message
                        if msg:
                            messages.append(f"{name}: {msg}")
                elif folder:
                    existing_by_source = index_permissions_by_source(folder.permissions)
                    for role_name, desired_permission in entry.permissions.items():
                        resolution = principal_results[role_name]
                        assert resolution.principal is not None
                        existing = existing_by_source.get(resolution.principal.id)

                        if desired_permission.is_remove:
                            if existing is None:
                                unchanged_roles.append(role_name)
                            else:
                                pending_roles.append(role_name)
                                pending_remove_roles.append(role_name)
                                permission_id = clean_text(existing.get("id") or existing.get("permissionId"))
                                self._log(
                                    f"  DIFF: {entry.path_text} | {role_name}: "
                                    f"SGNL=запись {permission_id or 'без ID'}; Excel=-; действие=DELETE\n"
                                )
                        elif existing and permission_flags_equal(existing, desired_permission.flags):
                            unchanged_roles.append(role_name)
                        else:
                            pending_roles.append(role_name)
                            pending_upsert_roles.append(role_name)
                            self._log(
                                f"  DIFF: {entry.path_text} | {role_name}: "
                                f"действие=UPSERT {desired_permission.code}\n"
                            )

                    if pending_roles:
                        status = "К применению"
                        messages.append(
                            f"Изменений: {len(pending_roles)}; без изменений: {len(unchanged_roles)}"
                        )
                        if pending_remove_roles:
                            messages.append("Снять права: " + ", ".join(pending_remove_roles[:10]))
                        if pending_upsert_roles:
                            messages.append("Создать/перезаписать: " + ", ".join(pending_upsert_roles[:10]))
                        if unchanged_roles:
                            messages.append("Пропуск — уже соответствует Excel: " + ", ".join(unchanged_roles[:10]))
                    else:
                        status = "Без изменений"
                        messages.append(
                            f"Все {len(unchanged_roles)} команд уже выполнены: "
                            "права совпадают либо запись для '-' уже отсутствует"
                        )

                plan.append(
                    PlanRow(
                        entry=entry,
                        folder=folder,
                        status=status,
                        message=" ".join(m for m in messages if m),
                        principals=principal_results,
                        pending_roles=pending_roles,
                        unchanged_roles=unchanged_roles,
                    )
                )

            apply_count = sum(1 for row in plan if row.can_apply)
            unchanged_count = sum(1 for row in plan if row.no_changes)
            error_count = sum(1 for row in plan if not row.is_valid)
            pending_entries = sum(row.pending_permission_count for row in plan if row.can_apply)
            skipped_entries = sum(row.unchanged_permission_count for row in plan if row.is_valid)
            self._log(
                f"Предпросмотр готов. Папок к изменению: {apply_count}; "
                f"без изменений: {unchanged_count}; с ошибками: {error_count}; всего: {len(plan)}\n"
            )
            self._log(
                f"Операций к выполнению: {pending_entries}; уже соответствует Excel: {skipped_entries}\n"
            )
            self.done.emit({"plan": plan, "folder_index": index})
        except Exception:
            self.failed.emit(traceback.format_exc())


class ExcelStructureLoadWorker(QtCore.QObject):
    log = Signal(str)
    done = Signal(object)
    failed = Signal(str)

    def __init__(self, excel_path: str, sheet_name: str = ""):
        super().__init__()
        self.excel_path = excel_path
        self.sheet_name = clean_text(sheet_name)

    @Slot()
    def run(self):
        try:
            self.log.emit("Загрузка папочной структуры из Excel\n")
            self.log.emit(f"Excel: {self.excel_path}\n")
            self.log.emit(f"Лист: {self.sheet_name or 'автоматически'}\n")
            source_rows = PermissionExcelParser(self.excel_path, self.sheet_name).parse_paths()
            required = build_required_folder_rows(source_rows)
            plan: List[FolderCreatePlanRow] = []
            for entry in required:
                parent_parts = entry.path_parts[:-1]
                plan.append(
                    FolderCreatePlanRow(
                        row_number=entry.row_number,
                        path_parts=list(entry.path_parts),
                        status="Из Excel",
                        parent_path=display_path(parent_parts) or "Root",
                        message="Загружено из выбранного листа; сверка с SGNL ещё не выполнена",
                    )
                )
            self.log.emit(f"Уникальных папок из Excel: {len(plan)}\n")
            self.done.emit({"plan": plan, "count": len(plan), "sheet_name": self.sheet_name})
        except Exception:
            self.failed.emit(traceback.format_exc())


class FolderStructurePreviewWorker(QtCore.QObject):
    log = Signal(str)
    done = Signal(object)
    failed = Signal(str)

    def __init__(self, docs_url: str, hub_url: str, token: str, project_id: str, excel_path: str, sheet_name: str = "", cookies: Optional[Dict[str, str]] = None):
        super().__init__()
        self.docs_url = docs_url
        self.hub_url = hub_url
        self.token = token
        self.project_id = project_id
        self.excel_path = excel_path
        self.sheet_name = clean_text(sheet_name)
        self.cookies = cookies or {}

    @Slot()
    def run(self):
        try:
            self.log.emit("=" * 72 + "\n")
            self.log.emit("Предпросмотр создания папочной структуры\n")
            self.log.emit(f"Excel: {self.excel_path}\n")
            self.log.emit(f"Лист Excel: {self.sheet_name or 'автоматически'}\n")
            self.log.emit(f"Project ID: {self.project_id}\n")
            self.log.emit("=" * 72 + "\n\n")

            source_rows = PermissionExcelParser(self.excel_path, self.sheet_name).parse_paths()
            required = build_required_folder_rows(source_rows)
            self.log.emit(f"Уникальных требуемых папок: {len(required)}\n")

            client = SgnlClient(self.docs_url, self.hub_url, self.token, cookies=self.cookies)
            tree = client.get_folder_tree(self.project_id)
            index = FolderStructureIndex.from_tree(tree)
            self.log.emit(f"Папок уже существует: {max(0, len(index.by_path) - 1)}\n")

            available_paths = set(index.by_path.keys())
            plan: List[FolderCreatePlanRow] = []
            for entry in required:
                norm_path = normalize_path_parts(entry.path_parts)
                parent_parts = entry.path_parts[:-1]
                parent_norm = normalize_path_parts(parent_parts)
                existing = index.by_path.get(norm_path)
                if existing:
                    actual_name = clean_text(existing.get("name"))
                    message = "Уже есть в проекте"
                    if actual_name and actual_name != entry.path_parts[-1]:
                        message += f"; фактическое имя: {actual_name}"
                    plan.append(FolderCreatePlanRow(
                        row_number=entry.row_number,
                        path_parts=list(entry.path_parts),
                        status="Существует",
                        parent_path=display_path(parent_parts) or "Root",
                        folder_id=clean_text(existing.get("id")),
                        message=message,
                    ))
                    available_paths.add(norm_path)
                elif parent_norm in available_paths:
                    plan.append(FolderCreatePlanRow(
                        row_number=entry.row_number,
                        path_parts=list(entry.path_parts),
                        status="Создать",
                        parent_path=display_path(parent_parts) or "Root",
                        message="Папка отсутствует и будет создана",
                    ))
                    available_paths.add(norm_path)
                else:
                    plan.append(FolderCreatePlanRow(
                        row_number=entry.row_number,
                        path_parts=list(entry.path_parts),
                        status="Ошибка",
                        parent_path=display_path(parent_parts) or "Root",
                        message="Не найден родительский путь",
                    ))

            to_create = sum(1 for row in plan if row.status == "Создать")
            existing_count = sum(1 for row in plan if row.status == "Существует")
            errors = sum(1 for row in plan if row.status == "Ошибка")
            self.log.emit(f"Итог предпросмотра: создать {to_create}; уже есть {existing_count}; ошибок {errors}\n")
            self.done.emit({"plan": plan, "to_create": to_create, "existing": existing_count, "errors": errors})
        except Exception:
            self.failed.emit(traceback.format_exc())


class CreateFoldersWorker(QtCore.QObject):
    log = Signal(str)
    done = Signal(object)
    failed = Signal(str)

    def __init__(self, docs_url: str, hub_url: str, token: str, project_id: str, plan: Sequence[FolderCreatePlanRow], cookies: Optional[Dict[str, str]] = None):
        super().__init__()
        self.docs_url = docs_url
        self.hub_url = hub_url
        self.token = token
        self.project_id = project_id
        self.plan = list(plan)
        self.cookies = cookies or {}

    @Slot()
    def run(self):
        try:
            client = SgnlClient(self.docs_url, self.hub_url, self.token, cookies=self.cookies)
            tree = client.get_folder_tree(self.project_id)
            index = FolderStructureIndex.from_tree(tree)
            stats = {"created": 0, "existing": 0, "failed": 0, "total": len(self.plan)}
            self.log.emit("=" * 72 + "\n")
            self.log.emit("Создание папочной структуры\n")
            self.log.emit("=" * 72 + "\n\n")

            for number, row in enumerate(self.plan, 1):
                norm_path = normalize_path_parts(row.path_parts)
                current = index.by_path.get(norm_path)
                if current:
                    stats["existing"] += 1
                    self.log.emit(f"[{number}/{len(self.plan)}] SKIP: {row.path_text} — уже существует\n")
                    continue
                if row.status == "Ошибка":
                    stats["failed"] += 1
                    self.log.emit(f"[{number}/{len(self.plan)}] ERROR: {row.path_text} — {row.message}\n")
                    continue

                parent_parts = row.path_parts[:-1]
                parent = index.get(parent_parts)
                parent_id = clean_text(parent.get("id")) if parent else ""
                if not parent_id:
                    stats["failed"] += 1
                    self.log.emit(f"[{number}/{len(self.plan)}] ERROR: {row.path_text} — родитель ещё не создан\n")
                    continue

                try:
                    folder_id = client.create_folder(row.name, parent_id)
                    index.add(row.path_parts, folder_id)
                    stats["created"] += 1
                    self.log.emit(f"[{number}/{len(self.plan)}] CREATE: {row.path_text} ({folder_id})\n")
                except Exception as exc:
                    # За время между preview и apply папку мог создать другой пользователь.
                    try:
                        refreshed = FolderStructureIndex.from_tree(client.get_folder_tree(self.project_id))
                        existing = refreshed.get(row.path_parts)
                    except Exception:
                        refreshed = None
                        existing = None
                    if existing:
                        index = refreshed
                        stats["existing"] += 1
                        self.log.emit(f"[{number}/{len(self.plan)}] SKIP: {row.path_text} — появилась параллельно\n")
                    else:
                        stats["failed"] += 1
                        self.log.emit(f"[{number}/{len(self.plan)}] ERROR: {row.path_text} — {exc}\n")

            self.log.emit("\nИтог\n")
            self.log.emit(f"  создано: {stats['created']}\n")
            self.log.emit(f"  уже существовало: {stats['existing']}\n")
            self.log.emit(f"  ошибок: {stats['failed']}\n")
            self.done.emit(stats)
        except Exception:
            self.failed.emit(traceback.format_exc())


class ApplyWorker(QtCore.QObject):
    log = Signal(str)
    done = Signal(object)
    failed = Signal(str)

    def __init__(
        self,
        docs_url: str,
        hub_url: str,
        token: str,
        project_id: str,
        plan: Sequence[PlanRow],
        cookies: Optional[Dict[str, str]] = None,
    ):
        super().__init__()
        self.docs_url = docs_url
        self.hub_url = hub_url
        self.token = token
        self.project_id = project_id
        self.plan = list(plan)
        self.cookies = cookies or {}

    def _log(self, text: str):
        self.log.emit(text)

    @Slot()
    def run(self):
        try:
            client = SgnlClient(self.docs_url, self.hub_url, self.token, cookies=self.cookies)

            # Перед изменением читаем актуальные записи. Повторный запуск продолжит
            # работу: совпадающие upsert и уже удалённые permissionId будут пропущены.
            self._log("Перед применением повторно загружаю актуальное дерево прав SGNL\n")
            current_tree = client.get_permissions_tree(self.project_id)
            current_index = FolderIndex.from_tree(current_tree)

            rows = [row for row in self.plan if row.is_valid]
            invalid_rows = len(self.plan) - len(rows)
            stats = {
                "folders": 0,
                "updated_entries": 0,
                "removed_entries": 0,
                "already_actual_entries": 0,
                "already_actual_rows": 0,
                "failed": 0,
                "skipped": invalid_rows,
                "total_rows": len(self.plan),
            }
            self._log("=" * 72 + "\n")
            self._log("Синхронизация прав SGNL с Excel\n")
            self._log("  П/С/З/Р: POST .../upsert\n")
            self._log("  '-': DELETE /api/permissions по permissionId\n")
            self._log(
                f"Строк для повторной проверки: {len(rows)}; "
                f"ошибочных строк пропущено: {invalid_rows}\n"
            )
            self._log("=" * 72 + "\n\n")

            for idx, row in enumerate(rows, 1):
                current_folder, folder_message = current_index.find(row.entry.path_parts)
                if current_folder is None:
                    stats["failed"] += 1
                    self._log(
                        f"[{idx}/{len(rows)}] ERROR: {row.entry.path_text} — "
                        f"папка исчезла или путь изменился: {folder_message}\n"
                    )
                    continue

                self._log(f"[{idx}/{len(rows)}] Excel строка {row.entry.row_number}: {row.entry.path_text}\n")
                self._log(f"  SGNL: {current_folder.path_text}\n")
                self._log(f"  folderId: {current_folder.id}\n")

                upsert_payload, delete_items, skipped_messages = self._build_operations(row, current_folder)
                stats["already_actual_entries"] += len(skipped_messages)
                for message in skipped_messages:
                    self._log(f"  = SKIP {message}\n")

                missing_delete_ids = [item for item in delete_items if not clean_text(item.get("permissionId"))]
                delete_items = [item for item in delete_items if clean_text(item.get("permissionId"))]

                if not upsert_payload and not delete_items and not missing_delete_ids:
                    stats["already_actual_rows"] += 1
                    stats["skipped"] += 1
                    self._log("  SKIP: изменений для этой папки нет\n")
                    continue

                row_changed = False
                row_failed = False

                if missing_delete_ids:
                    row_failed = True
                    for item in missing_delete_ids:
                        self._log(
                            f"  ERROR: нельзя снять права у {item.get('principalName')}: "
                            "SGNL не вернул permissionId\n"
                        )

                if upsert_payload:
                    try:
                        client.upsert_folder_permissions(current_folder.id, upsert_payload)
                        stats["updated_entries"] += len(upsert_payload)
                        row_changed = True
                        self._merge_payload_into_folder(current_folder, upsert_payload)
                        for item in upsert_payload:
                            principal_name = self._principal_name_by_payload(row, item)
                            flags_text = self._flags_to_text(item)
                            action = "UPDATE" if clean_text(item.get("id")) else "CREATE"
                            self._log(f"  + {action} {principal_name}: {flags_text}\n")
                    except Exception as exc:
                        recovered = False
                        remaining_names: List[str] = []
                        try:
                            refreshed_tree = client.get_permissions_tree(self.project_id)
                            refreshed_index = FolderIndex.from_tree(refreshed_tree)
                            refreshed_folder, _ = refreshed_index.find(row.entry.path_parts)
                            if refreshed_folder is not None:
                                remaining_upserts, _remaining_deletes, _skipped = self._build_operations(
                                    row, refreshed_folder
                                )
                                remaining_names = [
                                    self._principal_name_by_payload(row, item)
                                    for item in remaining_upserts
                                ]
                                if not remaining_upserts:
                                    current_index = refreshed_index
                                    current_folder = refreshed_folder
                                    recovered = True
                        except Exception as verify_exc:
                            self._log(f"  WARNING: не удалось проверить upsert после ошибки: {verify_exc}\n")

                        if recovered:
                            stats["updated_entries"] += len(upsert_payload)
                            row_changed = True
                            self._log(
                                "  OK: соединение оборвалось, но повторная проверка "
                                "подтвердила применение upsert\n"
                            )
                        else:
                            row_failed = True
                            self._log(f"  ERROR UPSERT: {exc}\n")
                            if remaining_names:
                                self._log("  Не подтверждены: " + ", ".join(remaining_names) + "\n")

                if delete_items:
                    permission_ids = [clean_text(item.get("permissionId")) for item in delete_items]
                    try:
                        client.delete_permissions(self.project_id, permission_ids)
                        stats["removed_entries"] += len(permission_ids)
                        row_changed = True
                        self._remove_deleted_permissions(current_folder, delete_items)
                        for item in delete_items:
                            self._log(
                                f"  - REMOVE {item.get('principalName')}: "
                                f"permissionId={item.get('permissionId')}\n"
                            )
                    except Exception as exc:
                        recovered = False
                        remaining_names: List[str] = []
                        try:
                            refreshed_tree = client.get_permissions_tree(self.project_id)
                            refreshed_index = FolderIndex.from_tree(refreshed_tree)
                            refreshed_folder, _ = refreshed_index.find(row.entry.path_parts)
                            if refreshed_folder is not None:
                                _remaining_upserts, remaining_deletes, _skipped = self._build_operations(
                                    row, refreshed_folder
                                )
                                remaining_delete_ids = {
                                    clean_text(item.get("permissionId"))
                                    for item in remaining_deletes
                                }
                                requested_ids = set(permission_ids)
                                remaining_names = [
                                    clean_text(item.get("principalName"))
                                    for item in remaining_deletes
                                    if clean_text(item.get("permissionId")) in requested_ids
                                ]
                                if not (requested_ids & remaining_delete_ids):
                                    current_index = refreshed_index
                                    current_folder = refreshed_folder
                                    recovered = True
                        except Exception as verify_exc:
                            self._log(f"  WARNING: не удалось проверить DELETE после ошибки: {verify_exc}\n")

                        if recovered:
                            stats["removed_entries"] += len(permission_ids)
                            row_changed = True
                            self._log(
                                "  OK: соединение оборвалось, но повторная проверка "
                                "подтвердила снятие прав\n"
                            )
                        else:
                            row_failed = True
                            self._log(f"  ERROR DELETE: {exc}\n")
                            if remaining_names:
                                self._log("  Не сняты права: " + ", ".join(remaining_names) + "\n")

                if row_changed:
                    stats["folders"] += 1
                if row_failed:
                    stats["failed"] += 1

            self._log("\nИтог\n")
            self._log(f"  изменено папок: {stats['folders']}\n")
            self._log(f"  создано/перезаписано записей: {stats['updated_entries']}\n")
            self._log(f"  удалено записей прав: {stats['removed_entries']}\n")
            self._log(f"  операций уже было выполнено: {stats['already_actual_entries']}\n")
            self._log(f"  папок без изменений: {stats['already_actual_rows']}\n")
            self._log(f"  ошибок: {stats['failed']}\n")
            self._log(f"  пропущено строк: {stats['skipped']}\n")
            self.done.emit(stats)
        except Exception:
            self.failed.emit(traceback.format_exc())

    def _build_operations(
        self,
        row: PlanRow,
        folder: FolderItem,
    ) -> Tuple[List[dict], List[dict], List[str]]:
        existing_by_source = index_permissions_by_source(folder.permissions)
        upsert_payload: List[dict] = []
        delete_items: List[dict] = []
        skipped_messages: List[str] = []

        for role_name, perm in row.entry.permissions.items():
            res = row.principals.get(role_name)
            if not res or not res.principal:
                continue
            source_id = res.principal.id
            existing = existing_by_source.get(source_id)

            if perm.is_remove:
                if existing is None:
                    skipped_messages.append(f"{res.principal.name}: запись права уже отсутствует")
                    continue
                delete_items.append(
                    {
                        "permissionId": clean_text(existing.get("id") or existing.get("permissionId")),
                        "sourceId": source_id,
                        "principalName": res.principal.name,
                    }
                )
                continue

            if existing and permission_flags_equal(existing, perm.flags):
                skipped_messages.append(
                    f"{res.principal.name}: все пять флагов совпадают с Excel"
                )
                continue

            desired_flags = {
                key: bool(perm.flags.get(key))
                for key in PERMISSION_FLAG_KEYS
            }
            item = {
                **desired_flags,
                "sourceId": source_id,
                "sourceType": res.principal.source_type,
            }
            if existing:
                existing_id = clean_text(existing.get("id") or existing.get("permissionId"))
                if existing_id:
                    item["id"] = existing_id
            upsert_payload.append(item)

        return upsert_payload, delete_items, skipped_messages

    @staticmethod
    def _merge_payload_into_folder(folder: FolderItem, payload: Sequence[dict]) -> None:
        """Обновляет локальный снимок после успешного upsert."""
        existing_by_source = index_permissions_by_source(folder.permissions)
        for sent in payload:
            source_id = clean_text(sent.get("sourceId"))
            if not source_id:
                continue
            target = existing_by_source.get(source_id)
            if target is None:
                target = {}
                folder.permissions.append(target)
                existing_by_source[source_id] = target
            target.update(sent)

    @staticmethod
    def _remove_deleted_permissions(folder: FolderItem, delete_items: Sequence[dict]) -> None:
        deleted_ids = {
            clean_text(item.get("permissionId"))
            for item in delete_items
            if clean_text(item.get("permissionId"))
        }
        deleted_sources = {
            clean_text(item.get("sourceId"))
            for item in delete_items
            if clean_text(item.get("sourceId"))
        }
        folder.permissions = [
            item
            for item in folder.permissions
            if clean_text(item.get("id") or item.get("permissionId")) not in deleted_ids
            and permission_source_id(item) not in deleted_sources
        ]

    def _principal_name_by_payload(self, row: PlanRow, payload_item: dict) -> str:
        source_id = clean_text(payload_item.get("sourceId"))
        for role_name, res in row.principals.items():
            if res.principal and res.principal.id == source_id:
                return res.principal.name
        return source_id

    @staticmethod
    def _flags_to_text(item: dict) -> str:
        parts = []
        for key, label in [
            ("read", "read"),
            ("download", "download"),
            ("create", "create"),
            ("update", "update"),
            ("delete", "delete"),
        ]:
            if item.get(key):
                parts.append(label)
        return ", ".join(parts) if parts else "нет прав"


class LoadAttributeTypesWorker(QtCore.QObject):
    log = Signal(str)
    done = Signal(object)
    failed = Signal(str)

    def __init__(self, docs_url: str, hub_url: str, token: str, project_id: str,
                 cookies: Optional[Dict[str, str]] = None):
        super().__init__()
        self.docs_url = docs_url
        self.hub_url = hub_url
        self.token = token
        self.project_id = project_id
        self.cookies = cookies or {}

    @Slot()
    def run(self):
        try:
            client = SgnlClient(self.docs_url, self.hub_url, self.token, cookies=self.cookies)
            self.log.emit("Загрузка атрибутов проекта\n")
            attributes = client.get_attribute_types(self.project_id)
            active = [item for item in attributes if not bool(item.get("deleted"))]
            self.log.emit(f"  всего записей: {len(attributes)}; активных: {len(active)}\n")
            self.done.emit({"attributes": attributes})
        except Exception:
            self.failed.emit(traceback.format_exc())


class CreateAttributeTypesWorker(QtCore.QObject):
    log = Signal(str)
    done = Signal(object)
    failed = Signal(str)

    def __init__(self, docs_url: str, hub_url: str, token: str, project_id: str,
                 drafts: Sequence[dict], cookies: Optional[Dict[str, str]] = None):
        super().__init__()
        self.docs_url = docs_url
        self.hub_url = hub_url
        self.token = token
        self.project_id = project_id
        self.drafts = [dict(item) for item in drafts]
        self.cookies = cookies or {}

    @Slot()
    def run(self):
        try:
            client = SgnlClient(self.docs_url, self.hub_url, self.token, cookies=self.cookies)
            existing = client.get_attribute_types(self.project_id)
            existing_names = {
                normalize_text(item.get("name"))
                for item in existing
                if not bool(item.get("deleted")) and clean_text(item.get("name"))
            }
            stats = {"created": 0, "skipped": 0, "failed": 0, "created_names": [], "errors": []}
            self.log.emit("=" * 72 + "\n")
            self.log.emit("Создание атрибутов SGNL\n")
            self.log.emit(f"В очереди: {len(self.drafts)}\n")
            self.log.emit("=" * 72 + "\n\n")

            for index, draft in enumerate(self.drafts, 1):
                name = clean_text(draft.get("name"))
                norm_name = normalize_text(name)
                if norm_name in existing_names:
                    stats["skipped"] += 1
                    self.log.emit(f"[{index}/{len(self.drafts)}] SKIP: '{name}' уже существует\n")
                    continue
                try:
                    attribute_id = client.create_attribute_type(
                        project_id=self.project_id,
                        name=name,
                        attribute_type=clean_text(draft.get("type")),
                        is_required=bool(draft.get("isRequired")),
                        data_type=clean_text(draft.get("dataType")),
                        values=list(draft.get("values") or []),
                    )
                    existing_names.add(norm_name)
                    stats["created"] += 1
                    stats["created_names"].append(name)
                    self.log.emit(f"[{index}/{len(self.drafts)}] OK: '{name}' — {attribute_id}\n")
                except Exception as exc:
                    stats["failed"] += 1
                    stats["errors"].append(f"{name}: {exc}")
                    self.log.emit(f"[{index}/{len(self.drafts)}] ERROR: '{name}' — {exc}\n")

            refreshed = client.get_attribute_types(self.project_id)
            self.log.emit("\nИтог\n")
            self.log.emit(f"  создано: {stats['created']}\n")
            self.log.emit(f"  уже существовало: {stats['skipped']}\n")
            self.log.emit(f"  ошибок: {stats['failed']}\n")
            self.done.emit({"stats": stats, "attributes": refreshed})
        except Exception:
            self.failed.emit(traceback.format_exc())


# -----------------------------------------------------------------------------
# UI
# -----------------------------------------------------------------------------
class ModernSection(QtWidgets.QWidget):
    def __init__(self, icon: str, title: str, collapsible: bool = True, parent=None):
        super().__init__(parent)
        self._collapsible = collapsible
        self._expanded = True
        layout = QtWidgets.QVBoxLayout(self)
        layout.setContentsMargins(0, 0, 0, 0)
        layout.setSpacing(0)
        self.card = QtWidgets.QFrame(self)
        self.card.setObjectName("card")
        card_layout = QtWidgets.QVBoxLayout(self.card)
        card_layout.setContentsMargins(0, 0, 0, 0)
        card_layout.setSpacing(0)
        self.header = QtWidgets.QFrame(self.card)
        self.header.setObjectName("sectionHeader")
        header_layout = QtWidgets.QHBoxLayout(self.header)
        header_layout.setContentsMargins(12, 8, 12, 6)
        header_layout.setSpacing(8)
        self.icon_label = QtWidgets.QLabel(icon)
        self.icon_label.setObjectName("sectionIcon")
        self.title_label = QtWidgets.QLabel(title)
        self.title_label.setObjectName("sectionTitle")
        self.arrow_label = QtWidgets.QLabel("⌄" if collapsible else "")
        self.arrow_label.setObjectName("sectionArrow")
        header_layout.addWidget(self.icon_label)
        header_layout.addWidget(self.title_label)
        header_layout.addStretch(1)
        header_layout.addWidget(self.arrow_label)
        self.body = QtWidgets.QFrame(self.card)
        self.body.setObjectName("sectionBody")
        self.grid = QtWidgets.QGridLayout(self.body)
        self.grid.setContentsMargins(12, 0, 12, 10)
        self.grid.setHorizontalSpacing(10)
        self.grid.setVerticalSpacing(7)
        card_layout.addWidget(self.header)
        card_layout.addWidget(self.body)
        layout.addWidget(self.card)
        if collapsible:
            self.header.setCursor(QtCore.Qt.PointingHandCursor)
            self.header.mousePressEvent = self._toggle_from_mouse

    def _toggle_from_mouse(self, event):
        self.set_expanded(not self._expanded)
        event.accept()

    def set_expanded(self, expanded: bool):
        if not self._collapsible:
            return
        self._expanded = expanded
        self.body.setVisible(expanded)
        self.arrow_label.setText("⌄" if expanded else "›")


class LogDialog(QtWidgets.QDialog):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.setWindowTitle("Log")
        self.resize(920, 560)
        layout = QtWidgets.QVBoxLayout(self)
        layout.setContentsMargins(8, 8, 8, 8)
        self.text = QtWidgets.QTextEdit()
        self.text.setObjectName("logArea")
        self.text.setReadOnly(True)
        self.text.setFont(QtGui.QFont("Cascadia Code", 10))
        layout.addWidget(self.text, 1)

    def set_log_text(self, value: str):
        self.text.setPlainText(value)
        self.text.moveCursor(QtGui.QTextCursor.End)
        self.text.ensureCursorVisible()

    def append(self, value: str):
        self.text.moveCursor(QtGui.QTextCursor.End)
        self.text.insertPlainText(value)
        self.text.ensureCursorVisible()


class MainWindow(QtWidgets.QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("Larix CDE — SIGNAL Docs")
        self.resize(1200, 700)
        self.setMinimumSize(1020, 620)

        self.current_plan: List[PlanRow] = []
        self.current_folder_plan: List[FolderCreatePlanRow] = []
        self.folder_plan_verified = False
        self.attribute_drafts: List[dict] = []
        self.current_attribute_types: List[dict] = []
        self._attribute_auto_loaded_project = ""
        self.active_threads: List[QtCore.QThread] = []
        self.active_workers: List[QtCore.QObject] = []
        self._log_lines: List[str] = []
        self.log_dialog: Optional[LogDialog] = None
        self._apply_after_preview = False
        self._loaded_token = ""
        self._session_cookies: Dict[str, str] = {}
        self._oauth_request: Dict[str, str] = {}
        self._oauth_email = ""
        self._oauth_password = ""

        base_dir = Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parent))
        self.asset_dir = shared_asset_dir(base_dir)
        self.is_dark_theme = initial_dark_theme(False)
        icon_file = asset_path(self.asset_dir, "icon.ico")
        if icon_file:
            self.setWindowIcon(QtGui.QIcon(icon_file))

        self._setup_ui()
        install_status_bar(self)
        mark_destructive_buttons(self)
        self._apply_stylesheet()
        install_window_state(self, "cde_tool")

    def _make_inline_icon(self, file_name: str, size: int = 18) -> QtWidgets.QLabel:
        label = QtWidgets.QLabel()
        label.setObjectName("inlineIcon")
        label.setFixedSize(size, size)
        label.setProperty("iconAsset", file_name)
        label.setProperty("iconSize", size)
        icon = themed_icon(self.asset_dir, file_name, self.is_dark_theme)
        label.setPixmap(icon.pixmap(QtCore.QSize(size, size)))
        label.setAlignment(QtCore.Qt.AlignCenter)
        return label

    def _add_card_header(self, layout: QtWidgets.QVBoxLayout, title: str, icon_name: str,
                         status: Optional[QtWidgets.QLabel] = None):
        row = QtWidgets.QHBoxLayout()
        row.setSpacing(8)
        row.addWidget(self._make_inline_icon(icon_name, 18), 0, QtCore.Qt.AlignVCenter)
        title_label = QtWidgets.QLabel(title)
        title_label.setObjectName("cardTitle")
        row.addWidget(title_label)
        row.addStretch(1)
        if status is not None:
            row.addWidget(status, 0, QtCore.Qt.AlignVCenter)
        layout.addLayout(row)

    def _set_button_icon(self, button: QtWidgets.QAbstractButton, file_name: str, size: int = 16):
        if button is None:
            return
        button.setProperty("iconAsset", file_name)
        button.setIcon(themed_icon(self.asset_dir, file_name, self.is_dark_theme))
        button.setIconSize(QtCore.QSize(size, size))

    def _refresh_action_icons(self):
        mappings = [
            ("btn_login", "free-icon-login-2623062.png"),
            ("btn_load_lists", "free-icon-refresh-5234214.png"),
            ("btn_check", "access_icon_variant_1.png"),
            ("btn_excel", "upload.png"),
            ("btn_load_sheets", "free-icon-refresh-5234214.png"),
            ("btn_log", "information.png"),
            ("btn_permissions_template", "Excel.png"),
            ("btn_preview", "comparison.png"),
            ("btn_apply", "krug_galka.png"),
            ("btn_folder_log", "information.png"),
            ("btn_folder_template", "Excel.png"),
            ("btn_folder_load_excel", "Excel.png"),
            ("btn_folder_preview", "comparison.png"),
            ("btn_folder_create", "folder_icon_variant_1.png"),
            ("btn_attr_log", "information.png"),
            ("btn_attr_template", "Excel.png"),
            ("btn_attr_refresh", "free-icon-refresh-5234214.png"),
            ("btn_attr_add", "arrow-right.png"),
            ("btn_attr_create", "krug_galka.png"),
            ("btn_auth_cancel", "arrow-left.png"),
            ("btn_service_toggle", "links.png"),
        ]
        for attr, icon_name in mappings:
            button = getattr(self, attr, None)
            if button is not None:
                self._set_button_icon(button, icon_name)
        self._update_sensitive_icons()
        for label in self.findChildren(QtWidgets.QLabel):
            asset = label.property("iconAsset")
            if not asset:
                continue
            try:
                size = int(label.property("iconSize") or 18)
            except Exception:
                size = 18
            label.setPixmap(themed_icon(self.asset_dir, str(asset), self.is_dark_theme).pixmap(QtCore.QSize(size, size)))

    def _update_sensitive_icons(self):
        if hasattr(self, "btn_password_eye"):
            visible = self.ed_password.echoMode() == QtWidgets.QLineEdit.Normal
            icon_name = "free-icon-hide-11238328.png" if visible else "free-icon-eye-2455724.png"
            self._set_button_icon(self.btn_password_eye, icon_name, 15)
            self.btn_password_eye.setToolTip("Скрыть пароль" if visible else "Показать пароль")
        if hasattr(self, "btn_toggle_token"):
            visible = self.ed_token.echoMode() == QtWidgets.QLineEdit.Normal
            icon_name = "free-icon-hide-11238328.png" if visible else "free-icon-eye-2455724.png"
            self._set_button_icon(self.btn_toggle_token, icon_name, 15)
            self.btn_toggle_token.setToolTip("Скрыть access_token" if visible else "Показать access_token")

    def _fit_workspace_tabs(self):
        """Подгоняет высоту блока вкладок под содержимое текущего раздела."""
        if not hasattr(self, "tabs"):
            return
        page = self.tabs.currentWidget()
        if page is None:
            return
        layout = page.layout()
        if layout is not None:
            layout.activate()
        page.adjustSize()
        content_h = page.sizeHint().height()
        tab_h = self.tabs.tabBar().sizeHint().height()
        # Небольшой запас на рамку QTabWidget и внутренние отступы.
        target_h = max(220, content_h + tab_h + 10)
        self.tabs.setFixedHeight(target_h)

    def _schedule_workspace_fit(self, *_args):
        QtCore.QTimer.singleShot(0, self._fit_workspace_tabs)

    def _toggle_theme(self, dark: Optional[bool] = None):
        self.is_dark_theme = (not self.is_dark_theme) if dark is None else bool(dark)
        persist_dark_theme(self.is_dark_theme)
        self._apply_stylesheet()
        if self.current_plan:
            self._fill_plan_table(self.current_plan)
        if self.current_folder_plan:
            self._fill_folder_table(self.current_folder_plan)

    def _toggle_password_visibility(self):
        if self.ed_password.echoMode() == QtWidgets.QLineEdit.Password:
            self.ed_password.setEchoMode(QtWidgets.QLineEdit.Normal)
        else:
            self.ed_password.setEchoMode(QtWidgets.QLineEdit.Password)
        self._update_sensitive_icons()

    def _toggle_service_panel(self, checked: bool):
        self.service_panel.setVisible(bool(checked))
        self.btn_service_toggle.setText("Скрыть служебные параметры" if checked else "Служебные параметры")

    def _setup_ui(self):
        self.setFont(QtGui.QFont("Segoe UI", 9))
        central = QtWidgets.QWidget()
        central.setObjectName("centralRoot")
        self.setCentralWidget(central)
        central_layout = QtWidgets.QVBoxLayout(central)
        central_layout.setContentsMargins(0, 0, 0, 0)
        central_layout.setSpacing(0)

        scroll = QtWidgets.QScrollArea()
        scroll.setObjectName("mainScroll")
        scroll.setWidgetResizable(True)
        scroll.setFrameShape(QtWidgets.QFrame.NoFrame)
        scroll.setHorizontalScrollBarPolicy(QtCore.Qt.ScrollBarAsNeeded)
        scroll.setVerticalScrollBarPolicy(QtCore.Qt.ScrollBarAsNeeded)
        central_layout.addWidget(scroll)

        page = QtWidgets.QWidget()
        page.setObjectName("page")
        scroll.setWidget(page)
        root = QtWidgets.QVBoxLayout(page)
        root.setContentsMargins(24, 20, 24, 12)
        root.setSpacing(12)

        # Header: повторяет структуру исходного Larix-интерфейса.
        header = QtWidgets.QHBoxLayout()
        header.setContentsMargins(0, 0, 0, 4)
        header.setSpacing(12)
        title_col = QtWidgets.QVBoxLayout()
        title_col.setSpacing(2)
        title = QtWidgets.QLabel("SIGNAL Docs")
        title.setObjectName("pageTitle")
        subtitle = QtWidgets.QLabel("Структура проекта, ролевая матрица и пользовательские атрибуты")
        subtitle.setObjectName("pageSubtitle")
        title_col.addWidget(title)
        title_col.addWidget(subtitle)
        header.addLayout(title_col, 1)
        self.theme_toggle = ThemeToggle(self.asset_dir, self)
        self.theme_toggle.setToolTip("Светлая / тёмная тема")
        self.theme_toggle.toggled.connect(self._toggle_theme)
        self.back_to_manager_button = add_standard_header_controls(header, self, self.theme_toggle)
        root.addLayout(header)

        # Общая авторизация — одна на все разделы.
        auth_card = QtWidgets.QFrame()
        auth_card.setObjectName("card")
        auth_layout = QtWidgets.QVBoxLayout(auth_card)
        auth_layout.setContentsMargins(14, 11, 14, 13)
        auth_layout.setSpacing(8)
        self.lbl_auth = QtWidgets.QLabel("Не авторизован")
        self.lbl_auth.setObjectName("badStatus")
        self._add_card_header(auth_layout, "Подключение к SGNL", "free-icon-login-2623062.png", self.lbl_auth)

        auth_grid = QtWidgets.QGridLayout()
        auth_grid.setHorizontalSpacing(12)
        auth_grid.setVerticalSpacing(6)
        auth_grid.setColumnStretch(0, 1)
        auth_grid.setColumnStretch(1, 1)
        auth_grid.setColumnMinimumWidth(2, 120)

        self.ed_email = QtWidgets.QLineEdit()
        self.ed_email.setPlaceholderText("Email SGNL")
        self.ed_password = QtWidgets.QLineEdit()
        self.ed_password.setPlaceholderText("Пароль")
        self.ed_password.setEchoMode(QtWidgets.QLineEdit.Password)
        password_box = QtWidgets.QWidget()
        password_layout = QtWidgets.QHBoxLayout(password_box)
        password_layout.setContentsMargins(0, 0, 0, 0)
        password_layout.setSpacing(4)
        password_layout.addWidget(self.ed_password, 1)
        self.btn_password_eye = QtWidgets.QToolButton()
        self.btn_password_eye.setObjectName("passwordEye")
        self.btn_password_eye.setFixedSize(34, 34)
        self.btn_password_eye.clicked.connect(self._toggle_password_visibility)
        password_layout.addWidget(self.btn_password_eye)

        self.btn_login = QtWidgets.QPushButton("Войти")
        self.btn_login.setObjectName("loginButton")
        self.btn_login.setCursor(QtCore.Qt.PointingHandCursor)
        self.btn_login.setFixedHeight(36)
        self.btn_login.setMinimumWidth(120)

        email_label = self._field_label("Email")
        password_label = self._field_label("Пароль")
        auth_grid.addWidget(email_label, 0, 0)
        auth_grid.addWidget(password_label, 0, 1)
        auth_grid.addWidget(self.ed_email, 1, 0)
        auth_grid.addWidget(password_box, 1, 1)
        auth_grid.addWidget(self.btn_login, 1, 2)

        self.cmb_company = QtWidgets.QComboBox()
        self.cmb_company.setEditable(True)
        self.cmb_company.setPlaceholderText("Сначала войдите")
        if self.cmb_company.lineEdit():
            self.cmb_company.lineEdit().setPlaceholderText("Сначала войдите")
        self.cmb_project = QtWidgets.QComboBox()
        self.cmb_project.setEditable(True)
        self.cmb_project.setPlaceholderText("Сначала выберите компанию")
        if self.cmb_project.lineEdit():
            self.cmb_project.lineEdit().setPlaceholderText("Сначала выберите компанию")
        self.btn_load_lists = QtWidgets.QPushButton("Обновить списки")
        self.btn_load_lists.setObjectName("secondaryAction")
        self.btn_load_lists.setMinimumWidth(120)
        auth_grid.addWidget(self._field_label("Компания"), 2, 0)
        auth_grid.addWidget(self._field_label("Проект"), 2, 1)
        auth_grid.addWidget(self.cmb_company, 3, 0)
        auth_grid.addWidget(self.cmb_project, 3, 1)
        auth_grid.addWidget(self.btn_load_lists, 3, 2)
        auth_layout.addLayout(auth_grid)

        # Системные ID API нужны логике, но не должны загромождать основной экран.
        self.ed_project_id = QtWidgets.QLineEdit("")
        self.ed_project_id.setVisible(False)
        self.ed_company_id = QtWidgets.QLineEdit("")
        self.ed_company_id.setVisible(False)

        self.btn_service_toggle = QtWidgets.QPushButton("Служебные параметры")
        self.btn_service_toggle.setObjectName("serviceToggle")
        self.btn_service_toggle.setCheckable(True)
        self.btn_service_toggle.setChecked(False)
        self.btn_service_toggle.toggled.connect(self._toggle_service_panel)
        auth_layout.addWidget(self.btn_service_toggle, 0, QtCore.Qt.AlignLeft)

        self.service_panel = QtWidgets.QFrame()
        self.service_panel.setObjectName("servicePanel")
        service_grid = QtWidgets.QGridLayout(self.service_panel)
        service_grid.setContentsMargins(10, 9, 10, 9)
        service_grid.setHorizontalSpacing(10)
        service_grid.setVerticalSpacing(6)
        self.ed_docs_url = QtWidgets.QLineEdit(DEFAULT_DOCS_URL)
        self.ed_hub_url = QtWidgets.QLineEdit(DEFAULT_HUB_URL)
        self.ed_token = QtWidgets.QLineEdit()
        self.ed_token.setPlaceholderText("access_token заполняется автоматически после входа")
        self.ed_token.setEchoMode(QtWidgets.QLineEdit.Password)
        self.btn_toggle_token = QtWidgets.QToolButton()
        self.btn_toggle_token.setObjectName("eyeButton")
        self.btn_toggle_token.setFixedSize(34, 34)
        token_box = QtWidgets.QWidget()
        token_layout = QtWidgets.QHBoxLayout(token_box)
        token_layout.setContentsMargins(0, 0, 0, 0)
        token_layout.setSpacing(4)
        token_layout.addWidget(self.ed_token, 1)
        token_layout.addWidget(self.btn_toggle_token)
        self.btn_check = QtWidgets.QPushButton("Проверить доступ")
        self.btn_check.setObjectName("secondaryAction")
        service_grid.addWidget(self._field_label("Docs URL"), 0, 0)
        service_grid.addWidget(self.ed_docs_url, 0, 1)
        service_grid.addWidget(self._field_label("Hub URL"), 0, 2)
        service_grid.addWidget(self.ed_hub_url, 0, 3)
        service_grid.addWidget(self._field_label("Token"), 1, 0)
        service_grid.addWidget(token_box, 1, 1, 1, 2)
        service_grid.addWidget(self.btn_check, 1, 3)
        service_grid.setColumnStretch(1, 1)
        service_grid.setColumnStretch(3, 1)
        self.service_panel.hide()
        auth_layout.addWidget(self.service_panel)
        root.addWidget(auth_card)

        # Excel общий для ролевой матрицы и папочной структуры. На вкладке атрибутов он скрывается.
        excel_card = QtWidgets.QFrame()
        self.excel_card = excel_card
        excel_card.setObjectName("card")
        excel_layout = QtWidgets.QVBoxLayout(excel_card)
        excel_layout.setContentsMargins(14, 11, 14, 13)
        excel_layout.setSpacing(9)
        self._add_card_header(excel_layout, "Excel и параметры", "Excel.png")
        excel_grid = QtWidgets.QGridLayout()
        excel_grid.setHorizontalSpacing(10)
        excel_grid.setVerticalSpacing(6)
        self.ed_excel = QtWidgets.QLineEdit()
        self.ed_excel.setPlaceholderText("Выберите Excel-файл с папочной структурой и ролевой матрицей")
        self.ed_excel.setReadOnly(True)
        self.btn_excel = QtWidgets.QPushButton("Загрузить файл")
        self.btn_excel.setObjectName("secondaryAction")
        self.cmb_sheet = QtWidgets.QComboBox()
        self.cmb_sheet.setPlaceholderText("Сначала выберите Excel-файл")
        self.cmb_sheet.setEnabled(False)
        self.btn_load_sheets = QtWidgets.QPushButton("Обновить листы")
        self.btn_load_sheets.setObjectName("secondaryAction")
        excel_grid.addWidget(self._field_label("Excel-файл"), 0, 0)
        excel_grid.addWidget(self.ed_excel, 1, 0, 1, 2)
        excel_grid.addWidget(self.btn_excel, 1, 2)
        excel_grid.addWidget(self._field_label("Лист Excel"), 2, 0)
        excel_grid.addWidget(self.cmb_sheet, 3, 0, 1, 2)
        excel_grid.addWidget(self.btn_load_sheets, 3, 2)
        excel_grid.setColumnStretch(0, 1)
        excel_grid.setColumnStretch(1, 1)
        excel_layout.addLayout(excel_grid)

        rights_box = QtWidgets.QFrame()
        rights_box.setObjectName("rightsLegendBox")
        rights_layout = QtWidgets.QHBoxLayout(rights_box)
        rights_layout.setContentsMargins(10, 7, 10, 7)
        rights_layout.setSpacing(8)
        rights_layout.addWidget(self._make_inline_icon("information.png", 16), 0, QtCore.Qt.AlignTop)
        self.lbl_rights = QtWidgets.QLabel(
            "Пусто — не менять ·  - — снять права ·  П — просмотр ·  С — скачивание ·  "
            "З — загрузка/создание ·  Р — полный доступ"
        )
        self.lbl_rights.setObjectName("rightsLegend")
        self.lbl_rights.setWordWrap(True)
        rights_layout.addWidget(self.lbl_rights, 1)
        excel_layout.addWidget(rights_box)
        root.addWidget(excel_card)

        # Разделы оформлены тем же верхним переключателем, что и в исходном архиве.
        self.tabs = QtWidgets.QTabWidget()
        self.tabs.setTabBar(NoWheelTabBar())
        self.tabs.setObjectName("workspaceTabs")
        self.tabs.setDocumentMode(True)
        self.tabs.setMovable(False)
        self.tabs.tabBar().setDrawBase(False)
        self.tabs.tabBar().setExpanding(True)
        self.tabs.tabBar().setUsesScrollButtons(False)
        # Вкладки не должны растягиваться на всю оставшуюся высоту окна.
        # Их высота подстраивается под содержимое текущего раздела, а если
        # страница целиком не помещается — прокручивается общий QScrollArea.
        self.tabs.setSizePolicy(QtWidgets.QSizePolicy.Expanding, QtWidgets.QSizePolicy.Fixed)
        root.addWidget(self.tabs, 0, QtCore.Qt.AlignTop)

        permissions_tab = QtWidgets.QWidget()
        permissions_tab.setObjectName("tabPage")
        permissions_layout = QtWidgets.QVBoxLayout(permissions_tab)
        permissions_layout.setContentsMargins(2, 12, 2, 2)
        permissions_layout.setSpacing(12)

        permissions_card = QtWidgets.QFrame()
        permissions_card.setObjectName("card")
        permissions_card.setSizePolicy(QtWidgets.QSizePolicy.Expanding, QtWidgets.QSizePolicy.Maximum)
        permissions_card_layout = QtWidgets.QVBoxLayout(permissions_card)
        permissions_card_layout.setContentsMargins(14, 12, 14, 14)
        permissions_card_layout.setSpacing(9)
        self._add_card_header(permissions_card_layout, "Предпросмотр прав", "access_icon_variant_1.png")

        info_banner = QtWidgets.QFrame()
        info_banner.setObjectName("infoBanner")
        info_banner.setSizePolicy(QtWidgets.QSizePolicy.Expanding, QtWidgets.QSizePolicy.Maximum)
        info_layout = QtWidgets.QHBoxLayout(info_banner)
        info_layout.setContentsMargins(10, 7, 10, 7)
        info_layout.setSpacing(8)
        info_layout.addWidget(self._make_inline_icon("information.png", 16), 0, QtCore.Qt.AlignTop)
        hint = QtWidgets.QLabel(
            "Проверьте права перед применением. Программа покажет только те строки, где нужно что-то изменить."
        )
        hint.setObjectName("infoText")
        hint.setWordWrap(True)
        info_layout.addWidget(hint, 1)
        permissions_card_layout.addWidget(info_banner)

        self.tbl_plan = QtWidgets.QTableWidget(0, 7)
        self.tbl_plan.setObjectName("matrixTable")
        self.tbl_plan.setHorizontalHeaderLabels(["Статус", "Строка", "Путь Excel", "Папка SGNL", "К применению / всего", "Права", "Комментарий"])
        self.tbl_plan.verticalHeader().setVisible(False)
        self.tbl_plan.setEditTriggers(QtWidgets.QAbstractItemView.NoEditTriggers)
        self.tbl_plan.setSelectionBehavior(QtWidgets.QAbstractItemView.SelectRows)
        self.tbl_plan.setAlternatingRowColors(True)
        self.tbl_plan.setWordWrap(False)
        configure_preview_table(self.tbl_plan, min_height=135, max_height=220)
        header = self.tbl_plan.horizontalHeader()
        header.setSectionResizeMode(0, QtWidgets.QHeaderView.ResizeToContents)
        header.setSectionResizeMode(1, QtWidgets.QHeaderView.ResizeToContents)
        header.setSectionResizeMode(2, QtWidgets.QHeaderView.Stretch)
        header.setSectionResizeMode(3, QtWidgets.QHeaderView.Stretch)
        header.setSectionResizeMode(4, QtWidgets.QHeaderView.ResizeToContents)
        header.setSectionResizeMode(5, QtWidgets.QHeaderView.Stretch)
        header.setSectionResizeMode(6, QtWidgets.QHeaderView.Stretch)
        permissions_card_layout.addWidget(self.tbl_plan)

        permissions_footer = QtWidgets.QHBoxLayout()
        self.btn_log = QtWidgets.QPushButton("Открыть лог")
        self.btn_log.setObjectName("logButton")
        self.btn_permissions_template = QtWidgets.QPushButton("Скачать шаблон Excel")
        self.btn_permissions_template.setObjectName("downloadButton")
        self.btn_permissions_details = make_details_button()
        self.btn_permissions_details.setEnabled(False)
        self.btn_permissions_details.clicked.connect(
            lambda: show_table_details(
                self, self.tbl_plan, "Ролевая матрица — Подробнее",
                "Полный результат сверки Excel с актуальными правами SGNL."
            )
        )
        self.btn_preview = QtWidgets.QPushButton("Сверить с SGNL")
        self.btn_preview.setObjectName("secondaryAction")
        self.btn_apply = QtWidgets.QPushButton("Применить права")
        self.btn_apply.setObjectName("orangeAction")
        self.btn_apply.setToolTip("Если предпросмотр ещё не построен, программа сначала выполнит проверку автоматически.")
        permissions_footer.addWidget(self.btn_log)
        permissions_footer.addWidget(self.btn_permissions_template)
        permissions_footer.addStretch(1)
        permissions_footer.addWidget(self.btn_permissions_details)
        permissions_footer.addWidget(self.btn_preview)
        permissions_footer.addWidget(self.btn_apply)
        permissions_card_layout.addLayout(permissions_footer)
        permissions_layout.addWidget(permissions_card, 0, QtCore.Qt.AlignTop)
        self.tabs.addTab(permissions_tab, "Ролевая матрица")

        folders_tab = QtWidgets.QWidget()
        folders_tab.setObjectName("tabPage")
        folders_layout = QtWidgets.QVBoxLayout(folders_tab)
        folders_layout.setContentsMargins(2, 12, 2, 2)
        folders_layout.setSpacing(12)

        folders_card = QtWidgets.QFrame()
        folders_card.setObjectName("card")
        folders_card.setSizePolicy(QtWidgets.QSizePolicy.Expanding, QtWidgets.QSizePolicy.Maximum)
        folders_card_layout = QtWidgets.QVBoxLayout(folders_card)
        folders_card_layout.setContentsMargins(14, 12, 14, 14)
        folders_card_layout.setSpacing(9)
        self._add_card_header(folders_card_layout, "Предпросмотр папочной структуры", "folder_icon_variant_1.png")

        folder_banner = QtWidgets.QFrame()
        folder_banner.setObjectName("infoBanner")
        folder_banner.setSizePolicy(QtWidgets.QSizePolicy.Expanding, QtWidgets.QSizePolicy.Maximum)
        folder_banner_layout = QtWidgets.QHBoxLayout(folder_banner)
        folder_banner_layout.setContentsMargins(10, 7, 10, 7)
        folder_banner_layout.setSpacing(8)
        folder_banner_layout.addWidget(self._make_inline_icon("information.png", 16), 0, QtCore.Qt.AlignTop)
        folder_hint = QtWidgets.QLabel(
            "Проверьте структуру перед созданием. Будут добавлены только отсутствующие папки; существующие останутся без изменений."
        )
        folder_hint.setObjectName("infoText")
        folder_hint.setWordWrap(True)
        folder_banner_layout.addWidget(folder_hint, 1)
        folders_card_layout.addWidget(folder_banner)

        self.tbl_folders = QtWidgets.QTableWidget(0, 6)
        self.tbl_folders.setObjectName("matrixTable")
        self.tbl_folders.setHorizontalHeaderLabels(["Статус", "Строка", "Путь Excel", "Родитель", "ID папки", "Комментарий"])
        self.tbl_folders.verticalHeader().setVisible(False)
        self.tbl_folders.setEditTriggers(QtWidgets.QAbstractItemView.NoEditTriggers)
        self.tbl_folders.setSelectionBehavior(QtWidgets.QAbstractItemView.SelectRows)
        self.tbl_folders.setAlternatingRowColors(True)
        self.tbl_folders.setWordWrap(False)
        configure_preview_table(self.tbl_folders, min_height=135, max_height=220)
        folder_header = self.tbl_folders.horizontalHeader()
        folder_header.setSectionResizeMode(0, QtWidgets.QHeaderView.ResizeToContents)
        folder_header.setSectionResizeMode(1, QtWidgets.QHeaderView.ResizeToContents)
        folder_header.setSectionResizeMode(2, QtWidgets.QHeaderView.Stretch)
        folder_header.setSectionResizeMode(3, QtWidgets.QHeaderView.Stretch)
        folder_header.setSectionResizeMode(4, QtWidgets.QHeaderView.ResizeToContents)
        folder_header.setSectionResizeMode(5, QtWidgets.QHeaderView.Stretch)
        folders_card_layout.addWidget(self.tbl_folders)

        folders_footer = QtWidgets.QHBoxLayout()
        self.btn_folder_log = QtWidgets.QPushButton("Открыть лог")
        self.btn_folder_log.setObjectName("logButton")
        self.btn_folder_template = QtWidgets.QPushButton("Скачать шаблон Excel")
        self.btn_folder_template.setObjectName("downloadButton")
        self.btn_folder_load_excel = QtWidgets.QPushButton("Загрузить из Excel")
        self.btn_folder_load_excel.setObjectName("downloadButton")
        self.btn_folder_details = make_details_button()
        self.btn_folder_details.setEnabled(False)
        self.btn_folder_details.clicked.connect(
            lambda: show_table_details(
                self, self.tbl_folders, "Папочная структура — Подробнее",
                "Полный результат сверки требуемой структуры Excel с проектом SGNL."
            )
        )
        self.btn_folder_preview = QtWidgets.QPushButton("Сверить с SGNL")
        self.btn_folder_preview.setObjectName("secondaryAction")
        self.btn_folder_create = QtWidgets.QPushButton("Создать папки")
        self.btn_folder_create.setObjectName("orangeAction")
        folders_footer.addWidget(self.btn_folder_log)
        folders_footer.addWidget(self.btn_folder_template)
        folders_footer.addStretch(1)
        folders_footer.addWidget(self.btn_folder_details)
        folders_footer.addWidget(self.btn_folder_load_excel)
        folders_footer.addWidget(self.btn_folder_preview)
        folders_footer.addWidget(self.btn_folder_create)
        folders_card_layout.addLayout(folders_footer)
        folders_layout.addWidget(folders_card, 0, QtCore.Qt.AlignTop)
        self.tabs.addTab(folders_tab, "Папочная структура")

        # Атрибуты проекта — отдельный раздел, использующий ту же авторизацию и выбранный проект.
        attributes_tab = QtWidgets.QWidget()
        attributes_tab.setObjectName("tabPage")
        attributes_layout = QtWidgets.QVBoxLayout(attributes_tab)
        attributes_layout.setContentsMargins(2, 12, 2, 2)
        attributes_layout.setSpacing(12)

        attr_create_card = QtWidgets.QFrame()
        attr_create_card.setObjectName("card")
        attr_create_card.setSizePolicy(QtWidgets.QSizePolicy.Expanding, QtWidgets.QSizePolicy.Maximum)
        attr_create_layout = QtWidgets.QVBoxLayout(attr_create_card)
        attr_create_layout.setContentsMargins(14, 12, 14, 14)
        attr_create_layout.setSpacing(10)
        self._add_card_header(attr_create_layout, "Создание атрибутов", "links.png")

        attr_banner = QtWidgets.QFrame()
        attr_banner.setObjectName("infoBanner")
        attr_banner.setSizePolicy(QtWidgets.QSizePolicy.Expanding, QtWidgets.QSizePolicy.Maximum)
        attr_banner_layout = QtWidgets.QHBoxLayout(attr_banner)
        attr_banner_layout.setContentsMargins(10, 7, 10, 7)
        attr_banner_layout.setSpacing(8)
        attr_banner_layout.addWidget(self._make_inline_icon("information.png", 16), 0, QtCore.Qt.AlignTop)
        attr_hint = QtWidgets.QLabel(
            "Можно создать атрибут типа «Текст», «Да / Нет», «Дата» или «Список». "
            "Добавьте нужные атрибуты в очередь, затем нажмите «Создать атрибуты»."
        )
        attr_hint.setObjectName("infoText")
        attr_hint.setWordWrap(True)
        attr_banner_layout.addWidget(attr_hint, 1)
        attr_create_layout.addWidget(attr_banner)

        attr_form = QtWidgets.QGridLayout()
        attr_form.setHorizontalSpacing(10)
        attr_form.setVerticalSpacing(6)
        self.ed_attr_name = QtWidgets.QLineEdit()
        self.ed_attr_name.setPlaceholderText("Например: Раздел РД")
        self.cmb_attr_type = QtWidgets.QComboBox()
        self.cmb_attr_type.addItem("Текст", "Text")
        self.cmb_attr_type.addItem("Да / Нет", "Bool")
        self.cmb_attr_type.addItem("Дата", "Date")
        self.cmb_attr_type.addItem("Список", "List")
        self.cmb_attr_data_type = QtWidgets.QComboBox()
        self.cmb_attr_data_type.addItem("Без области", "")
        self.cmb_attr_data_type.addItem("Элемент / файл", "Item")
        self.cmb_attr_data_type.addItem("Папка", "Folder")
        self.chk_attr_required = QtWidgets.QCheckBox("Обязательный атрибут")
        self.ed_attr_values = QtWidgets.QTextEdit()
        self.ed_attr_values.setPlaceholderText("Значения списка — по одному в строке или через ;")
        self.ed_attr_values.setFixedHeight(72)
        self.ed_attr_values.setEnabled(False)

        attr_form.addWidget(self._field_label("Название"), 0, 0)
        attr_form.addWidget(self._field_label("Тип"), 0, 1)
        attr_form.addWidget(self._field_label("Область"), 0, 2)
        attr_form.addWidget(self.ed_attr_name, 1, 0)
        attr_form.addWidget(self.cmb_attr_type, 1, 1)
        attr_form.addWidget(self.cmb_attr_data_type, 1, 2)
        attr_form.addWidget(self.chk_attr_required, 2, 0, 1, 2)
        attr_form.addWidget(self._field_label("Варианты для типа «Список»"), 3, 0, 1, 3)
        attr_form.addWidget(self.ed_attr_values, 4, 0, 1, 3)
        attr_form.setColumnStretch(0, 2)
        attr_form.setColumnStretch(1, 1)
        attr_form.setColumnStretch(2, 2)
        attr_create_layout.addLayout(attr_form)

        attr_queue_actions = QtWidgets.QHBoxLayout()
        self.btn_attr_add = QtWidgets.QPushButton("Добавить в очередь")
        self.btn_attr_add.setObjectName("secondaryAction")
        self.btn_attr_remove = QtWidgets.QPushButton("Удалить выбранные")
        self.btn_attr_remove.setObjectName("logButton")
        attr_queue_actions.addWidget(self.btn_attr_add)
        attr_queue_actions.addWidget(self.btn_attr_remove)
        attr_queue_actions.addStretch(1)
        attr_create_layout.addLayout(attr_queue_actions)

        self.tbl_attribute_queue = QtWidgets.QTableWidget(0, 6)
        self.tbl_attribute_queue.setObjectName("matrixTable")
        self.tbl_attribute_queue.setHorizontalHeaderLabels(
            ["Название", "Тип", "Область", "Обязательный", "Значения", "Статус"]
        )
        self.tbl_attribute_queue.verticalHeader().setVisible(False)
        self.tbl_attribute_queue.setEditTriggers(QtWidgets.QAbstractItemView.NoEditTriggers)
        self.tbl_attribute_queue.setSelectionBehavior(QtWidgets.QAbstractItemView.SelectRows)
        self.tbl_attribute_queue.setSelectionMode(QtWidgets.QAbstractItemView.ExtendedSelection)
        self.tbl_attribute_queue.setAlternatingRowColors(True)
        configure_preview_table(self.tbl_attribute_queue, min_height=110, max_height=170)
        self.tbl_attribute_queue.setSelectionMode(QtWidgets.QAbstractItemView.ExtendedSelection)
        attr_queue_header = self.tbl_attribute_queue.horizontalHeader()
        attr_queue_header.setSectionResizeMode(0, QtWidgets.QHeaderView.Stretch)
        attr_queue_header.setSectionResizeMode(1, QtWidgets.QHeaderView.ResizeToContents)
        attr_queue_header.setSectionResizeMode(2, QtWidgets.QHeaderView.Stretch)
        attr_queue_header.setSectionResizeMode(3, QtWidgets.QHeaderView.ResizeToContents)
        attr_queue_header.setSectionResizeMode(4, QtWidgets.QHeaderView.Stretch)
        attr_queue_header.setSectionResizeMode(5, QtWidgets.QHeaderView.ResizeToContents)
        attr_create_layout.addWidget(self.tbl_attribute_queue)
        attributes_layout.addWidget(attr_create_card)

        attr_existing_card = QtWidgets.QFrame()
        attr_existing_card.setObjectName("card")
        attr_existing_card.setSizePolicy(QtWidgets.QSizePolicy.Expanding, QtWidgets.QSizePolicy.Maximum)
        attr_existing_layout = QtWidgets.QVBoxLayout(attr_existing_card)
        attr_existing_layout.setContentsMargins(14, 12, 14, 14)
        attr_existing_layout.setSpacing(9)
        self._add_card_header(attr_existing_layout, "Атрибуты текущего проекта", "comparison.png")

        self.tbl_attribute_existing = QtWidgets.QTableWidget(0, 6)
        self.tbl_attribute_existing.setObjectName("matrixTable")
        self.tbl_attribute_existing.setHorizontalHeaderLabels(
            ["Название", "Тип", "Область", "Обязательный", "Значения", "ID"]
        )
        self.tbl_attribute_existing.verticalHeader().setVisible(False)
        self.tbl_attribute_existing.setEditTriggers(QtWidgets.QAbstractItemView.NoEditTriggers)
        self.tbl_attribute_existing.setSelectionBehavior(QtWidgets.QAbstractItemView.SelectRows)
        self.tbl_attribute_existing.setAlternatingRowColors(True)
        configure_preview_table(self.tbl_attribute_existing, min_height=120, max_height=190)
        attr_existing_header = self.tbl_attribute_existing.horizontalHeader()
        attr_existing_header.setSectionResizeMode(0, QtWidgets.QHeaderView.Stretch)
        attr_existing_header.setSectionResizeMode(1, QtWidgets.QHeaderView.ResizeToContents)
        attr_existing_header.setSectionResizeMode(2, QtWidgets.QHeaderView.Stretch)
        attr_existing_header.setSectionResizeMode(3, QtWidgets.QHeaderView.ResizeToContents)
        attr_existing_header.setSectionResizeMode(4, QtWidgets.QHeaderView.Stretch)
        attr_existing_header.setSectionResizeMode(5, QtWidgets.QHeaderView.ResizeToContents)
        attr_existing_layout.addWidget(self.tbl_attribute_existing)

        attr_footer = QtWidgets.QHBoxLayout()
        self.btn_attr_log = QtWidgets.QPushButton("Открыть лог")
        self.btn_attr_log.setObjectName("logButton")
        self.btn_attr_template = QtWidgets.QPushButton("Скачать шаблон Excel")
        self.btn_attr_template.setObjectName("downloadButton")
        self.btn_attr_refresh = QtWidgets.QPushButton("Обновить список")
        self.btn_attr_refresh.setObjectName("secondaryAction")
        self.btn_attr_create = QtWidgets.QPushButton("Создать атрибуты")
        self.btn_attr_create.setObjectName("orangeAction")
        attr_footer.addWidget(self.btn_attr_log)
        attr_footer.addWidget(self.btn_attr_template)
        attr_footer.addStretch(1)
        attr_footer.addWidget(self.btn_attr_refresh)
        attr_footer.addWidget(self.btn_attr_create)
        attr_existing_layout.addLayout(attr_footer)
        attributes_layout.addWidget(attr_existing_card, 0, QtCore.Qt.AlignTop)
        self.attributes_tab_index = self.tabs.addTab(attributes_tab, "Атрибуты")

        self.statusBar().showMessage("Готово")

        self.btn_login.clicked.connect(self._login)
        self.btn_load_lists.clicked.connect(self._load_catalogs)
        self.btn_toggle_token.clicked.connect(self._toggle_token)
        self.btn_check.clicked.connect(self._check_access)
        self.btn_excel.clicked.connect(self._pick_excel)
        self.btn_load_sheets.clicked.connect(lambda: self._load_excel_sheets(auto_load_structure=True))
        self.cmb_sheet.currentIndexChanged.connect(self._on_excel_sheet_changed)
        self.btn_log.clicked.connect(self._show_log)
        self.btn_folder_log.clicked.connect(self._show_log)
        self.btn_permissions_template.clicked.connect(lambda: self._save_excel_template("permissions"))
        self.btn_folder_template.clicked.connect(lambda: self._save_excel_template("folders"))
        self.btn_attr_template.clicked.connect(lambda: self._save_excel_template("attributes"))
        self.btn_preview.clicked.connect(self._preview)
        self.btn_apply.clicked.connect(self._apply_permissions)
        self.btn_folder_load_excel.clicked.connect(self._load_folder_structure_from_excel)
        self.btn_folder_preview.clicked.connect(self._preview_folder_structure)
        self.btn_folder_create.clicked.connect(self._create_folder_structure)
        self.btn_attr_log.clicked.connect(self._show_log)
        self.btn_attr_refresh.clicked.connect(self._load_attribute_types)
        self.btn_attr_add.clicked.connect(self._add_attribute_draft)
        self.btn_attr_remove.clicked.connect(self._remove_selected_attribute_drafts)
        self.btn_attr_create.clicked.connect(self._create_attributes)
        self.cmb_attr_type.currentIndexChanged.connect(self._on_attribute_type_changed)
        self.tabs.currentChanged.connect(self._on_workspace_tab_changed)
        self.tabs.currentChanged.connect(self._schedule_workspace_fit)
        self._schedule_workspace_fit()
        self.cmb_company.currentIndexChanged.connect(self._on_company_changed)
        self.cmb_project.currentIndexChanged.connect(self._on_project_changed)
        self.ed_excel.textChanged.connect(self._invalidate_plan)
        self.ed_project_id.textChanged.connect(self._invalidate_plan)
        self.ed_company_id.textChanged.connect(self._invalidate_plan)
        self.ed_docs_url.textChanged.connect(self._invalidate_plan)
        self.ed_hub_url.textChanged.connect(self._invalidate_plan)

        # Резервный интерактивный вход показывается только когда SGNL требует CAPTCHA.
        self._build_auth_overlay(central)

    def _build_auth_overlay(self, parent: QtWidgets.QWidget):
        self.auth_overlay = QtWidgets.QFrame(parent)
        self.auth_overlay.setObjectName("authOverlay")
        self.auth_overlay.setAutoFillBackground(True)
        self.auth_overlay.hide()

        layout = QtWidgets.QVBoxLayout(self.auth_overlay)
        layout.setContentsMargins(18, 14, 18, 14)
        layout.setSpacing(10)

        top = QtWidgets.QHBoxLayout()
        title = QtWidgets.QLabel("Проверка CAPTCHA SGNL")
        title.setObjectName("authTitle")
        self.lbl_auth_page_status = QtWidgets.QLabel("Открываю защищённую форму…")
        self.lbl_auth_page_status.setObjectName("hintLabel")
        self.btn_auth_cancel = QtWidgets.QPushButton("Назад")
        self.btn_auth_cancel.setObjectName("secondaryButton")
        self.btn_auth_cancel.clicked.connect(self._cancel_auth_overlay)
        top.addWidget(title)
        top.addWidget(self.lbl_auth_page_status, 1)
        top.addWidget(self.btn_auth_cancel)
        layout.addLayout(top)

        hint = QtWidgets.QLabel(
            "SGNL запросил CAPTCHA. Пройдите проверку и нажмите кнопку входа в форме ниже. "
            "Логин и пароль подставляются автоматически. Это встроенная часть программы — "
            "Chrome или Edge не запускаются."
        )
        hint.setObjectName("hintLabel")
        hint.setWordWrap(True)
        layout.addWidget(hint)

        if WEBENGINE_AVAILABLE and QtWebEngineWidgets is not None:
            self.auth_view = QtWebEngineWidgets.QWebEngineView(self.auth_overlay)
            self.auth_view.urlChanged.connect(self._on_embedded_auth_url_changed)
            self.auth_view.loadFinished.connect(self._on_embedded_auth_load_finished)
            layout.addWidget(self.auth_view, 1)
        else:
            self.auth_view = None
            error = QtWidgets.QLabel(
                "Для CAPTCHA нужен компонент Qt WebEngine.\n\n"
                "PySide6: pip install -U PySide6\n"
                "PyQt5: pip install PyQtWebEngine"
            )
            error.setWordWrap(True)
            error.setObjectName("badStatus")
            layout.addWidget(error, 1, QtCore.Qt.AlignCenter)

        self.auth_overlay.setGeometry(parent.rect())

    def resizeEvent(self, event):
        super().resizeEvent(event)
        overlay = getattr(self, "auth_overlay", None)
        central = self.centralWidget()
        if overlay is not None and central is not None:
            overlay.setGeometry(central.rect())

    def _field_label(self, text: str) -> QtWidgets.QLabel:
        label = QtWidgets.QLabel(text)
        label.setObjectName("fieldLabel")
        return label

    def _apply_stylesheet(self):
        self.setStyleSheet(build_qss(self.is_dark_theme, self.asset_dir) + shared_style_overrides(self.is_dark_theme))
        if hasattr(self, "theme_toggle"):
            self.theme_toggle.blockSignals(True)
            self.theme_toggle.setChecked(self.is_dark_theme, animate=False)
            self.theme_toggle.blockSignals(False)
        self._refresh_action_icons()
        if self.log_dialog is not None:
            self.log_dialog.setStyleSheet(self.styleSheet())
        self._set_native_titlebar_theme(self.is_dark_theme)
        QtCore.QTimer.singleShot(0, lambda: self._set_native_titlebar_theme(self.is_dark_theme))

    def _set_native_titlebar_theme(self, dark: bool):
        apply_windows_titlebar_theme(self, bool(dark))

    def _toggle_token(self):
        if self.ed_token.echoMode() == QtWidgets.QLineEdit.Password:
            self.ed_token.setEchoMode(QtWidgets.QLineEdit.Normal)
        else:
            self.ed_token.setEchoMode(QtWidgets.QLineEdit.Password)
        self._update_sensitive_icons()

    def _log(self, text: str):
        chunk = str(text)
        if chunk and not chunk.endswith("\n"):
            chunk += "\n"
        self._log_lines.append(chunk)
        if self.log_dialog and self.log_dialog.isVisible():
            self.log_dialog.append(chunk)

    def _show_log(self):
        if self.log_dialog is None:
            self.log_dialog = LogDialog(self)
        self.log_dialog.set_log_text("".join(self._log_lines) if self._log_lines else "Лог пока пуст.\n")
        self.log_dialog.show()
        self.log_dialog.raise_()
        self.log_dialog.activateWindow()

    def _ensure_log_visible(self):
        if self.log_dialog is None:
            self.log_dialog = LogDialog(self)
        if not self.log_dialog.isVisible():
            self.log_dialog.set_log_text("".join(self._log_lines) if self._log_lines else "Лог пока пуст.\n")
            self.log_dialog.show()
        self.log_dialog.raise_()
        self.log_dialog.activateWindow()

    def _set_status_label(self, text: str, ok: bool):
        self.lbl_auth.setText(text)
        self.lbl_auth.setObjectName("goodStatus" if ok else "badStatus")
        self.lbl_auth.style().unpolish(self.lbl_auth)
        self.lbl_auth.style().polish(self.lbl_auth)

    def _show_error(self, title: str, details: str):
        self._log(f"ERROR: {title}\n{details}\n")
        lower = details.casefold()
        if "winerror 10060" in lower or "connecttimeout" in lower or "connection to docs.sgnl.pro timed out" in lower:
            message = (
                "Не удалось установить соединение с docs.sgnl.pro. В этой версии включено использование "
                "системного proxy/VPN и автоматические повторные попытки. Если ошибка повторится, проверь, "
                "что VPN активен и python.exe разрешён в брандмауэре/антивирусе. Полная трассировка сохранена в логе."
            )
        elif "не удалось подключиться к" in lower:
            last_line = next((line.strip() for line in reversed(details.splitlines()) if line.strip()), details[:1200])
            message = last_line[:1800] + "\n\nПолная трассировка сохранена в логе."
        else:
            message = details[:4000]
        QtWidgets.QMessageBox.critical(self, title, message)

    def _set_busy(self, busy: bool, message: str = ""):
        widgets = [
            self.ed_docs_url,
            self.ed_hub_url,
            self.ed_email,
            self.ed_password,
            self.btn_login,
            self.cmb_company,
            self.cmb_project,
            self.btn_load_lists,
            self.ed_project_id,
            self.ed_company_id,
            self.ed_token,
            self.btn_toggle_token,
            self.btn_check,
            self.ed_excel,
            self.btn_excel,
            self.cmb_sheet,
            self.btn_load_sheets,
            self.btn_permissions_template,
            self.btn_preview,
            self.btn_apply,
            self.btn_folder_template,
            self.btn_folder_load_excel,
            self.btn_folder_details,
            self.btn_folder_preview,
            self.btn_folder_create,
            self.ed_attr_name,
            self.cmb_attr_type,
            self.cmb_attr_data_type,
            self.chk_attr_required,
            self.ed_attr_values,
            self.btn_attr_add,
            self.btn_attr_remove,
            self.btn_attr_template,
            self.btn_attr_refresh,
            self.btn_attr_create,
        ]
        for widget in widgets:
            widget.setEnabled(not busy)
        if message:
            self.statusBar().showMessage(message)
        if busy:
            if QtWidgets.QApplication.overrideCursor() is None:
                QtWidgets.QApplication.setOverrideCursor(QtCore.Qt.WaitCursor)
        else:
            while QtWidgets.QApplication.overrideCursor() is not None:
                QtWidgets.QApplication.restoreOverrideCursor()

    def _run_worker(self, worker: QtCore.QObject):
        thread = QtCore.QThread(self)
        worker.moveToThread(thread)
        thread.started.connect(worker.run)
        for signal_name in ("done", "failed", "loaded", "captcha_required"):
            signal = getattr(worker, signal_name, None)
            if signal is not None:
                try:
                    signal.connect(thread.quit)
                except Exception:
                    pass
        thread.finished.connect(lambda: self._cleanup_thread(thread, worker))
        self.active_threads.append(thread)
        self.active_workers.append(worker)
        thread.start()

    def _cleanup_thread(self, thread: QtCore.QThread, worker: QtCore.QObject):
        worker.deleteLater()
        thread.deleteLater()
        try:
            self.active_threads.remove(thread)
        except ValueError:
            pass
        try:
            self.active_workers.remove(worker)
        except ValueError:
            pass

    def _invalidate_plan(self):
        self.current_plan = []
        self.current_folder_plan = []
        self.folder_plan_verified = False
        self.tbl_plan.setRowCount(0)
        if hasattr(self, "btn_permissions_details"):
            self.btn_permissions_details.setEnabled(False)
        if hasattr(self, "tbl_folders"):
            self.tbl_folders.setRowCount(0)
        if hasattr(self, "btn_folder_details"):
            self.btn_folder_details.setEnabled(False)

    # ------------------------------------------------------------------
    # Data / actions
    # ------------------------------------------------------------------
    def _combo_id_by_text(self, combo: QtWidgets.QComboBox, text: str) -> str:
        norm = normalize_text(text)
        for idx in range(combo.count()):
            if normalize_text(combo.itemText(idx)) == norm:
                return clean_text(combo.itemData(idx))
        return ""

    def _current_company_id(self) -> str:
        data = self.cmb_company.currentData()
        value = clean_text(data) or self._combo_id_by_text(self.cmb_company, self.cmb_company.currentText()) or clean_text(self.ed_company_id.text())
        if value:
            self.ed_company_id.setText(value)
        return value

    def _current_project_id(self) -> str:
        data = self.cmb_project.currentData()
        value = clean_text(data) or self._combo_id_by_text(self.cmb_project, self.cmb_project.currentText()) or clean_text(self.ed_project_id.text())
        if value:
            self.ed_project_id.setText(value)
        return value

    def _set_combo_value(self, combo: QtWidgets.QComboBox, name: str, item_id: str):
        name = clean_text(name) or clean_text(item_id)
        item_id = clean_text(item_id)
        if not name and not item_id:
            return
        for idx in range(combo.count()):
            same_id = item_id and clean_text(combo.itemData(idx)) == item_id
            same_name = normalize_text(combo.itemText(idx)) == normalize_text(name)
            if same_id or same_name:
                combo.setCurrentIndex(idx)
                if item_id and not clean_text(combo.itemData(idx)):
                    combo.setItemData(idx, item_id)
                return
        combo.addItem(name, item_id)
        combo.setCurrentIndex(combo.count() - 1)

    def _on_company_changed(self, *_args):
        self._current_company_id()
        self._invalidate_plan()

    def _on_project_changed(self, *_args):
        self._current_project_id()
        self._invalidate_plan()
        self.current_attribute_types = []
        self._attribute_auto_loaded_project = ""
        if hasattr(self, "tbl_attribute_existing"):
            self.tbl_attribute_existing.setRowCount(0)
        if hasattr(self, "attributes_tab_index") and self.tabs.currentIndex() == self.attributes_tab_index:
            self._on_workspace_tab_changed(self.attributes_tab_index)

    @staticmethod
    def _attribute_type_label(value: str) -> str:
        return {
            "Text": "Текст",
            "Bool": "Да / Нет",
            "Date": "Дата",
            "List": "Список",
        }.get(clean_text(value), clean_text(value) or "—")

    @staticmethod
    def _attribute_data_type_label(value: str) -> str:
        return {
            "Item": "Элемент / файл",
            "Folder": "Папка",
            "": "Без области",
        }.get(clean_text(value), clean_text(value) or "Без области")

    def _on_workspace_tab_changed(self, index: int):
        is_attributes = hasattr(self, "attributes_tab_index") and index == self.attributes_tab_index
        if hasattr(self, "excel_card"):
            self.excel_card.setVisible(not is_attributes)
        if not is_attributes:
            return
        project_id = self._current_project_id()
        token = clean_text(self.ed_token.text())
        if project_id and token and self._attribute_auto_loaded_project != project_id:
            self._attribute_auto_loaded_project = project_id
            self._load_attribute_types(silent=True)

    def _on_attribute_type_changed(self, *_args):
        is_list = clean_text(self.cmb_attr_type.currentData()) == "List"
        self.ed_attr_values.setEnabled(is_list)
        if not is_list:
            self.ed_attr_values.clear()

    def _parse_attribute_values(self) -> List[str]:
        text = self.ed_attr_values.toPlainText()
        values: List[str] = []
        seen = set()
        for part in re.split(r"[;\n\r]+", text):
            value = clean_text(part)
            norm = normalize_text(value)
            if value and norm not in seen:
                seen.add(norm)
                values.append(value)
        return values

    def _add_attribute_draft(self) -> bool:
        name = clean_text(self.ed_attr_name.text())
        attribute_type = clean_text(self.cmb_attr_type.currentData())
        data_type = clean_text(self.cmb_attr_data_type.currentData())
        values = self._parse_attribute_values() if attribute_type == "List" else []
        if not name:
            QtWidgets.QMessageBox.warning(self, "Нет названия", "Укажи название атрибута.")
            return False
        if attribute_type == "List" and not values:
            QtWidgets.QMessageBox.warning(
                self,
                "Нет вариантов списка",
                "Для типа «Список» укажи хотя бы одно значение — по строкам или через ';'.",
            )
            return False
        norm_name = normalize_text(name)
        if any(normalize_text(item.get("name")) == norm_name for item in self.attribute_drafts):
            QtWidgets.QMessageBox.warning(self, "Дубликат", f"Атрибут «{name}» уже есть в очереди.")
            return False
        if any(
            not bool(item.get("deleted")) and normalize_text(item.get("name")) == norm_name
            for item in self.current_attribute_types
        ):
            QtWidgets.QMessageBox.warning(self, "Уже существует", f"Активный атрибут «{name}» уже есть в проекте.")
            return False
        self.attribute_drafts.append({
            "name": name,
            "type": attribute_type,
            "dataType": data_type,
            "isRequired": bool(self.chk_attr_required.isChecked()),
            "values": values,
        })
        self._fill_attribute_queue()
        self.ed_attr_name.clear()
        self.ed_attr_values.clear()
        self.ed_attr_name.setFocus()
        return True

    def _remove_selected_attribute_drafts(self):
        rows = sorted({index.row() for index in self.tbl_attribute_queue.selectionModel().selectedRows()}, reverse=True)
        if not rows:
            return
        for row in rows:
            if 0 <= row < len(self.attribute_drafts):
                self.attribute_drafts.pop(row)
        self._fill_attribute_queue()

    def _fill_attribute_queue(self):
        existing_names = {
            normalize_text(item.get("name"))
            for item in self.current_attribute_types
            if isinstance(item, dict) and not bool(item.get("deleted")) and clean_text(item.get("name"))
        }
        self.tbl_attribute_queue.setRowCount(len(self.attribute_drafts))
        for row, draft in enumerate(self.attribute_drafts):
            values = ", ".join(draft.get("values") or [])
            already_exists = normalize_text(draft.get("name")) in existing_names
            cells = [
                clean_text(draft.get("name")),
                self._attribute_type_label(draft.get("type")),
                self._attribute_data_type_label(draft.get("dataType")),
                "Да" if draft.get("isRequired") else "Нет",
                values or "—",
                "Уже существует" if already_exists else "Готов к созданию",
            ]
            for col, value in enumerate(cells):
                item = QtWidgets.QTableWidgetItem(value)
                if col == 5:
                    color = STATUS_DANGER if already_exists else STATUS_SUCCESS
                    item.setForeground(QtGui.QBrush(QtGui.QColor(color)))
                self.tbl_attribute_queue.setItem(row, col, item)
        self.tbl_attribute_queue.resizeRowsToContents()

    def _fill_attribute_existing(self, attributes: Sequence[dict]):
        active = [item for item in attributes if isinstance(item, dict) and not bool(item.get("deleted"))]
        active.sort(key=lambda item: (normalize_text(item.get("name")), clean_text(item.get("id"))))
        self.tbl_attribute_existing.setRowCount(len(active))
        for row, attribute in enumerate(active):
            list_values = [
                clean_text(item.get("name"))
                for item in (attribute.get("list") or [])
                if isinstance(item, dict) and not bool(item.get("deleted")) and clean_text(item.get("name"))
            ]
            cells = [
                clean_text(attribute.get("name")),
                self._attribute_type_label(attribute.get("type")),
                self._attribute_data_type_label(attribute.get("dataType")),
                "Да" if attribute.get("isRequired") else "Нет",
                ", ".join(list_values) or "—",
                clean_text(attribute.get("id")),
            ]
            for col, value in enumerate(cells):
                self.tbl_attribute_existing.setItem(row, col, QtWidgets.QTableWidgetItem(value))
        self.tbl_attribute_existing.resizeRowsToContents()

    def _load_attribute_types(self, *_args, silent: bool = False):
        project_id = self._current_project_id()
        if not project_id:
            if not silent:
                QtWidgets.QMessageBox.warning(self, "Проект не выбран", "Сначала выбери проект SGNL.")
            return
        if not clean_text(self.ed_token.text()):
            if not silent:
                QtWidgets.QMessageBox.warning(self, "Нет токена", "Сначала выполни вход в программе.")
            return
        self._set_busy(True, "Загрузка атрибутов...")
        worker = LoadAttributeTypesWorker(
            docs_url=clean_text(self.ed_docs_url.text()),
            hub_url=clean_text(self.ed_hub_url.text()),
            token=clean_text(self.ed_token.text()),
            project_id=project_id,
            cookies=self._session_cookies,
        )
        worker.log.connect(self._log)
        worker.done.connect(self._on_attribute_types_loaded)
        worker.failed.connect(self._on_attribute_worker_failed)
        self._run_worker(worker)

    def _on_attribute_types_loaded(self, payload: dict):
        self._set_busy(False, "Атрибуты загружены")
        self.current_attribute_types = list(payload.get("attributes") or [])
        self._attribute_auto_loaded_project = self._current_project_id()
        self._fill_attribute_existing(self.current_attribute_types)
        self._fill_attribute_queue()
        active_count = sum(1 for item in self.current_attribute_types if not bool(item.get("deleted")))
        self.statusBar().showMessage(f"Атрибуты проекта загружены: {active_count}")

    def _create_attributes(self):
        # Если пользователь заполнил форму, но не нажал «Добавить в очередь», добавляем её автоматически.
        if clean_text(self.ed_attr_name.text()):
            if not self._add_attribute_draft():
                return
        if not self.attribute_drafts:
            QtWidgets.QMessageBox.information(self, "Очередь пуста", "Добавь хотя бы один атрибут в очередь.")
            return
        project_id = self._current_project_id()
        if not project_id:
            QtWidgets.QMessageBox.warning(self, "Проект не выбран", "Сначала выбери проект SGNL.")
            return
        if not clean_text(self.ed_token.text()):
            QtWidgets.QMessageBox.warning(self, "Нет токена", "Сначала выполни вход в программе.")
            return
        answer = QtWidgets.QMessageBox.question(
            self,
            "Создать атрибуты",
            f"Будет обработано атрибутов: {len(self.attribute_drafts)}.\n\n"
            "Перед созданием программа повторно проверит существующие имена и пропустит дубликаты. Продолжить?",
            QtWidgets.QMessageBox.Yes | QtWidgets.QMessageBox.No,
            QtWidgets.QMessageBox.No,
        )
        if answer != QtWidgets.QMessageBox.Yes:
            return
        self._ensure_log_visible()
        self._set_busy(True, "Создание атрибутов...")
        worker = CreateAttributeTypesWorker(
            docs_url=clean_text(self.ed_docs_url.text()),
            hub_url=clean_text(self.ed_hub_url.text()),
            token=clean_text(self.ed_token.text()),
            project_id=project_id,
            drafts=self.attribute_drafts,
            cookies=self._session_cookies,
        )
        worker.log.connect(self._log)
        worker.done.connect(self._on_attributes_created)
        worker.failed.connect(self._on_attribute_worker_failed)
        self._run_worker(worker)

    def _on_attributes_created(self, payload: dict):
        self._set_busy(False, "Атрибуты обработаны")
        stats = dict(payload.get("stats") or {})
        created_names = {normalize_text(name) for name in stats.get("created_names") or []}
        # Убираем только успешно созданные. Ошибочные/пропущенные остаются видимыми для контроля.
        self.attribute_drafts = [
            draft for draft in self.attribute_drafts
            if normalize_text(draft.get("name")) not in created_names
        ]
        self._fill_attribute_queue()
        self.current_attribute_types = list(payload.get("attributes") or [])
        self._fill_attribute_existing(self.current_attribute_types)
        QtWidgets.QMessageBox.information(
            self,
            "Готово",
            f"Создано атрибутов: {stats.get('created', 0)}\n"
            f"Уже существовало: {stats.get('skipped', 0)}\n"
            f"Ошибок: {stats.get('failed', 0)}",
        )

    def _on_attribute_worker_failed(self, error: str):
        self._attribute_auto_loaded_project = ""
        self._set_busy(False, "Ошибка")
        self._show_error("Ошибка работы с атрибутами", error)

    def _login(self):
        docs_url = clean_text(self.ed_docs_url.text())
        hub_url = clean_text(self.ed_hub_url.text())
        email = clean_text(self.ed_email.text())
        password = self.ed_password.text()

        if not docs_url or not hub_url:
            QtWidgets.QMessageBox.warning(self, "Не хватает данных", "Заполни Docs URL и Hub URL.")
            return
        if not email or not password:
            QtWidgets.QMessageBox.warning(self, "Не хватает данных", "Укажи email и пароль SGNL.")
            return

        self._ensure_log_visible()
        self._set_busy(True, "Авторизация через SGNL...")
        self._log("Запускаю авторизацию без браузера по потоку из HAR\n")

        worker = LoginWorker(
            docs_url=docs_url,
            hub_url=hub_url,
            email=email,
            password=password,
        )
        worker.log.connect(self._log)
        worker.done.connect(self._on_login_done)
        worker.captcha_required.connect(self._on_captcha_required)
        worker.failed.connect(self._on_login_failed)
        self._run_worker(worker)

    def _on_captcha_required(self, message: str):
        self._set_busy(False, "Требуется CAPTCHA")
        self._set_status_label("SGNL запросил CAPTCHA", False)
        self._log((message or "SGNL запросил CAPTCHA") + "\n")

        if not WEBENGINE_AVAILABLE or getattr(self, "auth_view", None) is None:
            self._show_error(
                "Нужен Qt WebEngine",
                "SGNL потребовал CAPTCHA. Для её прохождения внутри программы установите компонент:\n\n"
                "pip install -U PySide6\n\n"
                "После установки перезапустите программу и повторите вход. "
                "Внешний браузер открываться не будет.",
            )
            return

        if self.log_dialog is not None and self.log_dialog.isVisible():
            self.log_dialog.hide()

        self._oauth_email = clean_text(self.ed_email.text())
        self._oauth_password = self.ed_password.text()
        client = SgnlClient(clean_text(self.ed_docs_url.text()), clean_text(self.ed_hub_url.text()))
        self._oauth_request = client.build_authorization_request()

        self.lbl_auth_page_status.setText("SGNL запросил CAPTCHA — пройдите проверку ниже")
        self.auth_overlay.setGeometry(self.centralWidget().rect())
        self.auth_overlay.show()
        self.auth_overlay.raise_()
        self.auth_view.setFocus()
        self.auth_view.setUrl(QtCore.QUrl(self._oauth_request["authorize_url"]))
        self.statusBar().showMessage("Ожидание прохождения CAPTCHA")

    def _cancel_auth_overlay(self):
        self._oauth_request = {}
        view = getattr(self, "auth_view", None)
        if view is not None:
            try:
                view.setUrl(QtCore.QUrl("about:blank"))
            except Exception:
                pass
        overlay = getattr(self, "auth_overlay", None)
        if overlay is not None:
            overlay.hide()
        self.statusBar().showMessage("Авторизация отменена")

    def _on_embedded_auth_url_changed(self, qurl):
        if not self._oauth_request:
            return
        url = qurl.toString()
        parsed = urlparse(url)
        expected = urlparse(self._oauth_request.get("redirect_uri", ""))
        if normalize_text(parsed.hostname or "") != normalize_text(expected.hostname or ""):
            return
        if not parsed.path.startswith(expected.path):
            return

        query = parse_qs(parsed.query)
        error = clean_text(query.get("error", [""])[0])
        if error:
            description = clean_text(query.get("error_description", [""])[0])
            self._cancel_auth_overlay()
            self._show_error("Ошибка OAuth", f"{error}: {description}".strip(": "))
            return

        code = clean_text(query.get("code", [""])[0])
        state = clean_text(query.get("state", [""])[0])
        if not code:
            return
        if state and state != clean_text(self._oauth_request.get("state")):
            self._cancel_auth_overlay()
            self._show_error(
                "Ошибка OAuth",
                "OAuth state не совпал; вход отменён из соображений безопасности.",
            )
            return

        self._finish_embedded_oauth(code, url)

    def _on_embedded_auth_load_finished(self, ok: bool):
        if not ok or getattr(self, "auth_view", None) is None:
            return

        host = self.auth_view.url().host().casefold()
        if host == "auth.sgnl.pro":
            self.lbl_auth_page_status.setText("Пройдите CAPTCHA и подтвердите вход")
        else:
            self.lbl_auth_page_status.setText("Выполняется авторизация…")

        if host != "auth.sgnl.pro" or (not self._oauth_email and not self._oauth_password):
            return

        email_json = json.dumps(self._oauth_email)
        password_json = json.dumps(self._oauth_password)
        script = f"""
        (() => {{
          const fill = () => {{
            const email = document.querySelector(
              'input[type="email"], input[name="email"], input[autocomplete="username"]'
            );
            const password = document.querySelector(
              'input[type="password"], input[name="password"], input[autocomplete="current-password"]'
            );
            const setValue = (element, value) => {{
              if (!element || !value) return;
              const prototype = Object.getPrototypeOf(element);
              const descriptor = Object.getOwnPropertyDescriptor(prototype, 'value');
              if (descriptor && descriptor.set) descriptor.set.call(element, value);
              else element.value = value;
              element.dispatchEvent(new Event('input', {{bubbles:true}}));
              element.dispatchEvent(new Event('change', {{bubbles:true}}));
            }};
            setValue(email, {email_json});
            setValue(password, {password_json});
          }};
          fill();
          [300, 800, 1600, 3000].forEach(delay => setTimeout(fill, delay));
        }})();
        """
        try:
            self.auth_view.page().runJavaScript(script)
        except Exception:
            pass

    def _finish_embedded_oauth(self, code: str, callback_url: str):
        oauth_request = dict(self._oauth_request)
        self._oauth_request = {}
        view = getattr(self, "auth_view", None)
        if view is not None:
            try:
                view.setUrl(QtCore.QUrl("about:blank"))
            except Exception:
                pass
        self.auth_overlay.hide()

        self._ensure_log_visible()
        self._set_busy(True, "Получение access_token...")
        worker = OAuthTokenWorker(
            docs_url=clean_text(self.ed_docs_url.text()),
            hub_url=clean_text(self.ed_hub_url.text()),
            code=code,
            verifier=oauth_request["verifier"],
            redirect_uri=oauth_request["redirect_uri"],
            callback_url=callback_url,
        )
        worker.log.connect(self._log)
        worker.done.connect(self._on_login_done)
        worker.failed.connect(self._on_login_failed)
        self._run_worker(worker)

    def _on_login_done(self, payload: dict):
        token = clean_text(payload.get("token"))
        if token:
            self.ed_token.setText(token)
            self._loaded_token = token
        self._session_cookies = dict(payload.get("cookies") or {})
        self._log(f"  cookies сессии сохранены: {len(self._session_cookies)}\n")
        self._set_busy(False, "Авторизация выполнена")
        self._set_status_label("Вход выполнен, загружаю компании/проекты", True)
        self._load_catalogs()

    def _on_login_failed(self, error: str):
        self._set_busy(False, "Ошибка")
        self._set_status_label("Ошибка входа", False)
        self._show_error("Ошибка входа", error)

    def _load_catalogs(self):
        if not clean_text(self.ed_docs_url.text()) or not clean_text(self.ed_hub_url.text()):
            QtWidgets.QMessageBox.warning(self, "Не хватает данных", "Заполни Docs URL и Hub URL.")
            return
        if not clean_text(self.ed_token.text()):
            QtWidgets.QMessageBox.warning(self, "Нет токена", "Сначала выполни вход в программе.")
            return
        self._ensure_log_visible()
        self._set_busy(True, "Загрузка компаний и проектов...")
        worker = LoadCatalogsWorker(
            docs_url=clean_text(self.ed_docs_url.text()),
            hub_url=clean_text(self.ed_hub_url.text()),
            token=clean_text(self.ed_token.text()),
            company_id=self._current_company_id(),
            cookies=self._session_cookies,
        )
        worker.log.connect(self._log)
        worker.done.connect(self._on_catalogs_loaded)
        worker.failed.connect(self._on_worker_failed_generic("Ошибка загрузки компаний/проектов"))
        self._run_worker(worker)

    def _on_catalogs_loaded(self, payload: dict):
        companies = payload.get("companies") or []
        projects = payload.get("projects") or []
        selected_company_id = clean_text(payload.get("company_id"))

        self.cmb_company.blockSignals(True)
        self.cmb_company.clear()
        for company in companies:
            self.cmb_company.addItem(clean_text(company.get("name")) or clean_text(company.get("id")), clean_text(company.get("id")))
        if selected_company_id:
            for idx in range(self.cmb_company.count()):
                if clean_text(self.cmb_company.itemData(idx)) == selected_company_id:
                    self.cmb_company.setCurrentIndex(idx)
                    break
        self.cmb_company.blockSignals(False)
        self._current_company_id()

        self.cmb_project.blockSignals(True)
        self.cmb_project.clear()
        for project in projects:
            self.cmb_project.addItem(clean_text(project.get("name")) or clean_text(project.get("id")), clean_text(project.get("id")))
        if self.cmb_project.count() == 0:
            self.cmb_project.setPlaceholderText("Сначала выберите компанию")
        if self.cmb_project.lineEdit():
            self.cmb_project.lineEdit().setPlaceholderText("Сначала выберите компанию")
        self.cmb_project.blockSignals(False)
        self._current_project_id()

        self._set_busy(False, "Списки загружены")
        self._set_status_label(f"Списки загружены: компаний {len(companies)}, проектов {len(projects)}", True)

    def _on_worker_failed_generic(self, title: str):
        def handler(error: str):
            self._set_busy(False, "Ошибка")
            self._show_error(title, error)
        return handler

    def _current_sheet_name(self) -> str:
        return clean_text(self.cmb_sheet.currentText())

    def _save_excel_template(self, template_kind: str):
        templates = {
            "permissions": ("SGNL_Шаблон_Ролевая_матрица.xlsx", "ролевой матрицы"),
            "folders": ("SGNL_Шаблон_Папочная_структура.xlsx", "папочной структуры"),
            "attributes": ("SGNL_Шаблон_Атрибуты.xlsx", "атрибутов"),
        }
        default_name, label = templates.get(template_kind, ("SGNL_Шаблон.xlsx", "Excel"))
        documents = QtCore.QStandardPaths.writableLocation(QtCore.QStandardPaths.DocumentsLocation)
        default_path = os.path.join(documents or os.getcwd(), default_name)
        path, _ = QtWidgets.QFileDialog.getSaveFileName(
            self,
            f"Сохранить шаблон {label}",
            default_path,
            "Excel (*.xlsx)",
        )
        if not path:
            return
        if not path.lower().endswith(".xlsx"):
            path += ".xlsx"
        try:
            build_excel_template(template_kind, path)
        except Exception:
            self._show_error("Ошибка создания Excel-шаблона", traceback.format_exc())
            return
        self.statusBar().showMessage(f"Шаблон сохранён: {path}")
        self._log(f"Excel-шаблон {label} сохранён: {path}\n")
        QtWidgets.QMessageBox.information(
            self,
            "Шаблон сохранён",
            f"Шаблон {label} сохранён:\n{path}",
        )

    def _pick_excel(self):
        path, _ = QtWidgets.QFileDialog.getOpenFileName(self, "Выберите Excel-файл", "", "Excel (*.xlsx *.xlsm)")
        if path:
            self.ed_excel.setText(path)
            self._load_excel_sheets(auto_load_structure=True)

    def _load_excel_sheets(self, *_args, auto_load_structure: bool = False):
        excel_path = clean_text(self.ed_excel.text())
        if not excel_path or not os.path.exists(excel_path):
            QtWidgets.QMessageBox.warning(self, "Excel не найден", "Выбери существующий Excel-файл.")
            return
        previous = self._current_sheet_name()
        try:
            names = PermissionExcelParser.list_sheet_names(excel_path)
        except Exception:
            self._show_error("Ошибка чтения Excel", traceback.format_exc())
            return
        self.cmb_sheet.blockSignals(True)
        self.cmb_sheet.clear()
        for name in names:
            self.cmb_sheet.addItem(name)
        target = previous if previous in names else ("2. Ролевая матрица" if "2. Ролевая матрица" in names else (names[0] if names else ""))
        if target:
            self.cmb_sheet.setCurrentText(target)
        self.cmb_sheet.setEnabled(bool(names))
        self.cmb_sheet.blockSignals(False)
        self._invalidate_plan()
        self._log(f"Excel-листы загружены: {', '.join(names)}\n")
        if names and auto_load_structure:
            QtCore.QTimer.singleShot(0, self._load_folder_structure_from_excel)

    def _on_excel_sheet_changed(self, *_args):
        self._invalidate_plan()
        if self._current_sheet_name() and os.path.exists(clean_text(self.ed_excel.text())):
            QtCore.QTimer.singleShot(0, self._load_folder_structure_from_excel)

    def _validate_excel_fields(self) -> bool:
        excel_path = clean_text(self.ed_excel.text())
        if not excel_path or not os.path.exists(excel_path):
            QtWidgets.QMessageBox.warning(self, "Excel не найден", "Выбери существующий Excel-файл.")
            return False
        if not self._current_sheet_name():
            self._load_excel_sheets()
            if not self._current_sheet_name():
                QtWidgets.QMessageBox.warning(self, "Не выбран лист", "Выбери лист Excel с папочной структурой.")
                return False
        return True

    def _validate_connection_fields(self) -> bool:
        if not clean_text(self.ed_docs_url.text()) or not clean_text(self.ed_hub_url.text()):
            QtWidgets.QMessageBox.warning(self, "Не хватает данных", "Заполни Docs URL и Hub URL.")
            return False
        if not self._current_company_id() or not self._current_project_id():
            QtWidgets.QMessageBox.warning(
                self,
                "Не выбран проект",
                "Выбери компанию и проект из списка. Если список пустой — нажми «Войти» или «Загрузить список».",
            )
            return False
        if not clean_text(self.ed_token.text()):
            QtWidgets.QMessageBox.warning(self, "Нет токена", "Сначала выполни вход в программе.")
            return False
        return True

    def _validate_common(self) -> bool:
        if not self._validate_connection_fields():
            return False
        return self._validate_excel_fields()

    def _validate_folder_common(self) -> bool:
        if not clean_text(self.ed_docs_url.text()) or not clean_text(self.ed_hub_url.text()):
            QtWidgets.QMessageBox.warning(self, "Не хватает данных", "Заполни Docs URL и Hub URL.")
            return False
        if not self._current_project_id():
            QtWidgets.QMessageBox.warning(self, "Не выбран проект", "Выбери проект из списка.")
            return False
        if not clean_text(self.ed_token.text()):
            QtWidgets.QMessageBox.warning(self, "Нет токена", "Сначала выполни вход в программе.")
            return False
        return self._validate_excel_fields()

    def _check_access(self):
        if not self._validate_connection_fields():
            return
        self._ensure_log_visible()
        self._set_busy(True, "Проверка доступа...")
        worker = CheckAccessWorker(
            docs_url=clean_text(self.ed_docs_url.text()),
            hub_url=clean_text(self.ed_hub_url.text()),
            token=clean_text(self.ed_token.text()),
            project_id=self._current_project_id(),
            company_id=self._current_company_id(),
            cookies=self._session_cookies,
        )
        worker.log.connect(self._log)
        worker.done.connect(self._on_check_done)
        worker.failed.connect(self._on_check_failed)
        self._run_worker(worker)

    def _on_check_done(self, payload: dict):
        self._set_busy(False, "Доступ проверен")
        self._set_status_label(f"API доступен: ролей {payload['roles']}, папок {payload['folders']}", True)

    def _on_check_failed(self, error: str):
        self._set_busy(False, "Ошибка")
        self._set_status_label("Ошибка доступа", False)
        self._show_error("Ошибка проверки доступа", error)

    def _make_preview_worker(self) -> PreviewWorker:
        return PreviewWorker(
            docs_url=clean_text(self.ed_docs_url.text()),
            hub_url=clean_text(self.ed_hub_url.text()),
            token=clean_text(self.ed_token.text()),
            project_id=self._current_project_id(),
            company_id=self._current_company_id(),
            excel_path=clean_text(self.ed_excel.text()),
            sheet_name=self._current_sheet_name(),
            cookies=self._session_cookies,
        )

    def _start_preview(self, apply_after_preview: bool = False):
        if not self._validate_common():
            return
        self._apply_after_preview = apply_after_preview
        self._invalidate_plan()
        self._ensure_log_visible()
        self._set_busy(True, "Проверка перед выдачей прав..." if apply_after_preview else "Построение предпросмотра...")
        worker = self._make_preview_worker()
        worker.log.connect(self._log)
        worker.done.connect(self._on_preview_done)
        worker.failed.connect(self._on_preview_failed)
        self._run_worker(worker)

    def _preview(self):
        self._start_preview(apply_after_preview=False)

    def _on_preview_done(self, payload: dict):
        self.current_plan = payload["plan"]
        self._fill_plan_table(self.current_plan)
        apply_count = sum(1 for row in self.current_plan if row.can_apply)
        unchanged_count = sum(1 for row in self.current_plan if row.no_changes)
        error_count = sum(1 for row in self.current_plan if not row.is_valid)
        pending_entries = sum(row.pending_permission_count for row in self.current_plan if row.can_apply)
        self._set_busy(False, "Предпросмотр готов")

        self._log(
            f"Результат проверки: к изменению папок {apply_count}; "
            f"без изменений {unchanged_count}; с ошибками {error_count}; "
            f"записей прав к отправке {pending_entries}.\n"
        )

        if not self._apply_after_preview:
            return

        self._apply_after_preview = False
        if error_count:
            if apply_count <= 0:
                QtWidgets.QMessageBox.warning(
                    self,
                    "Права не применены",
                    f"Автоматическая проверка нашла ошибок: {error_count}. "
                    f"Папок с отличающимися правами нет; без изменений: {unchanged_count}.",
                )
                return
            if self._confirm_partial_apply(ok_count=apply_count, error_count=error_count):
                self._log("Пользователь подтвердил синхронизацию только корректных строк.\n")
                self._start_apply_worker()
            else:
                self._log("Применение отменено пользователем.\n")
            return

        if apply_count <= 0:
            QtWidgets.QMessageBox.information(
                self,
                "Изменения не требуются",
                f"Все права SGNL полностью совпадают с Excel. Проверено папок: {unchanged_count}.",
            )
            return

        self._log("Проверка успешна. Начинаю точную синхронизацию отличающихся прав.\n")
        self._start_apply_worker()

    def _on_preview_failed(self, error: str):
        self.current_plan = []
        self._apply_after_preview = False
        self._fill_plan_table([])
        self._set_busy(False, "Ошибка")
        self._show_error("Ошибка предпросмотра", error)

    def _fill_plan_table(self, plan: Sequence[PlanRow]):
        self.tbl_plan.setRowCount(len(plan))
        for row_idx, row in enumerate(plan):
            folder_path = row.folder.path_text if row.folder else ""
            permission_parts: List[str] = []
            pending = set(row.pending_roles)
            unchanged = set(row.unchanged_roles)
            for role, perm in row.entry.permissions.items():
                if role in unchanged:
                    marker = "пропуск — уже соответствует Excel"
                elif role in pending and perm.is_remove:
                    marker = "снять права (DELETE)"
                elif role in pending:
                    marker = "создать/перезаписать (UPSERT)"
                else:
                    marker = "не проверено"
                permission_parts.append(f"{role}: {perm.text} [{marker}]")
            perms_text = "; ".join(permission_parts)
            values = [
                row.status,
                str(row.entry.row_number),
                row.entry.path_text,
                folder_path,
                f"{row.pending_permission_count} / {row.permission_count}",
                perms_text,
                row.message,
            ]
            for col_idx, value in enumerate(values):
                item = QtWidgets.QTableWidgetItem(value)
                item.setToolTip(value)
                if row.no_changes:
                    item.setBackground(QtGui.QColor("#173227" if self.is_dark_theme else "#EDF8F1"))
                elif row.can_apply:
                    item.setBackground(QtGui.QColor("#3A2B1A" if self.is_dark_theme else "#FFF3E4"))
                else:
                    item.setBackground(QtGui.QColor("#3A2020" if self.is_dark_theme else "#FDECEC"))
                self.tbl_plan.setItem(row_idx, col_idx, item)
        self.tbl_plan.resizeRowsToContents()
        fit_preview_height(self.tbl_plan)
        if hasattr(self, "btn_permissions_details"):
            self.btn_permissions_details.setEnabled(bool(plan))

    def _confirm_apply(self, checked_rows: Optional[int] = None, pending_entries: Optional[int] = None) -> bool:
        if checked_rows is None:
            rows_text = "Количество изменений будет определено после автоматической проверки."
            action_text = (
                "Программа сначала сравнит Excel с актуальными правами SGNL. "
                "Совпадающие записи будут пропущены."
            )
        else:
            rows_text = f"Папок с отсутствующими или отличающимися правами: {checked_rows}."
            if pending_entries is not None:
                rows_text += f" Записей прав к отправке: {pending_entries}."
            action_text = (
                "Будет использован предпросмотр, а непосредственно перед отправкой программа "
                "повторно проверит актуальные права SGNL."
            )
        confirm_text = (
            f"{rows_text}\n\n"
            f"{action_text}\n\n"
            "П/С/З/Р устанавливаются точными значениями из Excel, включая понижение уровня.\n"
            "Пустая ячейка не меняет право. Символ '-' удаляет существующую запись права через DELETE API.\n\n"
            "Синхронизировать права с Excel?"
        )
        answer = QtWidgets.QMessageBox.question(
            self,
            "Подтвердить применение прав",
            confirm_text,
            QtWidgets.QMessageBox.Yes | QtWidgets.QMessageBox.No,
            QtWidgets.QMessageBox.No,
        )
        return answer == QtWidgets.QMessageBox.Yes

    def _confirm_partial_apply(self, ok_count: int, error_count: int) -> bool:
        box = QtWidgets.QMessageBox(self)
        box.setIcon(QtWidgets.QMessageBox.Warning)
        box.setWindowTitle("Есть ошибки")
        box.setText("В предпросмотре есть ошибки.")
        box.setInformativeText(
            f"Строк с отличающимися правами: {ok_count}.\n"
            f"Строк с ошибками: {error_count}.\n\n"
            "Можно синхронизировать права только для корректных строк. "
            "Ошибочные строки будут пропущены и останутся без изменений.\n\n"
            "Пустая ячейка не меняет право. Символ '-' снимает право через DELETE API."
        )
        apply_button = box.addButton("Да, синхронизировать права", QtWidgets.QMessageBox.AcceptRole)
        cancel_button = box.addButton("Отмена", QtWidgets.QMessageBox.RejectRole)
        box.setDefaultButton(cancel_button)
        if hasattr(box, "exec"):
            box.exec()
        else:  # PyQt5 fallback
            box.exec_()
        return box.clickedButton() == apply_button

    def _start_apply_worker(self):
        self._ensure_log_visible()
        self._set_busy(True, "Выдача прав...")
        worker = ApplyWorker(
            docs_url=clean_text(self.ed_docs_url.text()),
            hub_url=clean_text(self.ed_hub_url.text()),
            token=clean_text(self.ed_token.text()),
            project_id=self._current_project_id(),
            plan=self.current_plan,
            cookies=self._session_cookies,
        )
        worker.log.connect(self._log)
        worker.done.connect(self._on_apply_done)
        worker.failed.connect(self._on_apply_failed)
        self._run_worker(worker)

    def _apply_permissions(self):
        if not self._validate_common():
            return
        if self.current_plan:
            apply_count = sum(1 for row in self.current_plan if row.can_apply)
            unchanged_count = sum(1 for row in self.current_plan if row.no_changes)
            error_count = sum(1 for row in self.current_plan if not row.is_valid)
            pending_entries = sum(row.pending_permission_count for row in self.current_plan if row.can_apply)

            if error_count:
                if apply_count <= 0:
                    QtWidgets.QMessageBox.warning(
                        self,
                        "Нет прав к применению",
                        f"Отличающихся прав не найдено. Без изменений: {unchanged_count}; "
                        f"строк с ошибками: {error_count}. Исправь ошибки и проверь снова.",
                    )
                    return
                if not self._confirm_partial_apply(ok_count=apply_count, error_count=error_count):
                    return
                self._log(
                    f"Частичное применение: папок к изменению {apply_count}; "
                    f"записей прав {pending_entries}; ошибочных строк пропущено {error_count}.\n"
                )
                self._start_apply_worker()
                return

            if apply_count <= 0:
                QtWidgets.QMessageBox.information(
                    self,
                    "Изменения не требуются",
                    f"Все права SGNL полностью совпадают с Excel. Проверено папок: {unchanged_count}.",
                )
                return

            if not self._confirm_apply(checked_rows=apply_count, pending_entries=pending_entries):
                return
            self._start_apply_worker()
            return

        if not self._confirm_apply(checked_rows=None):
            return
        self._start_preview(apply_after_preview=True)

    def _load_folder_structure_from_excel(self):
        if not self._validate_excel_fields():
            return
        self.current_folder_plan = []
        self.folder_plan_verified = False
        self.tbl_folders.setRowCount(0)
        if hasattr(self, "btn_folder_details"):
            self.btn_folder_details.setEnabled(False)
        self._set_busy(True, "Загрузка структуры из Excel...")
        worker = ExcelStructureLoadWorker(
            excel_path=clean_text(self.ed_excel.text()),
            sheet_name=self._current_sheet_name(),
        )
        worker.log.connect(self._log)
        worker.done.connect(self._on_excel_structure_loaded)
        worker.failed.connect(self._on_excel_structure_failed)
        self._run_worker(worker)

    def _on_excel_structure_loaded(self, payload: dict):
        self.current_folder_plan = list(payload.get("plan") or [])
        self.folder_plan_verified = False
        self._fill_folder_table(self.current_folder_plan)
        self._set_busy(False, f"Из Excel загружено папок: {payload.get('count', 0)}")

    def _on_excel_structure_failed(self, error: str):
        self.current_folder_plan = []
        self.folder_plan_verified = False
        self.tbl_folders.setRowCount(0)
        self._set_busy(False, "Ошибка")
        self._show_error("Ошибка загрузки структуры из Excel", error)

    def _preview_folder_structure(self):
        if not self._validate_folder_common():
            return
        self.current_folder_plan = []
        self.tbl_folders.setRowCount(0)
        if hasattr(self, "btn_folder_details"):
            self.btn_folder_details.setEnabled(False)
        self._ensure_log_visible()
        self._set_busy(True, "Проверка папочной структуры...")
        worker = FolderStructurePreviewWorker(
            docs_url=clean_text(self.ed_docs_url.text()),
            hub_url=clean_text(self.ed_hub_url.text()),
            token=clean_text(self.ed_token.text()),
            project_id=self._current_project_id(),
            excel_path=clean_text(self.ed_excel.text()),
            sheet_name=self._current_sheet_name(),
            cookies=self._session_cookies,
        )
        worker.log.connect(self._log)
        worker.done.connect(self._on_folder_preview_done)
        worker.failed.connect(self._on_folder_preview_failed)
        self._run_worker(worker)

    def _on_folder_preview_done(self, payload: dict):
        self.current_folder_plan = list(payload.get("plan") or [])
        self.folder_plan_verified = True
        self._fill_folder_table(self.current_folder_plan)
        self._set_busy(False, "Предпросмотр структуры готов")
        if payload.get("errors"):
            QtWidgets.QMessageBox.warning(
                self,
                "Есть ошибки в структуре",
                f"К созданию: {payload.get('to_create', 0)}\n"
                f"Уже существует: {payload.get('existing', 0)}\n"
                f"Ошибок: {payload.get('errors', 0)}\n\n"
                "Ошибочные строки созданы не будут.",
            )

    def _on_folder_preview_failed(self, error: str):
        self.current_folder_plan = []
        self.folder_plan_verified = False
        self.tbl_folders.setRowCount(0)
        self._set_busy(False, "Ошибка")
        self._show_error("Ошибка проверки папочной структуры", error)

    def _fill_folder_table(self, plan: Sequence[FolderCreatePlanRow]):
        self.tbl_folders.setRowCount(len(plan))
        for row_idx, row in enumerate(plan):
            values = [
                row.status,
                str(row.row_number),
                row.path_text,
                row.parent_path,
                row.folder_id,
                row.message,
            ]
            for col_idx, value in enumerate(values):
                item = QtWidgets.QTableWidgetItem(value)
                item.setToolTip(value)
                if row.status == "Существует":
                    item.setBackground(QtGui.QColor("#173227" if self.is_dark_theme else "#EDF8F1"))
                elif row.status == "Создать":
                    item.setBackground(QtGui.QColor("#3A2B1A" if self.is_dark_theme else "#FFF3E4"))
                elif row.status == "Из Excel":
                    item.setBackground(QtGui.QColor("#1F2937" if self.is_dark_theme else "#F3F6FA"))
                else:
                    item.setBackground(QtGui.QColor("#3A2020" if self.is_dark_theme else "#FDECEC"))
                self.tbl_folders.setItem(row_idx, col_idx, item)
        self.tbl_folders.resizeRowsToContents()
        fit_preview_height(self.tbl_folders)
        if hasattr(self, "btn_folder_details"):
            self.btn_folder_details.setEnabled(bool(plan))

    def _create_folder_structure(self):
        if not self._validate_folder_common():
            return
        if not self.current_folder_plan or not self.folder_plan_verified:
            QtWidgets.QMessageBox.information(
                self,
                "Сначала сверка",
                "Сначала нажми «Сверить с SGNL». Загрузка из Excel показывает только требуемую структуру, "
                "но ещё не определяет, какие папки уже существуют в выбранном проекте.",
            )
            return
        create_count = sum(1 for row in self.current_folder_plan if row.status == "Создать")
        error_count = sum(1 for row in self.current_folder_plan if row.status == "Ошибка")
        if create_count == 0:
            QtWidgets.QMessageBox.information(self, "Создавать нечего", "Все папки из Excel уже существуют.")
            return
        answer = QtWidgets.QMessageBox.question(
            self,
            "Подтвердить создание папок",
            f"Будет создано отсутствующих папок: {create_count}.\n"
            f"Ошибочных строк будет пропущено: {error_count}.\n\n"
            "Существующие папки не изменяются и не дублируются. Продолжить?",
            QtWidgets.QMessageBox.Yes | QtWidgets.QMessageBox.No,
            QtWidgets.QMessageBox.No,
        )
        if answer != QtWidgets.QMessageBox.Yes:
            return

        self._ensure_log_visible()
        self._set_busy(True, "Создание папок...")
        worker = CreateFoldersWorker(
            docs_url=clean_text(self.ed_docs_url.text()),
            hub_url=clean_text(self.ed_hub_url.text()),
            token=clean_text(self.ed_token.text()),
            project_id=self._current_project_id(),
            plan=self.current_folder_plan,
            cookies=self._session_cookies,
        )
        worker.log.connect(self._log)
        worker.done.connect(self._on_folder_create_done)
        worker.failed.connect(self._on_folder_create_failed)
        self._run_worker(worker)

    def _on_folder_create_done(self, stats: dict):
        self._set_busy(False, "Папочная структура создана")
        QtWidgets.QMessageBox.information(
            self,
            "Готово",
            f"Создано папок: {stats.get('created', 0)}\n"
            f"Уже существовало: {stats.get('existing', 0)}\n"
            f"Ошибок: {stats.get('failed', 0)}",
        )
        # Дерево изменилось — старый предпросмотр прав больше нельзя считать актуальным.
        self.current_plan = []
        self.folder_plan_verified = False
        self.tbl_plan.setRowCount(0)
        self._preview_folder_structure()

    def _on_folder_create_failed(self, error: str):
        self._set_busy(False, "Ошибка")
        self._show_error("Ошибка создания папочной структуры", error)

    def _on_apply_done(self, stats: dict):
        self._set_busy(False, "Готово")
        QtWidgets.QMessageBox.information(
            self,
            "Готово",
            f"Изменено папок: {stats.get('folders', 0)}\n"
            f"Создано/перезаписано записей: {stats.get('updated_entries', 0)}\n"
            f"Снято записей прав: {stats.get('removed_entries', 0)}\n"
            f"Уже соответствовало Excel: {stats.get('already_actual_entries', 0)}\n"
            f"Папок без изменений: {stats.get('already_actual_rows', 0)}\n"
            f"Ошибок: {stats.get('failed', 0)}\n"
            f"Пропущено строк: {stats.get('skipped', 0)}",
        )
        # После применения снимок прав устарел. Новый запуск обязан снова свериться с SGNL.
        self.current_plan = []
        self.tbl_plan.setRowCount(0)

    def _on_apply_failed(self, error: str):
        self._set_busy(False, "Ошибка")
        self._show_error("Ошибка выдачи прав", error)


def main() -> int:
    app = QtWidgets.QApplication(sys.argv)
    window = MainWindow()
    window.show()
    if hasattr(app, "exec"):
        return app.exec()
    return app.exec_()


if __name__ == "__main__":
    sys.exit(main())
