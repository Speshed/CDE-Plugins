from __future__ import annotations

import re

from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from shared.excel_style import apply_excel_style
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from ..api.custom_fields import *
from ..api.roles import *

__all__ = ["_safe_sheet_title", "_write_simple_xlsx", "write_params_import_template", "write_project_stages_import_template", "write_content_types_import_template", "write_routes_import_template", "write_objects_import_template", "_module_filter_matches", "write_roles_import_template"]


def _safe_sheet_title(title):
    title = re.sub(r"[\\/*?:\[\]]", " ", str(title))[:31].strip()
    return title or "Лист1"


def _write_simple_xlsx(filename, sheets):
    wb = Workbook()
    first = True
    header_fill = PatternFill(start_color="D9EAF7", end_color="D9EAF7", fill_type="solid")
    section_fill = PatternFill(start_color="EAF4EC", end_color="EAF4EC", fill_type="solid")
    required_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    header_font = Font(bold=True)
    title_font = Font(bold=True, size=14)
    wrap_top = Alignment(wrap_text=True, vertical="top")

    for sheet_title, rows, widths, options in sheets:
        ws = wb.active if first else wb.create_sheet(_safe_sheet_title(sheet_title))
        first = False
        ws.title = _safe_sheet_title(sheet_title)
        for row in rows:
            ws.append(row)
        for cell in ws[1]:
            cell.font = title_font
            cell.fill = section_fill
            cell.alignment = wrap_top
        header_row = options.get("header_row", 2)
        for cell in ws[header_row]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        for col_idx in options.get("required_cols", []):
            ws.cell(row=header_row, column=col_idx).fill = required_fill
        for row in ws.iter_rows(min_row=header_row + 1):
            for cell in row:
                cell.alignment = wrap_top
        for col_idx, width in widths.items():
            ws.column_dimensions[get_column_letter(col_idx)].width = width
        if options.get("freeze"):
            ws.freeze_panes = options["freeze"]
        if options.get("filter"):
            ws.auto_filter.ref = options["filter"]
    apply_excel_style(wb)
    wb.save(filename)
    return filename


def write_params_import_template(filename):
    rows = [
        ["Шаблон создания параметров Project Point"],
        ["Наименование атрибута", "Внутреннее имя атрибута", "Подсказка", "Тип данных", "Значения списка"],
        ["Номер договора", "ContractNumber", "Номер договора по документу", "Короткий текст", ""],
        ["Статус проверки", "CheckStatus", "Справочное значение", "Список", "На проверке; Принято; Отклонено"],
    ]
    _write_simple_xlsx(
        filename,
        [("Параметры", rows, {1: 34, 2: 30, 3: 46, 4: 22, 5: 48}, {"header_row": 2, "required_cols": [1, 2, 4], "freeze": "A3", "filter": "A2:E500"})]
    )
    return {"filename": filename}


def write_project_stages_import_template(filename):
    rows = [
        ["Шаблон создания / обновления видов документов Project Point"],
        ["Code", "Title", "DisplayParameters", "IsActive"],
        ["РД", "Рабочая документация", 1, True],
        ["ПД", "Проектная документация", 1, True],
    ]
    _write_simple_xlsx(
        filename,
        [(
            "Виды документов",
            rows,
            {1: 18, 2: 48, 3: 24, 4: 18},
            {"header_row": 2, "required_cols": [1, 2], "freeze": "A3", "filter": "A2:D500"},
        )],
    )
    return {"filename": filename}


def write_content_types_import_template(filename):
    headers = ["Code", "Title", "Description", "ProjectStages", "DisplayParameters", "CustomFields"] + BOOL_FIELDS
    sample = ["DOC-GEN", "Общий документ", "Описание типа", "Рабочая документация; Проектная документация", 0, "Номер договора; Статус проверки"] + [False] * len(BOOL_FIELDS)
    rows = [["Шаблон создания / обновления типов документов Project Point"], headers, sample]
    _write_simple_xlsx(
        filename,
        [("Типы документов", rows, {1: 18, 2: 34, 3: 42, 4: 24, 5: 20, 6: 42, **{i: 18 for i in range(7, 7 + len(BOOL_FIELDS))}}, {"header_row": 2, "required_cols": [1, 2], "freeze": "A3", "filter": f"A2:{get_column_letter(len(headers))}500"})]
    )
    return {"filename": filename}


def write_routes_import_template(filename):
    wb = Workbook()
    ws = wb.active
    ws.title = "Маршруты"
    rows = [
        ["Шаблон маршрутов согласования"],
        [],
        [],
        ["Маршрут", "Проект", "Администратор", "Тип согласования", "Этап 1", "Этап 1", "Этап 1", "Этап 1", "Этап 1", "Этап 1", "Этап 1", "Этап 1", "Этап 1", "Этап 1", "Этап 2", "Этап 2", "Этап 2", "Этап 3", "Этап 3", "Этап 4", "Этап 4", "Этап 5", "Этап 5", "Этап 6", "Этап 6"],
        ["Маршрут 1", "Все", "Иванов Иван Иванович", "Согласование", "Петров Петр Петрович", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", ""],
    ]
    for row in rows:
        ws.append(row)
    ws_cond = wb.create_sheet("Условия маршрута")
    cond_rows = [
        ["Шаблон условий маршрута"], [], [], [],
        ["Маршрут", "Вид документа", "Тип документа", "Проект", "Объект строительства", "Заказчик", "Разработчик", "Дисциплина", "Категория", "Цель выпуска"],
        ["Маршрут 1", "Все", "Все", "Все", "Все", "Все", "Все", "Все", "Все", "Все"],
    ]
    for row in cond_rows:
        ws_cond.append(row)
    ws_dur = wb.create_sheet("Продолжительность")
    dur_rows = [
        ["Шаблон продолжительности этапов"],
        ["Данные начинаются с 6-й строки. Не смещайте строку с данными выше/ниже без изменения кода."],
        ["Колонки 1-7 — длительности. Колонки 11-35 — флаги этапов."],
        [],
        ["Маршрут", "Этап 1, дней", "Этап 2, дней", "Этап 3, дней", "Этап 4, дней", "Этап 5, дней", "Этап 6, дней", "Утверждающий, дней", "", "", "", "Этап 1 AutoComplete", "Этап 1 Single", "Этап 1 Reject", "Этап 1 ЭП", "Этап 2 AutoComplete", "Этап 2 Single", "Этап 2 Reject", "Этап 2 ЭП", "Этап 3 AutoComplete", "Этап 3 Single", "Этап 3 Reject", "Этап 3 ЭП", "Этап 4 AutoComplete", "Этап 4 Single", "Этап 4 Reject", "Этап 4 ЭП", "Этап 5 AutoComplete", "Этап 5 Single", "Этап 5 Reject", "Этап 5 ЭП", "Этап 6 AutoComplete", "Этап 6 Single", "Этап 6 Reject", "Этап 6 ЭП", "Утверждающий ЭП"],
        ["Маршрут 1", 2, 2, 2, 2, 2, 2, 4, "", "", "", "Да", "Нет", "Нет", "Нет", "Да", "Нет", "Нет", "Нет", "Да", "Нет", "Нет", "Нет", "Да", "Нет", "Нет", "Нет", "Да", "Нет", "Нет", "Нет", "Да", "Нет", "Нет", "Нет", "Нет"],
    ]
    for row in dur_rows:
        ws_dur.append(row)
    header_fill = PatternFill(start_color="D9EAF7", end_color="D9EAF7", fill_type="solid")
    title_fill = PatternFill(start_color="EAF4EC", end_color="EAF4EC", fill_type="solid")
    for sheet, header_row, freeze in [(ws, 4, "A5"), (ws_cond, 5, "A6"), (ws_dur, 5, "A6")]:
        sheet.cell(row=1, column=1).font = Font(bold=True, size=14)
        sheet.cell(row=1, column=1).fill = title_fill
        for cell in sheet[header_row]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        sheet.freeze_panes = freeze
        for col_idx in range(1, min(sheet.max_column, 36) + 1):
            sheet.column_dimensions[get_column_letter(col_idx)].width = 22
    apply_excel_style(wb)
    wb.save(filename)
    return {"filename": filename}



def write_objects_import_template(filename):
    rows = [
        ["Шаблон создания объектов строительства Project Point"],
        ["Code", "Title", "ParentCode", "ParentTitle"],
        ["OBJ-001", "Корпус 1", "EXISTING-PARENT", ""],
        ["OBJ-001-01", "Секция 1", "OBJ-001", "Корпус 1"],
    ]
    _write_simple_xlsx(
        filename,
        [(
            "Объекты строительства",
            rows,
            {1: 22, 2: 42, 3: 24, 4: 42},
            {"header_row": 2, "required_cols": [1, 2], "freeze": "A3", "filter": "A2:D500"},
        )],
    )
    return {"filename": filename}

def _module_filter_matches(module_name, selected_modules=None):
    if not selected_modules:
        return True
    normalized_selected = {
        _normalize_module_filter(module)
        for module in selected_modules
        if _normalize_lookup_text(module)
    }
    if not normalized_selected:
        return True
    return _normalize_module_filter(module_name) in normalized_selected


def write_roles_import_template(all_permissions, filename, active_modules=None, selected_modules=None):
    """Create multi-sheet role import Excel template.

    Workbook layout:
    - one worksheet per visible Project Point module;
    - columns: Роль | Привилегия | Свои элементы | Моей орг. Единицы | Другие орг. Единицы;
    - allowed checkboxes are prefilled with ✗ and can be changed to ✓;
    - unavailable checkbox cells are gray and empty; if the user fills them, import stops with an error.
    """
    module_title_map = build_permission_module_title_map(active_modules)
    visible_permissions = [
        p for p in all_permissions
        if p.get("Id") is not None
        and is_role_permission_visible(p)
        and is_permission_module_visible(p, active_modules)
    ]
    visible_permissions.sort(
        key=lambda p: (
            _role_ui_module_sort_key(permission_module_display(p, module_title_map)),
            _role_ui_text_sort_key(p.get("Name")),
        )
    )

    grouped = {}
    for perm in visible_permissions:
        module_name = permission_module_display(perm, module_title_map)
        if not _module_filter_matches(module_name, selected_modules):
            continue
        grouped.setdefault(module_name, []).append(perm)

    ordered_modules = sorted(grouped, key=_role_ui_module_sort_key)

    wb = Workbook()
    default_sheet = wb.active
    wb.remove(default_sheet)

    header_fill = PatternFill(start_color="D9EAF7", end_color="D9EAF7", fill_type="solid")
    blocked_fill = PatternFill(start_color="BFBFBF", end_color="BFBFBF", fill_type="solid")
    alt_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    input_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    header_font = Font(bold=True)
    wrap_top = Alignment(wrap_text=True, vertical="top")
    center = Alignment(horizontal="center", vertical="center")
    headers = ["Роль", "Привилегия", "Свои элементы", "Моей орг. Единицы", "Другие орг. Единицы"]
    sample_role = "Новая роль"
    check_validation = DataValidation(type="list", formula1='"✓,✗,Да,Нет"', allow_blank=True)

    permissions_count = 0
    sheet_names = []
    used_sheet_titles = set()

    for module_name in ordered_modules:
        sheet_title = _safe_sheet_title(module_name)
        base_title = sheet_title
        counter = 2
        while sheet_title in used_sheet_titles:
            suffix = f" {counter}"
            sheet_title = _safe_sheet_title(base_title[:31 - len(suffix)] + suffix)
            counter += 1
        used_sheet_titles.add(sheet_title)
        sheet_names.append(sheet_title)

        ws = wb.create_sheet(sheet_title)
        ws.append(headers)
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.freeze_panes = "A2"

        local_dv = DataValidation(type="list", formula1='"✓,✗,Да,Нет"', allow_blank=True)
        ws.add_data_validation(local_dv)

        for perm in grouped[module_name]:
            row_idx = ws.max_row + 1
            permissions_count += 1
            ws.append([sample_role, _normalize_lookup_text(perm.get("Name")), "", "", ""])

            row_fill = alt_fill if row_idx % 2 == 0 else input_fill
            for col_idx in range(1, 6):
                ws.cell(row=row_idx, column=col_idx).fill = row_fill

            allowed = get_permission_allowed_flags(perm.get("Id"))
            for flag, col_idx in [("My", 3), ("MyOrganization", 4), ("OtherOrganization", 5)]:
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.alignment = center
                if allowed.get(flag, True):
                    cell.value = "✗"
                    cell.fill = row_fill
                    local_dv.add(cell)
                else:
                    cell.value = ""
                    cell.fill = blocked_fill

        ws.auto_filter.ref = f"A1:E{max(ws.max_row, 1)}"
        widths = {"A": 34, "B": 64, "C": 16, "D": 20, "E": 22}
        for col, width in widths.items():
            ws.column_dimensions[col].width = width
        ws.row_dimensions[1].height = 30
        for row_idx in range(2, ws.max_row + 1):
            ws.row_dimensions[row_idx].height = 34
            ws.cell(row=row_idx, column=1).alignment = wrap_top
            ws.cell(row=row_idx, column=2).alignment = wrap_top

    if not ordered_modules:
        ws = wb.create_sheet("Нет прав")
        ws.append(["Не найдено прав для выбранного модуля"])
        ws["A1"].font = Font(bold=True)
        ws.column_dimensions["A"].width = 80

    info = wb.create_sheet("Инструкция")
    info_rows = [
        ["Шаблон импорта ролей Project Point"],
        ["Как заполнять"],
        ["1. Каждый лист соответствует одному модулю Project Point."],
        ["2. В колонке 'Роль' укажите название создаваемой роли. Для одной роли оставьте одно и то же название на всех выбранных листах."],
        ["3. Для создания нескольких ролей можно скопировать строки и поменять значение в колонке 'Роль'."],
        ["4. В белых ячейках доступа используйте '✓'/'✗' или 'Да'/'Нет'. Пусто считается как 'Нет'."],
        ["5. Серые ячейки соответствуют недоступным чекбоксам интерфейса. Их нельзя заполнять: при импорте будет ошибка."],
        ["6. Перед импортом в программе выберите галочками только те листы-модули, которые нужно загрузить."],
        ["7. Роли создаются активными. Если нужно импортировать неактивную роль, добавьте колонку 'Активна' на лист модуля и укажите 'Нет'."],
    ]
    for row in info_rows:
        info.append(row)
    info["A1"].font = Font(bold=True, size=14)
    info["A2"].font = Font(bold=True)
    info.column_dimensions["A"].width = 125
    for row in info.iter_rows(min_row=1, max_row=info.max_row, min_col=1, max_col=1):
        for cell in row:
            cell.alignment = wrap_top
    info.sheet_state = "hidden"

    apply_excel_style(wb)
    wb.save(filename)
    return {
        "filename": filename,
        "permissions_count": permissions_count,
        "modules_count": len(ordered_modules),
        "sheets": sheet_names,
    }
