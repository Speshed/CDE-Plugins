from __future__ import annotations

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from shared.excel_style import apply_excel_style


def write_vitrocad_template(kind: str, filename) -> None:
    """Create one of the three VitroCAD import books with parser headers."""
    wb = Workbook()
    ws = wb.active
    if kind == "folders":
        ws.title = "Папочная структура"
        ws.append(["Уровень 1", "Уровень 2", "Уровень 3", "Тип папки"])
        ws.append(["01. Проектирование", "АР", "Рабочая документация", "Папка"])
    elif kind == "permissions":
        ws.title = "Матрица прав"
        ws.append(["Уровень 1", "Уровень 2", "Роль 1", "Роль 2"])
        ws.append(["01. Проектирование", "АР", "Просмотр", ""])
    elif kind == "schedule":
        ws.title = "План-график"
        ws.append([
            "Название", "Путь", "Тип задачи", "Исполнитель", "Статус задачи",
            "Дата начала (План)", "Дата окончания (План)", "Важность", "Описание",
            "Комментарий", "Продолжительность", "Плановые трудозатраты",
            "Максимальные трудозатраты", "Предшественники", "Потомки",
        ])
        ws.append(["Этап 1", "01. Проектирование", "Задача", "", "Не начата"] + [""] * 10)
    else:
        raise ValueError(f"Неизвестный тип шаблона VitroCAD: {kind}")
    for column in range(1, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(column)].width = 24
    apply_excel_style(wb)
    output = str(filename)
    if not output.lower().endswith(".xlsx"):
        output += ".xlsx"
    wb.save(output)
