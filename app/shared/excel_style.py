"""Common visual language for generated CDE Excel templates.

The palette intentionally mirrors the restrained pastel look of the approved
approval-route workbook: light blue / peach / green / yellow / gray bands,
black text, thin neutral borders and white data cells.
"""
from __future__ import annotations

from copy import copy
import re

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

# Backward-compatible names used by Larix.
# Palette sampled from the approved approval-route workbook supplied by the user.
# The goal is a quiet Office-style look rather than vivid product accent colors.
ORANGE = "D9D9D9"
DARK = "F2F2F2"
LIGHT = "F5F5F5"
BORDER = "8F8F8F"
MUTED = "686868"
TEXT = "3F3F3F"

TITLE_FILL = "F2F2F2"
HEADER_FILL = "D9D9D9"
SUBHEADER_FILL = "F2F2F2"
NOTE_FILL = "F5F5F5"
REQUIRED_FILL = "FFF3CD"
BLOCKED_FILL = "BFBFBF"
ALT_FILL = "F5F5F5"

# Exact pastel stage sequence visible in the supplied reference workbook.
STAGE_FILLS = (
    "D9E2F3",  # stage 1 — blue
    "FBE5D7",  # stage 2 — peach
    "E2EFD9",  # stage 3 — green
    "DEEBF6",  # stage 4 — light blue
    "FFF3CD",  # stage 5 — yellow
    "D6DCE4",  # stage 6 — gray-blue
    "E2EFD9",  # stage 7 — green
)

# Legacy fills used by the generators before the common style pass.
# They are normalized so no template keeps the old saturated accent colors.
_FILL_REPLACEMENTS = {
    # Old app accent palette.
    "F7921E": HEADER_FILL,
    "2F2F2F": TITLE_FILL,
    "FFF4E8": NOTE_FILL,
    # Previous muted generator palette.
    "E7E6E6": TITLE_FILL,
    "F7F7F7": NOTE_FILL,
    "FFF2CC": REQUIRED_FILL,
    # Legacy Larix embedded-template blues.
    "17365D": TITLE_FILL,
    "2F75B5": HEADER_FILL,
    "D9EAF7": HEADER_FILL,
    "EEF5FB": NOTE_FILL,
    "EAF4EC": NOTE_FILL,
}


def _rgb_value(cell) -> str:
    try:
        value = cell.fill.fgColor.rgb or ""
    except Exception:
        return ""
    value = str(value).upper()
    if len(value) == 8:
        value = value[-6:]
    return value


def _legacy_theme_fill(cell) -> str | None:
    """Map old non-neutral Office theme fills into the common muted palette."""
    try:
        color = cell.fill.fgColor
        if color.type != "theme" or color.theme in (None, 0, 1):
            return None
        tint = float(color.tint or 0.0)
    except Exception:
        return None
    # Dark or saturated theme fills become the neutral title band; progressively
    # lighter theme tints become header/note bands. This mainly normalizes the
    # older Larix role-matrix workbook without touching black/white theme grays.
    if tint < 0.20:
        return TITLE_FILL
    if tint < 0.65:
        return HEADER_FILL
    return NOTE_FILL


def _solid_fill(color: str) -> PatternFill:
    return PatternFill("solid", fgColor=color)


def _style_cell(cell, *, fill: str | None = None, bold: bool | None = None,
                font_size: int | float | None = None, color: str | None = None,
                horizontal: str | None = None, vertical: str = "center") -> None:
    if cell.__class__.__name__ == "MergedCell":
        return
    if fill:
        cell.fill = _solid_fill(fill)
    font = copy(cell.font)
    font.name = "Segoe UI"
    font.sz = font_size or 10
    if bold is not None:
        font.bold = bool(bold)
    if color:
        font.color = color
    cell.font = font
    alignment = copy(cell.alignment)
    alignment.vertical = vertical
    alignment.wrap_text = True
    if horizontal:
        alignment.horizontal = horizontal
    cell.alignment = alignment


def _nonempty_count(ws, row_index: int) -> int:
    return sum(1 for cell in ws[row_index] if cell.value not in (None, ""))


def _row_values(ws, row_index: int) -> list[str]:
    return [str(cell.value).strip() for cell in ws[row_index] if cell.value not in (None, "")]


def _header_row(ws) -> int | None:
    """Find the first real tabular header near the top of the sheet."""
    max_probe = min(ws.max_row, 8)
    first_count = _nonempty_count(ws, 1) if ws.max_row else 0
    if first_count >= 2:
        return 1
    for row_idx in range(2, max_probe + 1):
        if _nonempty_count(ws, row_idx) >= 2:
            return row_idx
    return None


def _has_stage_groups_in_first_row(ws) -> bool:
    if ws.max_row < 1:
        return False
    return any(re.fullmatch(r"\d+\s*Этап", value, flags=re.IGNORECASE) for value in _row_values(ws, 1))


def _top_text(ws, max_rows: int = 6) -> set[str]:
    values: set[str] = set()
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, max_rows)):
        for cell in row:
            if cell.value not in (None, ""):
                values.add(str(cell.value).strip())
    return values


def _is_approval_access_duration_sheet(ws) -> bool:
    values = _top_text(ws, 4)
    return "Доступ" in values and any("Продолжительность" in value for value in values)


def _is_approval_flags_sheet(ws) -> bool:
    values = _top_text(ws, 5)
    return (
        "Отмена процесса согласования" in values
        or "Обязательность проверки" in values
    )


def _style_approval_main_sheet(ws, border: Border) -> None:
    """Main route sheet: colored stage groups, otherwise white cells like the reference."""
    if ws.max_column >= 1 and ws.cell(1, 1).value not in (None, ""):
        _style_cell(ws.cell(1, 1), bold=True, color=TEXT, horizontal="left")

    for cell in ws[1]:
        if cell.__class__.__name__ == "MergedCell" or cell.value in (None, ""):
            continue
        match = re.fullmatch(r"(\d+)\s*Этап", str(cell.value).strip(), flags=re.IGNORECASE)
        if match:
            stage = max(1, int(match.group(1)))
            _style_cell(
                cell,
                fill=STAGE_FILLS[(stage - 1) % len(STAGE_FILLS)],
                bold=True,
                color=TEXT,
                horizontal="center",
            )

    if ws.max_row >= 2:
        for cell in ws[2]:
            if cell.__class__.__name__ != "MergedCell" and cell.value not in (None, ""):
                _style_cell(cell, bold=True, color=TEXT, horizontal="center")

    for row in ws.iter_rows():
        for cell in row:
            if cell.__class__.__name__ != "MergedCell" and cell.value is not None:
                cell.border = border


def _style_approval_flags_sheet(ws, border: Border) -> None:
    """5.2-like sheet: all stage bands use the same soft blue from the reference."""
    if ws.max_column >= 1 and ws.cell(1, 1).value not in (None, ""):
        _style_cell(ws.cell(1, 1), bold=True, color=TEXT, horizontal="left")
    for cell in ws[1]:
        if cell.__class__.__name__ == "MergedCell" or cell.value in (None, ""):
            continue
        if re.fullmatch(r"\d+\s*Этап", str(cell.value).strip(), flags=re.IGNORECASE):
            _style_cell(cell, fill=STAGE_FILLS[0], bold=True, color=TEXT, horizontal="center")
    if ws.max_row >= 2:
        for cell in ws[2]:
            if cell.__class__.__name__ != "MergedCell" and cell.value not in (None, ""):
                _style_cell(cell, bold=True, color=TEXT, horizontal="center")
    # Detailed flags in row 4 intentionally stay white in the approved workbook.
    if ws.max_row >= 4:
        for cell in ws[4]:
            if cell.__class__.__name__ != "MergedCell" and cell.value not in (None, ""):
                _style_cell(cell, bold=False, color=TEXT, horizontal="left")
    for row in ws.iter_rows():
        for cell in row:
            if cell.__class__.__name__ != "MergedCell" and cell.value is not None:
                cell.border = border


def _style_approval_access_duration_sheet(ws, border: Border) -> None:
    """5.1-like sheet: two light-gray group rows and a darker-gray field header row."""
    for row_idx in (1, 2):
        if row_idx > ws.max_row:
            continue
        for cell in ws[row_idx]:
            if cell.__class__.__name__ != "MergedCell" and cell.value not in (None, ""):
                _style_cell(cell, fill=TITLE_FILL, bold=True, color=TEXT, horizontal="center")
                cell.border = border
    if ws.max_row >= 3:
        for cell in ws[3]:
            if cell.__class__.__name__ != "MergedCell" and cell.value not in (None, ""):
                _style_cell(cell, fill=HEADER_FILL, bold=True, color=TEXT, horizontal="center")
                cell.border = border


def apply_excel_style(workbook) -> None:
    """Apply a quiet, consistent style without changing workbook structure/values."""
    thin = Side(style="thin", color=BORDER)
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for ws in workbook.worksheets:
        ws.sheet_view.showGridLines = False

        # Base typography/borders + conversion of old vivid generator fills.
        for row in ws.iter_rows():
            for cell in row:
                if cell.__class__.__name__ == "MergedCell" or cell.value is None:
                    continue
                rgb = _rgb_value(cell)
                replacement = _FILL_REPLACEMENTS.get(rgb) or _legacy_theme_fill(cell)
                if replacement:
                    cell.fill = _solid_fill(replacement)
                font = copy(cell.font)
                font.name = "Segoe UI"
                font.sz = 10
                # Old title/header styles often had white text for dark/orange fills.
                if replacement:
                    font.color = TEXT
                cell.font = font
                cell.border = border
                alignment = copy(cell.alignment)
                alignment.vertical = "center"
                alignment.wrap_text = True
                cell.alignment = alignment

        # Approval-route sheets get the same palette/layout language as the supplied reference.
        if _is_approval_flags_sheet(ws):
            _style_approval_flags_sheet(ws, border)
            continue
        if _is_approval_access_duration_sheet(ws):
            _style_approval_access_duration_sheet(ws, border)
            continue
        if _has_stage_groups_in_first_row(ws):
            _style_approval_main_sheet(ws, border)
            continue

        first_values = _row_values(ws, 1) if ws.max_row else []
        header_row = _header_row(ws)

        # A single first-row value is a sheet title, not a table header.
        if len(first_values) == 1:
            for cell in ws[1]:
                if cell.value not in (None, ""):
                    _style_cell(cell, fill=TITLE_FILL, bold=True, font_size=13, color=TEXT, vertical="center")
            ws.row_dimensions[1].height = max(ws.row_dimensions[1].height or 0, 24)

            # A one-value subtitle directly below the title gets a very light fill.
            if ws.max_row >= 2 and _nonempty_count(ws, 2) == 1 and header_row != 2:
                for cell in ws[2]:
                    if cell.value not in (None, ""):
                        _style_cell(cell, fill=NOTE_FILL, bold=False, font_size=9, color=MUTED, vertical="center")

        if header_row is not None:
            for cell in ws[header_row]:
                if cell.value in (None, ""):
                    continue
                # Required-field yellow and blocked gray carry meaning and stay intact.
                rgb = _rgb_value(cell)
                fill = rgb if rgb in {REQUIRED_FILL, BLOCKED_FILL} else HEADER_FILL
                _style_cell(cell, fill=fill, bold=True, color=TEXT, horizontal="center", vertical="center")
                cell.border = border
            ws.row_dimensions[header_row].height = max(ws.row_dimensions[header_row].height or 0, 28)

        # Some approval settings sheets have one more header tier with stage names
        # below the first grouped header. Keep that tier in the same quiet gray.
        for row_idx in range(2, min(ws.max_row, 4) + 1):
            values = _row_values(ws, row_idx)
            if any(re.fullmatch(r"\d+\s*Этап", value, flags=re.IGNORECASE) for value in values):
                if row_idx != header_row:
                    for cell in ws[row_idx]:
                        if cell.value in (None, ""):
                            continue
                        _style_cell(cell, fill=HEADER_FILL, bold=True, color=TEXT, horizontal="center", vertical="center")
                        cell.border = border
                    ws.row_dimensions[row_idx].height = max(ws.row_dimensions[row_idx].height or 0, 28)
