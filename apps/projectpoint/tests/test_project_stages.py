from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from projectpoint.api.project_stages import create_project_stage, update_project_stage
from projectpoint.excel.project_stages import project_stage_payload_equal, read_project_stages_excel
from projectpoint.excel.templates import write_project_stages_import_template


class _Response:
    def __init__(self, status_code=200, data=None):
        self.status_code = status_code
        self._data = data
        self.text = ""

    def json(self):
        return self._data


class _Session:
    def __init__(self):
        self.posts = []
        self.verify = True

    def post(self, url, **kwargs):
        self.posts.append((url, kwargs.get("json")))
        return _Response(200, {"Id": "new-id"})


class ProjectStagesTests(unittest.TestCase):
    def test_template_roundtrip(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "stages.xlsx"
            write_project_stages_import_template(str(path))
            wb = load_workbook(path, read_only=True)
            self.assertIn("Виды документов", wb.sheetnames)
            rows = read_project_stages_excel(str(path), "Виды документов")
            self.assertEqual(rows[0]["Code"], "РД")
            self.assertTrue(rows[0]["IsActive"])
            self.assertEqual(rows[0]["DisplayParameters"], 1)

    def test_duplicate_code_is_rejected(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "stages.xlsx"
            write_project_stages_import_template(str(path))
            from openpyxl import load_workbook as load
            wb = load(path)
            ws = wb["Виды документов"]
            ws.append(["рд", "Дубликат", 1, True])
            wb.save(path)
            with self.assertRaisesRegex(ValueError, "уже указан"):
                read_project_stages_excel(str(path), "Виды документов")

    def test_payload_equality(self):
        existing = {"Code": "РД", "Title": "Рабочая документация", "DisplayParameters": 1, "IsActive": True}
        desired = dict(existing)
        self.assertTrue(project_stage_payload_equal(existing, desired))
        desired["Title"] = "Новое имя"
        self.assertFalse(project_stage_payload_equal(existing, desired))

    def test_create_and_update_endpoints(self):
        sess = _Session()
        payload = {"Code": "001", "Title": "Тест", "DisplayParameters": 1, "IsActive": True}
        ok, status, _ = create_project_stage(sess, "token", "https://host", payload)
        self.assertTrue(ok)
        self.assertEqual(status, 200)
        self.assertTrue(sess.posts[0][0].endswith("/ru/api/Core/ProjectStageService/Create"))

        update = {"Id": "id-1", **payload}
        ok, _, _ = update_project_stage(sess, "token", "https://host", update)
        self.assertTrue(ok)
        self.assertTrue(sess.posts[1][0].endswith("/ru/api/Core/ProjectStageService/Update"))


if __name__ == "__main__":
    unittest.main()
