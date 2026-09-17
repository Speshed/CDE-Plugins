import io
import sys
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "app/plugins/larix"))
import Larix_User_Platform as larix


TEMPLATES = [
    (larix.STRUCTURE_TEMPLATE_B64, larix.STRUCTURE_TEMPLATE_SHA256, "structure"),
    (larix.USER_IMPORT_TEMPLATE_B64, larix.USER_IMPORT_TEMPLATE_SHA256, "users"),
    (larix.ROLE_MATRIX_TEMPLATE_B64, larix.ROLE_MATRIX_TEMPLATE_SHA256, "role"),
    (larix.REMARK_TYPE_TEMPLATE_B64, larix.REMARK_TYPE_TEMPLATE_SHA256, "remarks"),
    (larix.TASK_TYPE_TEMPLATE_B64, larix.TASK_TYPE_TEMPLATE_SHA256, "tasks"),
    (larix.APPROVAL_TEMPLATE_B64, larix.APPROVAL_TEMPLATE_SHA256, "approval"),
]


def _signature(book):
    return [
        (
            ws.title,
            ws.sheet_state,
            tuple(tuple(cell.value for cell in row) for row in ws.iter_rows()),
            tuple(str(rng) for rng in ws.merged_cells.ranges),
            tuple(str(dv.sqref) for dv in ws.data_validations.dataValidation),
            ws.freeze_panes,
            ws.auto_filter.ref,
        )
        for ws in book.worksheets
    ]


def test_all_larix_download_templates_preserve_structure_and_style(tmp_path):
    for encoded, sha, name in TEMPLATES:
        original = load_workbook(io.BytesIO(larix._decode_embedded_template(encoded, sha)), data_only=False)
        target = tmp_path / f"{name}.xlsx"
        saved = larix._write_embedded_template(encoded, sha, target)
        assert saved.endswith(".xlsx")
        styled = load_workbook(saved, data_only=False)
        assert _signature(original) == _signature(styled)
        assert any(cell.font.name == "Segoe UI" for ws in styled.worksheets for row in ws.iter_rows() for cell in row if cell.value is not None)


def test_renamed_role_and_approval_use_explicit_style_ids(tmp_path):
    role = load_workbook(
        larix._write_embedded_template(
            larix.ROLE_MATRIX_TEMPLATE_B64,
            larix.ROLE_MATRIX_TEMPLATE_SHA256,
            tmp_path / "renamed-role.xlsx",
            "role",
        )
    )
    sheet = role.worksheets[0]
    assert sheet["A1"].fill.fgColor.rgb.endswith("2F2F2F")
    assert sheet["A1"].font.name == "Segoe UI"
    assert sheet["A5"].fill.fgColor.rgb.endswith("F7921E")
    assert sheet["Q5"].fill.fgColor.rgb.endswith("F7921E")
    legend = [cell for row in sheet["S1:T5"] for cell in row if cell.value is not None]
    assert legend
    assert all(cell.fill.fgColor.rgb.endswith("FFF4E8") for cell in legend)
    assert all(cell.font.name == "Segoe UI" for cell in legend)
    assert all(cell.font.color.type == "rgb" and cell.font.color.rgb.endswith("666666") for cell in legend)

    approval = load_workbook(
        larix._write_embedded_template(
            larix.APPROVAL_TEMPLATE_B64,
            larix.APPROVAL_TEMPLATE_SHA256,
            tmp_path / "renamed-approval.xlsx",
            "approval",
        )
    )
    for ws in approval.worksheets:
        row_one = [cell for cell in ws[1] if cell.value is not None]
        assert row_one, ws.title
        assert all(cell.fill.fgColor.rgb.endswith("2F2F2F") for cell in row_one)
        assert all(cell.font.name == "Segoe UI" for cell in row_one)
        row_two = [cell for cell in ws[2] if cell.value is not None]
        assert row_two, ws.title
        assert all(cell.fill.fgColor.rgb.endswith("F7921E") for cell in row_two)
        assert all(cell.font.name == "Segoe UI" for cell in row_two)
