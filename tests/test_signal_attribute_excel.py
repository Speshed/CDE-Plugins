from pathlib import Path
import sys

from openpyxl import Workbook, load_workbook


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "app/plugins/signal"))
import SIGNAL


def _book(tmp_path, rows, *, headers=SIGNAL.ATTRIBUTE_EXCEL_HEADERS):
    path = tmp_path / "attributes.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "Атрибуты"
    for column, value in enumerate(headers, 1):
        sheet.cell(1, column, value)
    for row, values in enumerate(rows, 2):
        for column, value in enumerate(values, 1):
            sheet.cell(row, column, value)
    book.save(path)
    return path


def test_generated_attributes_template_parses_all_demo_rows(tmp_path):
    path = tmp_path / "generated.xlsx"
    SIGNAL.build_excel_template("attributes", str(path))

    result = SIGNAL.parse_attribute_excel(path)

    assert len(result.rows) == len(result.drafts) == 4
    assert result.errors == []
    assert [row["excel_row"] for row in result.rows] == [4, 5, 6, 7]
    assert result.drafts[-1]["type"] == "List"
    assert result.drafts[-1]["values"] == ["АР", "КР", "ОВ", "ВК"]


def test_parser_normalizes_drafts_and_list_values(tmp_path):
    path = _book(tmp_path, [["  Новый атрибут  ", " Да / Нет ", "  Без области ", "Да", ""]])

    result = SIGNAL.parse_attribute_excel(path)

    assert result.errors == []
    assert result.drafts == [{
        "name": "Новый атрибут",
        "type": "Bool",
        "dataType": "",
        "isRequired": True,
        "values": [],
    }]

    path = _book(tmp_path, [["Статус", "Список", "Папка", "Нет", " Новый;новый\nГотов "]])
    result = SIGNAL.parse_attribute_excel(path)
    assert result.drafts[0]["values"] == ["Новый", "Готов"]


def test_empty_workbook_and_missing_sheet_are_file_errors(tmp_path):
    path = _book(tmp_path, [])
    result = SIGNAL.parse_attribute_excel(path)
    assert result.rows == []
    assert any("нет ни одной непустой строки" in error for error in result.errors)

    path = tmp_path / "wrong-sheet.xlsx"
    book = Workbook()
    book.active.title = "Другое"
    book.save(path)
    result = SIGNAL.parse_attribute_excel(path)
    assert any("отсутствует лист" in error for error in result.errors)


def test_missing_headers_and_invalid_rows_are_reported_and_excluded(tmp_path):
    path = _book(tmp_path, [["A", "Список", "Папка", "Да", ""]], headers=("Название", "Тип"))
    result = SIGNAL.parse_attribute_excel(path)
    assert result.drafts == []
    assert any("обязательные заголовки" in error for error in result.errors)

    path = _book(tmp_path, [
        ["Без типа", "", "Папка", "Да", ""],
        ["Без вариантов", "Список", "Папка", "Нет", ""],
        ["Лишние варианты", "Текст", "Папка", "Нет", "x"],
        ["Плохая область", "Дата", "Раздел", "Нет", ""],
        ["Плохая обязательность", "Дата", "Папка", "Иногда", ""],
        ["", "Дата", "Папка", "Нет", ""],
    ])
    result = SIGNAL.parse_attribute_excel(path)
    assert len(result.rows) == 6
    assert result.drafts == []
    assert all(row["status"] == "Ошибка" and row["draft"] is None for row in result.rows)


def test_duplicate_names_are_case_insensitive(tmp_path):
    path = _book(tmp_path, [
        ["Название", "Текст", "Папка", "Нет", ""],
        [" название ", "Дата", "Папка", "Нет", ""],
    ])

    result = SIGNAL.parse_attribute_excel(path)

    assert len(result.drafts) == 1
    assert result.rows[1]["status"] == "Ошибка"
    assert any("повтор названия" in error for error in result.errors)


def test_parser_releases_source_workbook_handle(tmp_path):
    source = _book(tmp_path, [["Атрибут", "Текст", "Папка", "Нет", ""]])
    result = SIGNAL.parse_attribute_excel(source)
    assert result.errors == []

    renamed = tmp_path / "renamed.xlsx"
    source.rename(renamed)
    renamed.unlink()
    assert not source.exists()
