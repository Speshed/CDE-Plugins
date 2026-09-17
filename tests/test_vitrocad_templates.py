import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl import Workbook
from openpyxl.styles import PatternFill

ROOT = Path(__file__).resolve().parents[1] / "app"
VITROCAD = ROOT / "plugins" / "vitrocad"
if str(VITROCAD) not in sys.path:
    sys.path.insert(0, str(VITROCAD))

from templates import write_vitrocad_template
from modules.folder_creator import ExcelStructureParser
from modules.permissions import PermissionExcelParser
from modules.schedule_sync import ExcelReader
from modules.schedule_sync import FIELD_ALIASES, find_column
from shared.excel_style import apply_excel_style


def test_excel_style_keeps_role_header_and_required_fill():
    book = Workbook()
    sheet = book.active
    sheet.append(["Роль 1", "Роль 2", "Роль 3"])
    sheet.append(["Просмотр", "", ""])
    sheet.cell(1, 1).fill = PatternFill("solid", fgColor="FFF2CC")
    sheet.cell(1, 2).fill = PatternFill("solid", fgColor="D9EAF7")
    sheet.cell(2, 1).fill = PatternFill("solid", fgColor="F2F2F2")
    apply_excel_style(book)
    assert sheet.cell(1, 1).fill.fgColor.rgb.endswith("FFF2CC")
    assert sheet.cell(1, 2).fill.fgColor.rgb.endswith("F7921E")
    assert sheet.cell(1, 1).font.bold
    assert sheet.cell(2, 1).font.bold is not True
    assert sheet.cell(2, 1).fill.fgColor.rgb.endswith("F2F2F2")


def test_excel_style_marks_title_only_sheet():
    book = Workbook()
    sheet = book.active
    sheet.append(["Инструкция"])
    sheet.append(["Текст инструкции"])
    apply_excel_style(book)
    assert sheet["A1"].fill.fgColor.rgb.endswith("2F2F2F")


def test_vitrocad_folder_template_matches_parser(tmp_path):
    path = tmp_path / "folders.xlsx"
    write_vitrocad_template("folders", path)
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        assert book.active.cell(1, 1).value == "Уровень 1"
    finally:
        book.close()
    nodes = ExcelStructureParser(str(path)).parse()
    assert nodes and nodes[-1].path[-1] == "Рабочая документация"


def test_vitrocad_permissions_template_matches_parser(tmp_path):
    path = tmp_path / "permissions.xlsx"
    write_vitrocad_template("permissions", path)
    entries = PermissionExcelParser(str(path)).parse()
    assert entries and entries[0].path_parts == ["01. Проектирование", "АР"]


def test_vitrocad_schedule_template_matches_parser(tmp_path):
    path = tmp_path / "schedule.xlsx"
    write_vitrocad_template("schedule", path)
    sheets = ExcelReader.list_sheets(str(path))
    assert sheets == ["План-график"]
    data = ExcelReader.read(str(path), sheets[0])
    nodes, warnings = ExcelReader.prepare_nodes(data)
    assert not warnings
    assert nodes[-1].path == ("01. Проектирование", "Этап 1")
    for field, expected in {
        "name": "Название",
        "content_type_id": "Тип задачи",
        "assignedto": "Исполнитель",
        "task_status": "Статус задачи",
        "start_date_plan": "Дата начала (План)",
        "end_date_plan": "Дата окончания (План)",
    }.items():
        assert find_column(data.dataframe.columns, FIELD_ALIASES[field]) == expected
